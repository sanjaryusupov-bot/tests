"""
Дашборд сборки заказов: Google Таблица -> Streamlit -> Telegram WebApp.
Дизайн: «M Cosmetic Dark v2» (меню с иконками, поиск, KPI с мини-графиками, неоновый градиент).
Один файл. Запуск локально:  streamlit run app.py
Секреты (Streamlit Cloud -> Settings -> Secrets): блоки [gcp_service_account] и [metabase]
(в [metabase] можно добавить hub_table = 61, stock_table, magnit_table).
requirements.txt: streamlit, gspread, pandas, google-auth, openpyxl, requests, altair
Логотипы: положите logo.png (M Cosmetic) и logo2.png (партнёр) рядом с app.py.
Файлы на сервере (необязательные): change_history.json (колокольчик), visitors.json (окно «Посетители»).
"""
import hashlib
import threading
import time
import uuid
import zlib
from urllib.parse import quote, unquote
import base64
import html
import json
import re
from datetime import datetime, date, timedelta
from io import BytesIO
from pathlib import Path
from string import Template
from zoneinfo import ZoneInfo

import gspread
import numpy as np
import pandas as pd
import requests
import streamlit as st
import streamlit.components.v1 as components
from google.oauth2.service_account import Credentials

# ───────────────────────── Настройки ─────────────────────────
SHEET_ID = "12CxJMCBUHgkaj-_KbOs1aK7hx_jEYMyS3q4Hh0bHTGw"
GID = "1369918403"
NEW_GID = "750257398"        # вкладка «Новые заказы» (не маршрутизированные)
CACHE_SECONDS = 120          # как часто перечитывать таблицу
TZ = ZoneInfo("Asia/Tashkent")
HIDE_FROM_DAYS_PCT = 95      # в сводке по умолчанию прячем дни с готовностью от этого процента

# Логотипы: файлы рядом с app.py (PNG / SVG / JPG)
LOGO_FILES = ["logo.png", "logo.svg", "logo.jpg"]        # M Cosmetic
LOGO2_FILES = ["logo2.png", "logo2.svg", "logo2.jpg"]    # партнёр
LOGO_HEIGHT = 44      # высота окна с логотипом, px (хотите крупнее — 56 или 64)
LOGO_HOLD = 3.5       # сколько секунд показывать каждый логотип
LOGO_SLIDE = 0.7      # сколько секунд длится сдвиг вверх

# Магазины, которых НЕ должно быть в сводной «Маршрутизация»
EXCLUDE_STORES = [
    "Узуммаркет",
    "Яндекс Маркет МСА",
    "Е-ком МСА",
    "Импортные поставки",
    "Вайлдберриз МСА",
]

ALL = "all"   # «все даты» в выпадающем списке
# Архив отгрузок: по нему считаем, сколько штук в среднем уезжало на каждой машине
ARCHIVE_SHEET_ID = "1hKZ8ggNLW-OY1bV8xAW7PKl50Fof2co86oxGK92YPAA"
ARCHIVE_GID = "2106470004"
ARCHIVE_TAB = "Архив"
ARCHIVE_FROM = date(2026, 10, 1)
# Координаты магазинов (вкладка «Кординаты» в той же таблице)
COORDS_GID = "1482916811"

# Сохранение ручных рейсов: вкладка в основной таблице (нужен доступ «Редактор» для сервисного аккаунта)
PLAN_TAB = "Рейсы_план"
PLAN_FILE = "trip_plans.json"   # запасной вариант: файл на сервере

# Пустые короба (ответы формы)
BOXES_SHEET_ID = "1FZmR7vUXAWFBJgkrCHhM_y8LoAeSjutJLHZjTJRdYw4"
BOXES_GID = "2073580233"
BOXES_TAB = "Ответы на форму (1)"
MONTHS = ["январь", "февраль", "март", "апрель", "май", "июнь",
          "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]

# Категория и приоритет заказов (для Узуммаркет): вкладка «Лист1»
CAT_SHEET_ID = "1qaIGoF4PYPV2GZyvCN7D_kRQlchan6ZlAz2sl77TIhQ"
CAT_GID = "0"

# Metabase: ID таблиц (можно переопределить в Secrets: stock_table, magnit_table, hub_table)
MB_STOCK_T = 48     # «Stock Invoice Sku Cell» — история снимков остатков
MB_MAGNIT_T = 58    # «Magnit Stock»
MB_HUB_T = 61       # «Хаб отгрузка ББ» — приёмка коробок (биг-бэгов) на хабе
MB_STOCK_DAYS = 90  # за сколько последних дней можно выбирать период
MB_CAT_DAYS = 14    # за сколько последних дней собирать каталог товаров
MB_NOW_DAYS = 10    # (запасной параметр)
HUB_DAYS = 30       # за сколько последних дней читать «Хаб отгрузка ББ»
HUB_CARD_DATES = 3  # сколько дат показывать мелким шрифтом на карточке хаба

# Зоны хранения: где товар можно резервировать под заказы, а где нет
NO_RESERVE_ZONES = ["BP", "1DMG", "1M-BUF", "2DMG"]
RESERVE_ZONES = ["5M1", "1M1", "2M1", "1M2", "1ECO1", "1ZZ1", "3M1", "4M1", "HP2", "HP"]

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MODES = ["Все", "Не собрано", "Собрано", "Прочие заказы", "Маршрутизация", "Рейсы"]
ORDINALS = ["Первая", "Вторая", "Третья", "Четвёртая", "Пятая", "Шестая", "Седьмая", "Восьмая", "Девятая", "Десятая"]

PAGES = ["Сводка по дням", "Заказы", "Поиск заказа", "Пустые короба", "Движения по остаткам", "Аналитика"]
NAV_ICONS = {"Сводка по дням": "🏠", "Заказы": "📋", "Поиск заказа": "🔍",
             "Пустые короба": "📦", "Движения по остаткам": "🔄", "Аналитика": "📊"}
WEEKDAYS = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]
# Контакты внизу бокового меню
CONTACT_NAME = "Контакты"
CONTACT_PHONE = "+998900230423"
CONTACT_TG_URL = "https://t.me/sanyspv"
CONTACT_TG = "@sanyspv"
ADDRESS_URL = "https://yandex.uz/maps/-/CXb1NI87"
ADDRESS_TEXT = "просп. Ханабадтепа, 2/1"
LOGO_TOP_HEIGHT = 32  # высота логотипа в верхней панели, px
HISTORY_FILE = "change_history.json"   # история изменений для колокольчика (запасной файл на сервере)
HISTORY_KEEP = 300                      # сколько последних событий хранить
VISITORS_FILE = "visitors.json"         # история посетителей для окна «онлайн» (файл на сервере)
VIS_KEEP_DAYS = 30                      # сколько дней хранить посетителей и их сессии

# Палитра «M Cosmetic Dark»: тёмно-синий фон, малиново-фиолетовый неоновый градиент
LIGHT = dict(scheme="light", ink="#1B1B22", paper="#F4F4F6", card="#FFFFFF", line="#ECE7EA",
             mute="#6F6B76", done="#1E9E5C", left="#E8A33D", track="#EEE9EC", head="#F8F3F6",
             ok_bg="#E3F6EC", ok_fg="#1B7A48", no_bg="#FDEBD3", no_fg="#8A5206",
             car_bg="#FDE7F3", car_fg="#8E1B5E", band="#C2308B", hot_bg="#FFF1DA", note_bg="#FFF4D6",
             accent="#C2308B", accent_soft="#FDE7F3", accent_ink="#A8206F",
             side="#FFFFFF", g1="#C2308B", g2="#8B3DFF", glow="rgba(194,48,139,.12)",
             kpi_a="#FDE7F3", kpi_b="#FFF1DA", kpi_c="#E3EEFF", row="#FFFFFF")
DARK = dict(scheme="dark", ink="#EEF0FA", paper="#070A1A", card="#0D1126", line="#1E2548",
            mute="#8C93B5", done="#34C27A", left="#F2A23D", track="#1A1F3D", head="#121733",
            ok_bg="#10341F", ok_fg="#7FE3AE", no_bg="#3B2A0F", no_fg="#F5C77E",
            car_bg="#2A1535", car_fg="#F5A9D6", band="#8B3DFF", hot_bg="#33260F", note_bg="#33260F",
            accent="#E256A8", accent_soft="#3A1F4A", accent_ink="#F5A9D6",
            side="#090C1E", g1="#E0338F", g2="#8B3DFF", glow="rgba(224,51,143,.35)",
            kpi_a="#3A1255", kpi_b="#3A260F", kpi_c="#0F2A55", row="#10142B")

st.set_page_config(page_title="M Cosmetic · Сборка заказов", page_icon="📦", layout="wide",
                   initial_sidebar_state="expanded")

# ───────────────────────── Тема (день / ночь) ─────────────────────────
if "dark" not in st.session_state:
    st.session_state.dark = True          # тёмная тема по умолчанию (переключается кнопкой)
P = DARK if st.session_state.dark else LIGHT

# ───────────────────────── Иконки (линейные, как в макете) ─────────────────────────
_ICONS = {
    "home": "<path d='M3 10.5 12 3l9 7.5V20a1 1 0 0 1-1 1h-5v-6H9v6H4a1 1 0 0 1-1-1z'/>",
    "box": "<path d='m7.5 4.27 9 5.15'/><path d='M21 8a2 2 0 0 0-1-1.73l-7-4a2 2 0 0 0-2 0l-7 4A2 2 0 0 0 3 8v8a2 2 0 0 0 1 1.73l7 4a2 2 0 0 0 2 0l7-4A2 2 0 0 0 21 16Z'/><path d='m3.3 7 8.7 5 8.7-5'/><path d='M12 22V12'/>",
    "list": "<rect x='8' y='2' width='8' height='4' rx='1'/><path d='M16 4h2a2 2 0 0 1 2 2v14a2 2 0 0 1-2 2H6a2 2 0 0 1-2-2V6a2 2 0 0 1 2-2h2'/><path d='M12 11h4M12 16h4M8 11h.01M8 16h.01'/>",
    "search": "<circle cx='11' cy='11' r='8'/><path d='m21 21-4.3-4.3'/>",
    "swap": "<path d='M8 3 4 7l4 4'/><path d='M4 7h16'/><path d='m16 21 4-4-4-4'/><path d='M20 17H4'/>",
    "chart": "<path d='M3 3v16a2 2 0 0 0 2 2h16'/><path d='M18 17V9M13 17V5M8 17v-3'/>",
    "sun": "<circle cx='12' cy='12' r='4'/><path d='M12 2v2M12 20v2M4.93 4.93l1.41 1.41M17.66 17.66l1.41 1.41M2 12h2M20 12h2M6.34 17.66l-1.41 1.41M19.07 4.93l-1.41 1.41'/>",
    "moon": "<path d='M12 3a6 6 0 0 0 9 9 9 9 0 1 1-9-9Z'/>",
    "bell": "<path d='M6 8a6 6 0 0 1 12 0c0 7 3 9 3 9H3s3-2 3-9'/><path d='M10.3 21a1.94 1.94 0 0 0 3.4 0'/>",
    "calendar": "<rect x='3' y='4' width='18' height='18' rx='2'/><path d='M16 2v4M8 2v4M3 10h18'/>",
    "chev": "<path d='m9 18 6-6-6-6'/>",
    "download": "<path d='M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4'/><path d='m7 10 5 5 5-5'/><path d='M12 15V3'/>",
    "check": "<circle cx='12' cy='12' r='10'/><path d='m9 12 2 2 4-4'/>",
    "refresh": "<path d='M3 12a9 9 0 0 1 9-9 9.75 9.75 0 0 1 6.74 2.74L21 8'/><path d='M21 3v5h-5'/><path d='M21 12a9 9 0 0 1-9 9 9.75 9.75 0 0 1-6.74-2.74L3 16'/><path d='M8 16H3v5'/>",
    "clock": "<circle cx='12' cy='12' r='10'/><path d='M12 6v6l4 2'/>",
    "pin": "<path d='M20 10c0 6-8 12-8 12S4 16 4 10a8 8 0 0 1 16 0Z'/><circle cx='12' cy='10' r='3'/>",
    "send": "<path d='m22 2-7 20-4-9-9-4Z'/><path d='M22 2 11 13'/>",
    "phone": "<path d='M22 16.92v3a2 2 0 0 1-2.18 2 19.79 19.79 0 0 1-8.63-3.07 19.5 19.5 0 0 1-6-6 19.79 19.79 0 0 1-3.07-8.67A2 2 0 0 1 4.11 2h3a2 2 0 0 1 2 1.72c.13.96.36 1.9.7 2.81a2 2 0 0 1-.45 2.11L8.09 9.91a16 16 0 0 0 6 6l1.27-1.27a2 2 0 0 1 2.11-.45c.91.34 1.85.57 2.81.7A2 2 0 0 1 22 16.92Z'/>",
    "users": "<path d='M16 21v-2a4 4 0 0 0-4-4H6a4 4 0 0 0-4 4v2'/><circle cx='9' cy='7' r='4'/><path d='M22 21v-2a4 4 0 0 0-3-3.87'/><path d='M16 3.13a4 4 0 0 1 0 7.75'/>",
    "timer": "<path d='M10 2h4'/><path d='M12 14l3-3'/><circle cx='12' cy='14' r='8'/>",
    "hourglass": "<path d='M5 22h14'/><path d='M5 2h14'/><path d='M17 22v-4.172a2 2 0 0 0-.586-1.414L12 12l-4.414 4.414A2 2 0 0 0 7 17.828V22'/><path d='M7 2v4.172a2 2 0 0 0 .586 1.414L12 12l4.414-4.414A2 2 0 0 0 17 6.172V2'/>",
    "monitor": "<rect width='20' height='14' x='2' y='3' rx='2'/><path d='M8 21h8M12 17v4'/>",
    "smartphone": "<rect width='14' height='20' x='5' y='2' rx='2' ry='2'/><path d='M12 18h.01'/>",
    "globe": "<circle cx='12' cy='12' r='10'/><path d='M12 2a14.5 14.5 0 0 0 0 20 14.5 14.5 0 0 0 0-20'/><path d='M2 12h20'/>",
    "history": "<path d='M3 12a9 9 0 1 0 9-9 9.75 9.75 0 0 0-6.74 2.74L3 8'/><path d='M3 3v5h5'/><path d='M12 7v5l4 2'/>",
}
NAV_ICON_NAMES = ["home", "list", "search", "box", "swap", "chart"]


def _svg(name, color="currentColor", size=20, sw=2):
    return (f"<svg xmlns='http://www.w3.org/2000/svg' width='{size}' height='{size}' viewBox='0 0 24 24' fill='none' "
            f"stroke='{color}' stroke-width='{sw}' stroke-linecap='round' stroke-linejoin='round'>{_ICONS[name]}</svg>")


def ic(name, color="currentColor", size=20):
    """Иконка прямо в HTML."""
    return _svg(name, color, size)


def icon_uri(name, color="#000000"):
    """Иконка как data-URI (для CSS: фон или маска)."""
    return "data:image/svg+xml;utf8," + quote(_svg(name, color, 24), safe="")


def _icon_rule(selector, name, size=18):
    u = icon_uri(name)
    return (f'{selector}::before{{content:"";width:{size}px;height:{size}px;flex:none;background:currentColor;'
            f'-webkit-mask:url("{u}") center/contain no-repeat;mask:url("{u}") center/contain no-repeat;}}')


_EXCEL_SVG = ("<svg xmlns='http://www.w3.org/2000/svg' width='20' height='20' viewBox='0 0 20 20'>"
              "<rect width='20' height='20' rx='4' fill='#1D6F42'/>"
              "<path d='M6 6l8 8M14 6l-8 8' stroke='white' stroke-width='2' stroke-linecap='round'/></svg>")
SPARK_COL = {"all": "#E0338F", "left": "#F2A23D", "blue": "#3B9CFF", "ok": "#34C27A"}
KPI_ICON = {"all": "box", "left": "check", "blue": "box", "ok": "check"}


def spark_svg(vals, color, uid, w=104, h=36):
    v = [float(x) for x in vals]
    if len(v) < 2:
        return ""
    lo, hi = min(v), max(v)
    rng = (hi - lo) or 1.0
    xs = [i * (w - 4) / (len(v) - 1) + 2 for i in range(len(v))]
    ys = [h - 4 - (x - lo) / rng * (h - 10) for x in v]
    d = f"M{xs[0]:.1f},{ys[0]:.1f}"
    for i in range(1, len(v)):
        mx = (xs[i - 1] + xs[i]) / 2
        d += f" C{mx:.1f},{ys[i - 1]:.1f} {mx:.1f},{ys[i]:.1f} {xs[i]:.1f},{ys[i]:.1f}"
    area = d + f" L{xs[-1]:.1f},{h} L{xs[0]:.1f},{h} Z"
    return (f"<svg width='{w}' height='{h}' viewBox='0 0 {w} {h}'><defs>"
            f"<linearGradient id='sg{uid}' x1='0' y1='0' x2='0' y2='1'>"
            f"<stop offset='0' stop-color='{color}' stop-opacity='.45'/>"
            f"<stop offset='1' stop-color='{color}' stop-opacity='0'/></linearGradient></defs>"
            f"<path d='{area}' fill='url(#sg{uid})'/>"
            f"<path d='{d}' fill='none' stroke='{color}' stroke-width='2' stroke-linecap='round'/></svg>")


def trend_pct(vals, k=7):
    """Изменение суммы за последние k значений к предыдущим k (в процентах)."""
    vals = [float(x) for x in vals]
    if len(vals) < 2:
        return None
    k = min(k, len(vals) // 2)
    a, b = sum(vals[-k:]), sum(vals[-2 * k:-k])
    return (a - b) / b * 100 if b else None


WMS_STATUSES = [("PACKING", "#3B9CFF"), ("SORTED", "#34C27A"), ("SORTING", "#F2A23D"),
                ("PARTIALLY_RESERVED", "#E0338F"), ("RESERVED", "#8B3DFF")]


def status_breakdown(frame) -> list:
    """Сколько штук в каждом статусе WMS: [(статус, штук, цвет)]. Остальные статусы — одной строкой «ДРУГИЕ»."""
    s = (frame["Статус WMS"].astype(str).str.strip().str.upper().str.replace(r"[\s\-]+", "_", regex=True))
    out = [(name, int(frame.loc[s == name, "Штук"].sum()), col) for name, col in WMS_STATUSES]
    other = int(frame.loc[~s.isin([x for x, _ in WMS_STATUSES]), "Штук"].sum())
    if other:
        out.append(("ДРУГИЕ", other, "#8C93B5"))
    return out


# ───────────────────────── Грузовая «Газель» на линии отгрузки ─────────────────────────
# CSS общий: подключается и в основную страницу, и в карточку-«iframe» с таймером.
TRUCK_CSS = """
.lane { position:relative; padding-top:22px; margin-top:2px; }
.truck { position:absolute; top:0; left:clamp(54px, var(--p, 0%), 100%); width:50px; height:26px;
         transform:translateX(-100%); z-index:3; pointer-events:none;
         animation:tarrive 2.4s cubic-bezier(.2,.7,.2,1) both; }
.truck svg { display:block; width:100%; height:100%; overflow:visible; animation:tbob .42s ease-in-out infinite; }
.truck .wh { animation:wspin .5s linear infinite; }
.truck .sp { animation:tspeed .55s linear infinite; }
.truck.stop svg, .truck.stop .wh { animation:none; }
.truck.stop .sp { display:none; }
@keyframes tarrive { from { left:54px; } }
@keyframes tbob { 0%, 100% { transform:translateY(0); } 50% { transform:translateY(-1px); } }
@keyframes wspin { to { transform:rotate(360deg); } }
@keyframes tspeed { 0% { opacity:0; transform:translateX(5px); } 40% { opacity:.9; } 100% { opacity:0; transform:translateX(-8px); } }
"""


def truck_svg() -> str:
    """Грузовая «Газель»: будка, кабина со скошенным лобовым, колёса крутятся, сзади «ветер»."""
    ink, g1, g2 = P["ink"], P["g1"], P["g2"]

    def wheel(cx):
        return (f"<g class='wh' style='transform-origin:{cx}px 27px'>"
                f"<circle cx='{cx}' cy='27' r='5.6' fill='#20232C' stroke='{ink}' stroke-width='.8'/>"
                f"<circle cx='{cx}' cy='27' r='2.2' fill='#C9CEDA'/>"
                f"<path d='M{cx} 22.6v8.8M{cx - 4.4} 27h8.8' stroke='#C9CEDA' stroke-width='1'/></g>")

    return ("<svg viewBox='0 0 64 34' xmlns='http://www.w3.org/2000/svg'>"
            f"<g class='sp' stroke='{g1}' stroke-width='1.3' stroke-linecap='round'>"
            "<path d='M-3 8h-9'/><path d='M-1 14h-11'/><path d='M-4 20h-8'/></g>"
            f"<rect x='1' y='24' width='62' height='3' rx='1' fill='{ink}'/>"
            f"<rect x='1.5' y='2.5' width='38' height='22' rx='2.5' fill='#FFFFFF' stroke='{ink}' stroke-width='1.3'/>"
            f"<rect x='2' y='14' width='37' height='4' fill='{g1}'/><rect x='2' y='18' width='37' height='1.6' fill='{g2}'/>"
            f"<path d='M41 8h8.5l6 8.5H61a2 2 0 0 1 2 2V24.5H41z' fill='{g1}' stroke='{ink}' "
            f"stroke-width='1.3' stroke-linejoin='round'/>"
            f"<path d='M43 10.2h5.6l4.4 6.3H43z' fill='#BFE3FF' stroke='{ink}' stroke-width='.9' stroke-linejoin='round'/>"
            "<rect x='60.8' y='19.5' width='2.2' height='3' rx='.8' fill='#FFE27A'/>"
            + wheel(12) + wheel(51) + "</svg>")


def truck_html(pct, stop: bool) -> str:
    """Машинка на линии: pct — на каком проценте стоит нос машины, stop — отгрузка закончилась (машина стоит)."""
    return f"<div class='truck{' stop' if stop else ''}' style='--p:{int(pct)}%'>{truck_svg()}</div>"


CSS = Template("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
html, body, .stApp, .stApp button, .stApp input { font-family: 'Inter', sans-serif; }
.stApp { background:$paper; color:$ink; color-scheme:$scheme; }
#MainMenu, footer, [data-testid="stDecoration"], [data-testid="stToolbarActions"],
[data-testid="stAppDeployButton"], [data-testid="stMainMenu"] { display:none !important; }
.block-container { padding: 1.2rem 1rem 3rem; max-width: 1100px; }

.stApp p, .stApp label, .stApp li, .stApp [data-testid="stWidgetLabel"] p,
.stApp [data-testid="stMarkdownContainer"] p { color:$ink; }
.stApp [data-testid="stCaptionContainer"], .stApp [data-testid="stCaptionContainer"] p { color:$mute; }

/* кнопки и поля */
.stApp .stButton > button, .stApp .stDownloadButton > button {
  background:$card; color:$ink; border:1px solid $line; border-radius:12px; font-weight:600; }
.stApp .stButton > button:hover, .stApp .stDownloadButton > button:hover { border-color:$accent; color:$accent; }
.stApp .stButton > button p, .stApp .stDownloadButton > button p { color:inherit; }
.stApp [data-testid="stButtonGroup"] button { background:$card; color:$ink; border:1px solid $line; font-weight:600; }
.stApp [data-testid="stButtonGroup"] button p { color:inherit; }
.stApp [data-testid="stButtonGroup"] button[aria-checked="true"] { background:$accent_soft; color:$accent_ink; border-color:$accent_soft; }
.stApp [data-baseweb="select"] > div, .stApp [data-baseweb="input"], .stApp [data-baseweb="base-input"], .stApp input {
  background:$card !important; color:$ink !important; border-color:$line !important; border-radius:12px !important; }
.stApp [data-baseweb="select"] * { color:$ink !important; }
.stApp [data-baseweb="select"] svg { fill:$mute !important; }
.stApp [data-baseweb="tag"] { background:$accent !important; }
.stApp [data-baseweb="tag"] * { color:#fff !important; }
[data-baseweb="popover"] ul, [data-baseweb="popover"] li, [data-baseweb="menu"] { background:$card !important; color:$ink !important; }
[data-baseweb="popover"] li * { color:$ink !important; }
[data-baseweb="popover"] li:hover, [data-baseweb="popover"] li[aria-selected="true"] { background:$accent_soft !important; }
.stApp [data-testid="stExpander"] { background:$card; border:1px solid $line; border-radius:16px; }
.stApp [data-testid="stExpander"] summary p { color:$ink; font-weight:700; }
.stApp [data-testid="stAlert"] { background:$card; border:1px solid $line; border-radius:12px; }
.stApp textarea { background:$card !important; color:$ink !important; border-color:$line !important; border-radius:12px !important; }
.stApp [data-testid="stFileUploaderDropzone"] { background:$card; border:1px dashed $line; border-radius:12px; }
.stApp [data-testid="stFileUploaderDropzone"] * { color:$ink; }

/* шапка: логотипы + дата */
.brandbar { background:$card; border:1px solid $line; border-radius:16px; padding:10px 16px; margin-bottom:14px;
            display:flex; align-items:center; justify-content:flex-end; }
.brandbar .live { font-size:.78rem; color:$mute; display:flex; align-items:center; gap:6px; white-space:nowrap; }
.brandbar .live .dot { width:8px; height:8px; border-radius:50%; background:#1E9E5C; flex:none;
                       animation:livepulse 1.4s ease-in-out infinite; }
@keyframes livepulse {
  0%   { opacity:1;   box-shadow:0 0 0 0 rgba(30,158,92,.55); }
  50%  { opacity:.35; box-shadow:0 0 0 6px rgba(30,158,92,0); }
  100% { opacity:1;   box-shadow:0 0 0 0 rgba(30,158,92,0); }
}
.brand { display:flex; align-items:center; gap:10px; color:$accent; cursor:pointer; }
.brand img { display:block; }
.brand .bi { display:flex; color:$accent; }
.brand .bn { font-weight:800; font-size:1.15rem; letter-spacing:.04em; color:$accent; }
.top { display:flex; flex-direction:column; gap:2px; margin-bottom:6px; }
.top .crumb { font-size:.8rem; color:$mute; }
.top h1 { font-size:1.75rem; font-weight:800; letter-spacing:-0.02em; margin:0; color:$ink; }
.top .upd { font-size:.8rem; color:$mute; }

/* карточки-кнопки: невидимая кнопка накрывает всю карточку */
[class*="st-key-tap_"] { position:relative; }
[class*="st-key-tap_"] [data-testid="stElementContainer"]:has(button) { position:absolute; inset:0; z-index:5; margin:0; width:100%; height:100%; }
[class*="st-key-tap_"] [data-testid="stButton"] { width:100%; height:100%; display:block; }
[class*="st-key-tap_"] button { display:block; width:100%; height:100%; min-height:0; opacity:0; cursor:pointer; padding:0; }
[class*="st-key-tap_"]:hover .kpi, [class*="st-key-tap_"]:hover .day { border-color:$accent; }
[class*="st-key-tap_"] iframe { display:block; border:0; }

.kpi { background:$card; border:1px solid $line; border-radius:16px; padding:14px 16px; min-height:112px; }
.kpi .v { font-size:2.1rem; font-weight:800; letter-spacing:-0.02em; line-height:1.1; color:$ink; }
.kpi .l { font-size:.88rem; font-weight:700; margin-top:4px; color:$ink; }
.kpi .s { font-size:.8rem; color:$mute; margin-top:2px; }
.kpi.left { border-left:5px solid $left; }
.kpi.all { border-left:5px solid $accent; }
.kpi.ok { border-left:5px solid $done; }

.bar { height:10px; border-radius:99px; background:$track; overflow:hidden; }
.bar > i { display:block; height:100%; background:$done; border-radius:99px; }
.day { background:$card; border:1px solid $line; border-radius:16px; padding:12px 14px; margin-bottom:2px;
       display:grid; grid-template-columns: 92px 1fr 90px; gap:14px; align-items:center; }
.day .d { font-weight:800; font-size:1.05rem; line-height:1.15; color:$ink; }
.day .d small { display:block; font-weight:500; color:$mute; font-size:.78rem; }
.day .meta { font-size:.82rem; color:$mute; margin-top:6px; }
.day .meta b { color:$ink; }
.day .pct { text-align:right; font-weight:800; font-size:1.25rem; color:$ink; }
.day .pct small { display:block; font-weight:500; font-size:.76rem; color:$mute; }
.day.full { border-color:$done; }
@media (max-width: 640px) {
  .day { grid-template-columns: 1fr auto; }
  .day .mid { grid-column: 1 / -1; order: 3; }
}

.badge { display:inline-block; padding:2px 10px; border-radius:99px; font-size:.76rem; font-weight:700; white-space:nowrap; }
.badge.ok { background:$ok_bg; color:$ok_fg; }
.badge.no { background:$no_bg; color:$no_fg; }

.tbl-wrap { max-height:68vh; overflow:auto; border:1px solid $line; border-radius:16px; background:$card; }
table.tbl { border-collapse:collapse; width:100%; font-size:.85rem; color:$ink; }
.tbl th { position:sticky; top:0; background:$head; text-align:left; padding:10px 12px; font-weight:700; white-space:nowrap; color:$ink; }
.tbl td { padding:8px 12px; border-top:1px solid $line; vertical-align:top; color:$ink; }
.tbl .num { text-align:right; font-variant-numeric:tabular-nums; white-space:nowrap; }
.tbl td.city { font-weight:800; }
.tbl td.shop { font-weight:700; }
.tbl tr.total td { font-weight:800; background:$head; border-top:2px solid $line; }
.bxbar { height:8px; min-width:110px; border-radius:99px; background:$track; overflow:hidden; }
.bxbar i { display:block; height:100%; background:$accent; border-radius:99px; }
.sec { font-weight:800; font-size:1.05rem; margin:18px 0 8px; color:$ink; }

.hit { background:$card; border:1px solid $line; border-radius:16px; padding:12px 14px; margin-bottom:8px; color:$ink; }
.hit .n { font-weight:800; font-size:1.05rem; }
.hit .s { color:$mute; font-size:.85rem; margin:2px 0 8px; }
.note { background:$note_bg; border-left:5px solid $left; border-radius:10px; padding:10px 14px; margin:4px 0 10px; font-size:.9rem; color:$ink; }
.cars-wrap { overflow-x:auto; border:1px solid $line; border-radius:16px; background:$card; }
table.cars { border-collapse:collapse; width:100%; font-size:.88rem; color:$ink; }
.cars td { padding:6px 10px; vertical-align:top; color:$ink; border-top:1px solid $line; }
.cars tr.car td { border-top:none; }
.cars td.carname { background:$car_bg; color:$car_fg; font-weight:800; white-space:nowrap; width:1%; }
.chips { display:flex; flex-wrap:wrap; gap:6px; align-items:center; }
.chip { display:inline-flex; gap:4px; align-items:baseline; background:$track; color:$ink; border-radius:99px;
        padding:3px 10px; font-size:.8rem; font-weight:600; white-space:nowrap; }
.chip b { font-weight:800; }
.chip.ok { background:$ok_bg; color:$ok_fg; }
.chip.wait { background:$no_bg; color:$no_fg; }
.chips.big { margin:2px 0 10px; }
.chips.big .chip { font-size:.9rem; padding:5px 14px; background:$card; border:1px solid $line; }
.maps { display:inline-flex; gap:6px; flex-wrap:wrap; margin-left:auto; }
.cars a.mapbtn { display:inline-block; padding:4px 12px; border-radius:99px; font-size:.8rem; font-weight:700;
                 text-decoration:none; background:$accent; color:#fff !important; white-space:nowrap; }
.cars a.mapbtn:hover { background:$accent_ink; }
.load { flex:1; min-width:80px; max-width:220px; height:6px; border-radius:99px; background:$track; overflow:hidden; }
.load i { display:block; height:100%; background:$done; border-radius:99px; }
.load.hi i { background:$left; }
.loadpct { font-size:.78rem; font-weight:700; color:$mute; }
.ship { background:$card; border:1px solid $line; border-left:5px solid $done; border-radius:16px; padding:14px 16px; margin-bottom:12px; }
.ship .sh-top { display:flex; justify-content:space-between; align-items:baseline; gap:10px; }
.ship .sh-t { font-weight:800; font-size:1.05rem; color:$ink; }
.ship .sh-t small { font-weight:500; color:$mute; font-size:.8rem; margin-left:6px; }
.ship .sh-p { font-weight:800; font-size:1.7rem; letter-spacing:-0.02em; color:$ink; }
.ship .bar { height:14px; margin:8px 0; }
.ship .sh-m { font-size:.85rem; color:$mute; margin-bottom:8px; }
.ship .sh-m b { color:$ink; }
.sync { display:inline-flex; width:20px; height:20px; border-radius:50%; align-items:center; justify-content:center;
        margin-left:8px; vertical-align:middle; font-size:.75rem; font-weight:800; box-sizing:border-box; }
.sync.ok { background:#1E9E5C; color:#fff; }
.sync.spin { border:3px solid $track; border-top-color:$left; animation:syncspin .9s linear infinite; }
@keyframes syncspin { to { transform:rotate(360deg); } }
.cars td.carmeta { vertical-align:middle; }
.cars tr.band td { background:$band; padding:0; height:12px; border:none; }
.cars tr.band.thin td { height:7px; }
.cars tr.hot td { background:$hot_bg; font-weight:700; }
.cars td.shop { font-weight:800; white-space:nowrap; }
.cars td.num { text-align:right; font-variant-numeric:tabular-nums; white-space:nowrap; }

/* ═════════ «M Cosmetic Dark v2» ═════════ */
.stApp { background: radial-gradient(1100px 560px at 75% -8%, rgba(139,61,255,.16), transparent 60%),
                     radial-gradient(800px 480px at 8% 0%, rgba(224,51,143,.11), transparent 55%), $paper; }
.block-container { max-width:1180px; padding-top:1rem; }

/* ── боковая панель ── */
section[data-testid="stSidebar"] { width:250px !important; min-width:250px !important; background:$side;
                                   border-right:1px solid $line; }
[data-testid="stSidebar"] > div:first-child { padding-top:.6rem; }
[data-testid="stSidebar"] .stButton > button {
  justify-content:flex-start; gap:12px; text-align:left; background:transparent; border:1px solid transparent;
  color:$mute; border-radius:12px; padding:.66rem .9rem; font-weight:600; }
[data-testid="stSidebar"] .stButton > button:hover { background:$head; color:$ink; border-color:$line; }
[data-testid="stSidebar"] .stButton > button[kind="primary"] {
  background:linear-gradient(135deg,$g1,$g2); color:#fff; border:none; box-shadow:0 8px 24px $glow; }
[data-testid="stSidebar"] .stButton > button[kind="primary"] p { color:#fff; }
.sidecard { margin-top:34px; padding:12px 14px; border-radius:14px; border:1px solid rgba(52,194,122,.35);
            background:linear-gradient(135deg, rgba(30,158,92,.22), rgba(30,158,92,.04)); font-size:.75rem; color:$mute; }
.sidecard b { display:flex; align-items:center; gap:7px; color:$ink; font-size:.84rem; margin-bottom:3px; }
.sidecard i { width:7px; height:7px; border-radius:50%; background:#34C27A; box-shadow:0 0 8px #34C27A; display:inline-block; }
.contact { display:flex; align-items:center; gap:12px; margin-top:14px; padding:12px; border-radius:14px;
           border:1px solid $line; background:$head; }
.contact .av { width:40px; height:40px; border-radius:50%; flex:none; background:linear-gradient(135deg,$g1,$g2); color:#fff;
               display:flex; align-items:center; justify-content:center; font-weight:800; font-size:.82rem; }
.contact .ct { min-width:0; display:flex; flex-direction:column; gap:3px; }
.contact .ct b { color:$ink; font-size:.88rem; }
.contact a, .addr { color:$mute !important; text-decoration:none !important; font-size:.76rem; }
.contact a { display:inline-flex; align-items:center; gap:6px; }
.contact a.tg { color:$accent !important; font-weight:700; }
.contact a:hover, .addr:hover { color:$accent !important; }
.addr { display:flex; align-items:flex-start; gap:8px; margin-top:8px; padding:10px 12px; border-radius:14px;
        border:1px solid $line; background:$head; line-height:1.3; }
.addr svg, .contact a svg { flex:none; }

/* мини-логотип в верхней панели */
.st-key-topbar .brand { justify-content:flex-start; }
.lt-win.lt-one { overflow:visible; }
@keyframes ltfloat { 0%,100% { transform:translateY(0); filter:drop-shadow(0 0 0 rgba(224,51,143,0)); }
                     50% { transform:translateY(-3px); filter:drop-shadow(0 5px 10px rgba(224,51,143,.45)); } }
.lt-one img { animation:ltfloat 3s ease-in-out infinite; }

/* невидимые кнопки: их «нажимает» скрипт при клике на «N онлайн» и на логотип */
.st-key-onl_open, .st-key-go_home { position:absolute !important; left:-9999px; top:0; width:1px; height:1px;
                                    overflow:hidden; opacity:0; pointer-events:none; }

/* ── верхняя панель: поиск, онлайн, дата, колокольчик, тема ── */
.st-key-topbar { background:$card; border:1px solid $line; border-radius:16px; padding:8px 14px; margin-bottom:14px;
                 position:relative; z-index:60; }
.st-key-topbar [data-testid="stHorizontalBlock"] { align-items:center; gap:.75rem; }
.st-key-gsearch { position:relative; }
.st-key-gsearch input { padding-left:40px !important; padding-right:92px !important; height:40px;
  background-image:url("$u_search_s") !important; background-repeat:no-repeat !important;
  background-position:13px center !important; background-size:16px !important; }
.st-key-gsearch::after { content:"Ctrl + K"; position:absolute; right:12px; top:50%; transform:translateY(-50%);
  font-size:.7rem; font-weight:600; color:$mute; border:1px solid $line; border-radius:7px; padding:2px 7px;
  pointer-events:none; background:$head; }
.tb-right { display:flex; align-items:center; justify-content:flex-end; gap:14px; font-size:.8rem; color:$mute; white-space:nowrap; }
.tb-right .live { display:flex; align-items:center; gap:7px; }
.tb-right .dot { width:8px; height:8px; border-radius:50%; background:#34C27A; animation:livepulse 1.4s ease-in-out infinite; }
.tb-right .sep { width:1px; height:20px; background:$line; }
.tb-right .tb-date { display:flex; align-items:center; gap:8px; }
.tb-right .onl { cursor:pointer; padding:5px 10px; margin:-5px -10px; border-radius:10px; transition:background .15s, color .15s; }
.tb-right .onl:hover { background:$head; color:$ink; }
.tb-right .bell { position:relative; display:flex; cursor:pointer; }
.tb-right .bell::after { content:""; position:absolute; left:-24px; right:-12px; top:100%; height:16px; }
.tb-right .bell .bdot { position:absolute; top:-2px; right:-2px; width:8px; height:8px; border-radius:50%;
                        background:#FF3B7F; box-shadow:0 0 8px #FF3B7F; }
.bell-pop { display:none; position:absolute; right:-12px; top:calc(100% + 12px); width:370px; max-width:88vw;
            z-index:9999; padding:14px 16px; border-radius:16px; background:$card; border:1px solid $line;
            box-shadow:0 18px 50px rgba(0,0,0,.45); white-space:normal; text-align:left; color:$ink; cursor:default; }
.bell:hover .bell-pop, .bell-pop:hover { display:block; }
.bell-pop .bp-h { font-weight:800; font-size:.95rem; color:$ink; margin-bottom:8px; }
.bell-pop .bp-box { padding:8px 10px; border-radius:10px; background:$head; border:1px solid $line;
                    font-size:.74rem; color:$mute; line-height:1.55; margin-bottom:8px; }
.bell-pop .bp-box b { color:$ink; }
.bell-pop .bp-list { max-height:300px; overflow-y:auto; }
.bell-pop .ev { display:grid; grid-template-columns:78px 8px 1fr; gap:8px; align-items:baseline; padding:5px 0;
                border-top:1px solid $line; font-size:.76rem; color:$ink; }
.bell-pop .ev .t { color:$mute; font-size:.7rem; font-variant-numeric:tabular-nums; }
.bell-pop .ev i { width:8px; height:8px; border-radius:50%; display:inline-block; }
.bell-pop .ev small { color:$mute; font-size:.7rem; }
.bell-pop .empty { font-size:.76rem; color:$mute; padding:8px 0; }
.bell-pop .bp-f { margin-top:8px; font-size:.66rem; color:$mute; line-height:1.4; }

.st-key-theme_btn button { width:40px; height:40px; min-height:40px; padding:0; border-radius:12px; background:$head; gap:0; }
.st-key-theme_btn button p { display:none; }

/* ── окно «Посетители» (открывается по клику на «N онлайн») ── */
div[data-testid="stDialog"] div[role="dialog"] { background:$card !important; border:1px solid $line;
                                                 border-radius:18px; color:$ink; }
div[data-testid="stDialog"] h2, div[data-testid="stDialog"] button[aria-label="Close"] { color:$ink !important; }
.vk-wrap { display:grid; grid-template-columns:repeat(3, 1fr); gap:10px; margin:6px 0 14px; }
.vk { display:flex; align-items:center; justify-content:space-between; gap:10px; padding:12px 14px;
      border-radius:14px; background:$head; border:1px solid $line; }
.vk .vkl { font-size:.6rem; font-weight:700; letter-spacing:.06em; color:$mute; text-transform:uppercase; }
.vk .vkv { font-size:1.2rem; font-weight:800; color:$ink; margin-top:2px; white-space:nowrap; }
.vk .vki { width:34px; height:34px; border-radius:50%; flex:none; display:flex; align-items:center;
           justify-content:center; background:$track; color:$accent; }
.vtbl { border:1px solid $line; border-radius:14px; overflow-x:auto; background:$card; }
.vrow { display:grid; grid-template-columns:minmax(170px,1.5fr) auto minmax(150px,1fr) minmax(230px,1.7fr); gap:12px;
        align-items:center; padding:9px 14px; border-top:1px solid $line; min-width:660px; font-size:.8rem; color:$ink; }
.vrow.vh { border-top:none; background:$head; font-weight:800; font-size:.8rem; }
.vuser { display:flex; align-items:center; gap:10px; min-width:0; }
.vav { width:34px; height:34px; border-radius:50%; flex:none; overflow:hidden; display:flex; align-items:center;
       justify-content:center; color:#fff; font-weight:800; font-size:.72rem; }
.vname { color:$accent; font-weight:700; overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }
.vme { color:$mute; font-weight:500; margin-left:4px; }
.vpill { display:inline-flex; align-items:center; gap:6px; padding:3px 10px; border-radius:99px; font-size:.7rem;
         font-weight:700; white-space:nowrap; }
.vpill.s { background:$no_bg; color:$no_fg; }
.vpill.t { background:$track; color:$ink; }
.vpill.on { background:$ok_bg; color:$ok_fg; }
.vdet { display:flex; align-items:center; gap:10px; flex-wrap:wrap; }
.vd { display:inline-flex; color:$mute; cursor:help; }
.vtx { color:$mute; font-size:.72rem; white-space:nowrap; }
@media (max-width: 700px) { .vk-wrap { grid-template-columns:1fr; } }

/* ── заголовок и главные кнопки ── */
.top .crumb { display:flex; align-items:center; gap:6px; }
.top h1 { display:flex; align-items:center; gap:12px; font-size:1.55rem; }
.top h1::before { content:""; width:36px; height:36px; border-radius:11px; flex:none;
  background:url("$u_box_w") center/20px no-repeat, linear-gradient(135deg,$g1,$g2); box-shadow:0 6px 18px $glow; }
.stApp .stButton > button[kind="primary"] { background:linear-gradient(135deg,$g1,$g2); color:#fff; border:none;
                                            box-shadow:0 8px 22px $glow; }
.st-key-btn_refresh button, .st-key-btn_theme2 button { height:40px; border-radius:12px; gap:8px; }

/* ── вкладки-разделы ── */
.stApp [data-testid="stButtonGroup"] { background:$card; border:1px solid $line; border-radius:14px; padding:4px; gap:2px; }
.stApp [data-testid="stButtonGroup"] button { border:none; background:transparent; color:$mute; border-radius:10px; gap:7px; }
.stApp [data-testid="stButtonGroup"] button[aria-checked="true"] {
  background:linear-gradient(135deg,$g1,$g2); color:#fff; box-shadow:0 4px 16px $glow; }
.stApp [data-testid="stButtonGroup"] button[aria-checked="true"] p { color:#fff; }

/* ── карточки KPI: иконка-плитка, подпись, процент, мини-график (компактные) ── */
.kpi { display:grid; grid-template-columns:auto 1fr auto; column-gap:12px; align-items:center; min-height:76px;
       padding:10px 14px; border-radius:14px; overflow:hidden;
       border:1px solid rgba(224,51,143,.45); border-left:1px solid rgba(224,51,143,.45);
       background:linear-gradient(135deg, rgba(224,51,143,.30), $card 72%);
       box-shadow:0 8px 26px rgba(224,51,143,.16), inset 0 1px 0 rgba(255,255,255,.05); }
.kpi .ico { width:38px; height:38px; border-radius:11px; display:flex; align-items:center; justify-content:center;
            background:linear-gradient(135deg,$g1,$g2); box-shadow:0 4px 12px rgba(224,51,143,.4); }
.kpi .mid { min-width:0; }
.kpi .v { font-size:1.45rem; font-weight:800; letter-spacing:-.02em; line-height:1.1; }
.kpi .l { font-size:.78rem; font-weight:600; color:$mute; margin-top:2px; }
.kpi .s { font-size:.7rem; color:$mute; margin-top:1px; }
.kpi .sts { grid-column:1 / -1; display:flex; flex-wrap:wrap; gap:3px 10px; margin-top:8px; padding-top:6px;
            border-top:1px dashed $line; }
.kpi .sts .stt { width:100%; font-size:.56rem; font-weight:700; text-transform:uppercase; letter-spacing:.07em; color:$mute; }
.kpi .sts .stp { display:inline-flex; align-items:center; gap:5px; font-size:.62rem; color:$mute; white-space:nowrap; }
.kpi .sts .stp i { width:6px; height:6px; border-radius:50%; display:inline-block; }
.kpi .sts .stp b { color:$ink; font-weight:700; }
.st-key-tap_all .kpi, .st-key-tap_hub .kpi, .st-key-tap_pcs .kpi { min-height:112px; }
.kpi.nov .l { font-size:.88rem; font-weight:800; color:$ink; margin:0; }
.kpi.nov .sts { display:grid; grid-template-columns:1fr 1fr; gap:3px 14px; margin-top:8px; }
.kpi.nov .sts .stp { justify-content:space-between; font-size:.66rem; }
.kpi.nov .sts .stp span.nm { display:inline-flex; align-items:center; gap:5px; }
.kpi .trend { display:flex; flex-direction:column; align-items:flex-end; gap:2px; }
.kpi .delta { font-size:.7rem; font-weight:800; }
.kpi .delta.up { color:#34C27A; } .kpi .delta.dn { color:#F2A23D; }
.kpi.left { border-color:rgba(242,162,61,.45); border-left:1px solid rgba(242,162,61,.45);
            background:linear-gradient(135deg, rgba(242,162,61,.26), $card 72%); box-shadow:0 8px 26px rgba(242,162,61,.13); }
.kpi.left .ico { background:linear-gradient(135deg,#F6A93B,#E8702A); box-shadow:0 4px 12px rgba(242,162,61,.4); }
.kpi.blue { border-color:rgba(59,130,246,.5); border-left:1px solid rgba(59,130,246,.5);
            background:linear-gradient(135deg, rgba(59,130,246,.26), $card 72%); box-shadow:0 8px 26px rgba(59,130,246,.15); }
.kpi.blue .ico { background:linear-gradient(135deg,#3B9CFF,#2563EB); box-shadow:0 4px 12px rgba(59,130,246,.4); }
.kpi.ok { border-color:rgba(52,194,122,.45); border-left:1px solid rgba(52,194,122,.45);
          background:linear-gradient(135deg, rgba(52,194,122,.22), $card 72%); box-shadow:0 8px 26px rgba(52,194,122,.12); }
.kpi.ok .ico { background:linear-gradient(135deg,#34C27A,#1E9E5C); box-shadow:0 4px 12px rgba(52,194,122,.4); }
[class*="st-key-tap_"]:hover .kpi, [class*="st-key-tap_"]:hover .day { transform:translateY(-1px); filter:brightness(1.08); }
.kpi, .day { transition:transform .15s, filter .15s, border-color .15s; }

/* ── карточка «Хаб ББ»: мелкая таблица Дата / Принято / Не принято ── */
.hubt { grid-column:1 / -1; display:flex; flex-direction:column; gap:2px; margin-top:8px; padding-top:6px;
        border-top:1px dashed $line; }
.hubt .hr { display:grid; grid-template-columns:1.2fr 1fr 1fr; gap:8px; align-items:center; font-size:.66rem; color:$ink; }
.hubt .hr.hh { font-size:.56rem; font-weight:700; text-transform:uppercase; letter-spacing:.05em; color:$mute; }
.hubt .hr.hh span:nth-child(n+2), .hubt .hr b { text-align:right; }
.hubt .hr b { font-variant-numeric:tabular-nums; font-weight:800; }
.hubt .hr b.g { color:#34C27A; }
.hubt .hr b.r { color:#F2A23D; }
.hubt .none { font-size:.7rem; color:$mute; }

/* ── карточка отгрузки (компактная) ── */
.ship { border-radius:16px; border:1px solid rgba(139,61,255,.45); border-left:1px solid rgba(139,61,255,.45);
        background:linear-gradient(135deg, $head, $card); box-shadow:0 0 0 1px rgba(139,61,255,.12), 0 10px 30px $glow;
        padding:10px 14px; margin-bottom:8px; }
.ship .sh-t { display:flex; align-items:center; gap:8px; font-size:.95rem; }
.ship .sh-p { font-size:1.4rem; }
.ship .sh-m { font-size:.78rem; margin-bottom:0; margin-top:6px; }
.ship .bar { height:12px; margin:0; }
.pdot { width:8px; height:8px; border-radius:50%; background:$g1; box-shadow:0 0 8px $g1; display:inline-block; }
.bar { height:10px; background:$track; }
.bar > i { background-color:$g1;
  background-image:linear-gradient(135deg, rgba(255,255,255,.24) 25%, transparent 25%, transparent 50%,
                                   rgba(255,255,255,.24) 50%, rgba(255,255,255,.24) 75%, transparent 75%),
                   linear-gradient(90deg,$g1,$g2,#3B82F6);
  background-size:32px 32px, 100% 100%; animation:stripes 1s linear infinite; box-shadow:0 0 12px $glow; }
@keyframes stripes { from { background-position:0 0, 0 0; } to { background-position:32px 0, 0 0; } }

/* ── строки по дням: цветная полоска, плашка с числом, стрелка (компактные) ── */
.day { background:$row; border:1px solid $line; border-left:4px solid $g1; border-radius:12px; padding:8px 12px;
       grid-template-columns:70px 1fr auto 22px; gap:10px; }
.day .d { font-size:.92rem; }
.day .d small { font-size:.7rem; }
.day .meta { font-size:.72rem; margin-top:4px; }
.day .bar { height:9px; }
.day.urg { border-left-color:#E5384F; }
.day.amb { border-left-color:$left; }
.day.full { border-left-color:$done; }
.day .pct { background:$head; border:1px solid $line; border-radius:10px; padding:3px 10px; font-size:.95rem;
            min-width:50px; text-align:center; }
.day .pct small { font-size:.62rem; }
.day .chev { display:flex; align-items:center; justify-content:center; color:$mute; }
@media (max-width: 640px) {
  .day { grid-template-columns:1fr auto 20px; }
  .day .mid { grid-column:1 / -1; order:4; }
}
.badge.no { background:$no_bg; color:$no_fg; border:1px solid rgba(242,162,61,.35); }
.badge.ok { border:1px solid rgba(52,194,122,.35); }

/* ── таблицы и кнопки скачивания ── */
.tbl-wrap, .cars-wrap { border-radius:18px; box-shadow:0 10px 30px rgba(0,0,0,.18); }
.tbl th { background:$head; }
.stApp [data-testid="stExpander"], .hit, .note { border-radius:16px; }

/* компактная кнопка «Excel» */
.stApp .stDownloadButton > button { position:relative; background:$row; border-radius:10px;
  padding:.3rem 2.3rem .3rem .8rem; min-height:34px; font-size:.8rem; gap:0; }
.stApp .stDownloadButton > button p { font-size:.8rem; }
.stApp .stDownloadButton > button::before { content:""; width:16px; height:16px; flex:none; margin-right:8px;
  background:url("$u_excel") center/contain no-repeat; }
.stApp .stDownloadButton > button::after { content:""; position:absolute; right:10px; top:50%; margin-top:-7px;
  width:14px; height:14px; background:url("$u_download_s") center/contain no-repeat; }
.stApp .stDownloadButton > button:hover { border-color:$g1; color:$ink; box-shadow:0 0 0 1px $g1; }

/* ── убрать системные надписи загрузки и «серую» подсветку ── */
[data-testid="stSpinner"], .stSpinner, [data-testid="stStatusWidget"] { display:none !important; }
[data-stale="true"] { opacity:1 !important; }

/* ── кнопка сворачивания / раскрытия меню «<<» / «>>»: всегда видна ── */
header[data-testid="stHeader"] { display:block !important; background:transparent !important; height:auto !important;
  min-height:0 !important; overflow:visible !important; pointer-events:none !important; }
header[data-testid="stHeader"] * { pointer-events:none; }
[data-testid="stToolbar"] { background:transparent !important; }
[data-testid="stExpandSidebarButton"], [data-testid="stSidebarCollapsedControl"] {
  display:flex !important; visibility:visible !important; opacity:1 !important; pointer-events:auto !important;
  position:fixed !important; top:14px !important; left:12px !important; z-index:1000002 !important; }
[data-testid="stExpandSidebarButton"] *, [data-testid="stSidebarCollapsedControl"] * {
  pointer-events:auto !important; visibility:visible !important; opacity:1 !important; color:$ink !important; }
[data-testid="stExpandSidebarButton"], [data-testid="stSidebarCollapsedControl"] button {
  background:$head !important; border:1px solid $line !important; border-radius:10px !important; }
[data-testid="stSidebarHeader"] { overflow:visible !important; position:relative; z-index:20; }
[data-testid="stSidebarCollapseButton"], [data-testid="stSidebarCollapseButton"] button {
  display:flex !important; visibility:visible !important; opacity:1 !important; }
[data-testid="stSidebarCollapseButton"] button { background:$head !important; border:1px solid $line !important;
  border-radius:10px !important; }
[data-testid="stSidebarCollapseButton"] * { color:$ink !important; opacity:1 !important; }
/* когда меню свёрнуто, сдвигаем верхнюю панель, чтобы кнопка не перекрывала логотип */
.stApp:has(section[data-testid="stSidebar"][aria-expanded="false"]) .st-key-topbar { margin-left:46px; }

/* ── молния по рамке поиска при обновлении ── */
@property --bolt { syntax:"<angle>"; inherits:false; initial-value:0deg; }
.st-key-gsearch::before { content:""; position:absolute; inset:0; border-radius:12px; padding:2px;
  pointer-events:none; z-index:6; opacity:0; transition:opacity .5s ease;
  background:conic-gradient(from var(--bolt), transparent 0deg, transparent 240deg,
             rgba(176,60,214,.28) 285deg, rgba(200,90,235,.6) 335deg, rgba(236,180,252,.85) 355deg, transparent 360deg);
  -webkit-mask:linear-gradient(#000 0 0) content-box, linear-gradient(#000 0 0);
  -webkit-mask-composite:xor; mask-composite:exclude;
  filter:drop-shadow(0 0 3px rgba(176,60,214,.5)); }
.stApp[data-teststate="running"] .st-key-gsearch::before,
.stApp[data-test-script-state="running"] .st-key-gsearch::before {
  opacity:1; animation:boltspin 2.8s linear infinite; }
@keyframes boltspin { to { --bolt:360deg; } }
$truck_css
</style>
""")
_vals = dict(P)
_vals.update(u_search_s=icon_uri("search", P["mute"]), u_box_w=icon_uri("box", "#FFFFFF"),
             u_download_s=icon_uri("download", P["mute"]),
             u_excel="data:image/svg+xml;utf8," + quote(_EXCEL_SVG, safe=""),
             truck_css=TRUCK_CSS)
st.markdown(CSS.safe_substitute(_vals), unsafe_allow_html=True)

# иконки на кнопках: меню слева, вкладки разделов, тема, «Обновить»
_ir = []
for _i, _nm in enumerate(NAV_ICON_NAMES):
    _ir.append(_icon_rule(f".st-key-nav_{_i} button", _nm, 18))
    _ir.append(_icon_rule(f'.st-key-page [data-testid="stButtonGroup"] button:nth-of-type({_i + 1})', _nm, 15))
_theme_ic = "sun" if st.session_state.dark else "moon"
_ir.append(_icon_rule(".st-key-theme_btn button", _theme_ic, 18))
_ir.append(_icon_rule(".st-key-btn_theme2 button", _theme_ic, 17))
_ir.append(_icon_rule(".st-key-btn_refresh button", "refresh", 17))
st.markdown("<style>" + "".join(_ir) + "</style>", unsafe_allow_html=True)


# ───────────────────────── Данные ─────────────────────────
def _client():
    creds = Credentials.from_service_account_info(
        dict(st.secrets["gcp_service_account"]),
        scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"],
    )
    return gspread.authorize(creds)


def _client_rw():
    """Клиент с правом записи — только для сохранения ручных рейсов."""
    creds = Credentials.from_service_account_info(
        dict(st.secrets["gcp_service_account"]),
        scopes=["https://www.googleapis.com/auth/spreadsheets"],
    )
    return gspread.authorize(creds)


@st.cache_data(ttl=CACHE_SECONDS, show_spinner=False)
def load() -> pd.DataFrame:
    sheet = _client().open_by_key(SHEET_ID)
    ws = next(w for w in sheet.worksheets() if str(w.id) == GID)
    rows = ws.get_all_values()
    head = [(h.strip() or f"col{i}") for i, h in enumerate(rows[0])]
    raw = pd.DataFrame(rows[1:], columns=head, dtype=str).fillna("")

    def find(*keys):
        for c in raw.columns:
            if all(k in c.lower() for k in keys):
                return c
        return None

    c_order, c_shop, c_addr = find("№"), find("название"), find("адрес")
    c_city, c_qty = find("город"), find("кол-во")
    c_date, c_wms = find("дата"), find("wms")

    # колонка «в сборке» без заголовка: ищем по содержимому, запасной вариант — 9-я колонка (I)
    c_pick = next((c for c in raw.columns
                   if raw[c].str.lower().str.contains("в сборке").any()), raw.columns[8])

    # оставляем только строки с заказами (без разделителей вида «13/10/2026 (заказы на вт)»)
    raw = raw[raw[c_order].str.match(r"^\s*\w{1,3}-\d{4,}")].copy()

    df = pd.DataFrame({
        "Заказ": raw[c_order].str.strip(),
        "Магазин": raw[c_shop].str.strip(),
        "Адрес": raw[c_addr].str.strip() if c_addr else "",
        "Город": raw[c_city].str.strip(),
        "Штук": pd.to_numeric(raw[c_qty].str.replace(r"\s", "", regex=True), errors="coerce").fillna(0).astype(int),
        "Дата": pd.to_datetime(raw[c_date], dayfirst=True, errors="coerce").dt.date,
        "Сборка": raw[c_pick].str.strip(),
        "Статус WMS": raw[c_wms].str.strip(),
    })
    df["Собран"] = df["Сборка"].str.lower().str.contains("в сборке")
    return df.dropna(subset=["Дата"]).reset_index(drop=True)


@st.cache_data(ttl=CACHE_SECONDS, show_spinner=False)
def load_new_orders() -> pd.DataFrame:
    """Вкладка «Новые заказы» (не маршрутизированные):
    Номер_заказа / Статус_заказа / Дата_доставки / Товаров_зарезервировано / Точка_доставки / Магазин."""
    sheet = _client().open_by_key(SHEET_ID)
    ws = next((w for w in sheet.worksheets() if str(w.id) == NEW_GID), None) \
        or next((w for w in sheet.worksheets() if "новые" in w.title.lower()), None)
    if ws is None:
        raise ValueError("Не найдена вкладка «Новые заказы».")
    rows = ws.get_all_values()
    head = [(" ".join(h.split()) or f"col{i}") for i, h in enumerate(rows[0])]
    raw = pd.DataFrame(rows[1:], columns=head, dtype=str).fillna("")

    def find(*keys):
        return next((c for c in raw.columns if all(k in c.lower() for k in keys)), None)

    c_order = find("номер") or find("заказ")
    c_stat = find("статус")
    c_date = find("дата")
    c_qty = find("зарезерв") or find("товар") or find("кол-во")
    c_shop = find("магазин")
    if not all([c_order, c_date, c_qty]):
        raise ValueError("Во вкладке «Новые заказы» не найдены колонки: номер заказа, дата доставки, товаров зарезервировано.")

    raw = raw[raw[c_order].str.match(r"^\s*\w{1,3}-\d{4,}")].copy()
    d = pd.DataFrame({
        "Заказ": raw[c_order].str.strip(),
        "Статус": raw[c_stat].str.strip() if c_stat else "",
        "Дата": pd.to_datetime(raw[c_date].str.strip(), dayfirst=True, errors="coerce").dt.date,
        "Штук": pd.to_numeric(raw[c_qty].str.replace(r"\s", "", regex=True), errors="coerce").fillna(0).astype(int),
        "Магазин": raw[c_shop].str.strip() if c_shop else "",
    })
    return d.drop_duplicates("Заказ").reset_index(drop=True)


@st.cache_data(ttl=CACHE_SECONDS, show_spinner=False)
def load_status_orders() -> set:
    """Вкладка «Статус по отгрузкам» основной таблицы: множество номеров заказов, которые уже туда попали
    (ищем номера во всех ячейках вкладки, поэтому расположение колонок не важно)."""
    sheet = _client().open_by_key(SHEET_ID)
    ws = next((w for w in sheet.worksheets()
               if "статус" in w.title.lower() and "отгруз" in w.title.lower()), None)
    if ws is None:
        raise ValueError("Не найдена вкладка «Статус по отгрузкам».")
    pat = re.compile(r"^\s*\w{1,3}-\d{4,}\s*$")
    out = set()
    for row in ws.get_all_values():
        for cell in row:
            if pat.match(cell):
                out.add(cell.strip().upper())
    return out


@st.cache_data(ttl=CACHE_SECONDS, show_spinner=False)
def load_ship_tabs() -> dict:
    """Две вкладки таблицы отгрузок: «Маршруты New» (текущий день) и «Архив» (прошлые дни)."""
    sheet = _client().open_by_key(ARCHIVE_SHEET_ID)
    sheets = sheet.worksheets()

    def read(ws):
        rows = ws.get_all_values()
        head = [(" ".join(h.split()) or f"col{i}") for i, h in enumerate(rows[0])]
        return pd.DataFrame(rows[1:], columns=head, dtype=str).fillna("")

    arch = next((w for w in sheets if str(w.id) == ARCHIVE_GID), None) or sheet.worksheet(ARCHIVE_TAB)
    cur = next((w for w in sheets if "маршрут" in w.title.lower() and "new" in w.title.lower()), None)
    return {"archive": read(arch), "new": read(cur) if cur else None}


@st.cache_data(ttl=CACHE_SECONDS, show_spinner=False)
def load_boxes() -> pd.DataFrame:
    """Ответы формы: Отметка времени / Магазин / Кол-во полученных коробок."""
    sheet = _client().open_by_key(BOXES_SHEET_ID)
    ws = next((w for w in sheet.worksheets() if str(w.id) == BOXES_GID), None) or sheet.worksheet(BOXES_TAB)
    rows = ws.get_all_values()
    head = [(" ".join(h.split()) or f"col{i}") for i, h in enumerate(rows[0])]
    raw = pd.DataFrame(rows[1:], columns=head, dtype=str).fillna("")

    def find(*keys):
        return next((c for c in raw.columns if all(k in c.lower() for k in keys)), None)

    c_t, c_s = find("отметка"), find("магазин")
    c_q = find("кол-во") or find("коробок")
    if not all([c_t, c_s, c_q]):
        raise ValueError("Не найдены колонки: отметка времени, магазин, кол-во коробок.")
    ts = raw[c_t].str.strip()
    t = pd.to_datetime(ts, format="%d.%m.%Y %H:%M:%S", errors="coerce")
    miss = t.isna() & (ts != "")
    if miss.any():
        t[miss] = pd.to_datetime(ts[miss], dayfirst=True, errors="coerce")
    d = pd.DataFrame({
        "Время": t,
        "Магазин": raw[c_s].str.strip(),
        "Коробок": pd.to_numeric(raw[c_q].str.replace(r"\s", "", regex=True), errors="coerce").fillna(0).astype(int),
    })
    d = d[d["Время"].notna() & (d["Магазин"] != "")].sort_values("Время").reset_index(drop=True)
    d["Месяц"] = d["Время"].dt.strftime("%Y-%m")
    return d


@st.cache_data(ttl=CACHE_SECONDS, show_spinner=False)
def load_cat_prio() -> dict:
    """Лист1: Номер заказа 1С / Категория / Приоритет -> {номер заказа: (категория, приоритет)}."""
    sheet = _client().open_by_key(CAT_SHEET_ID)
    ws = next((w for w in sheet.worksheets() if str(w.id) == CAT_GID), None) or sheet.worksheet("Лист1")
    rows = ws.get_all_values()
    hi = next((i for i, r in enumerate(rows[:15]) if any("номер заказа" in c.lower() for c in r)), 0)
    head = [" ".join(h.split()).lower() for h in rows[hi]]

    def col(key, default):
        return next((i for i, h in enumerate(head) if key in h), default)

    i_o, i_c, i_p = col("номер заказа", 0), col("категори", 1), col("приоритет", 2)
    out = {}
    for r in rows[hi + 1:]:
        if len(r) <= i_o or not r[i_o].strip():
            continue
        cat = r[i_c].strip() if len(r) > i_c else ""
        pr = r[i_p].strip() if len(r) > i_p else ""
        out[r[i_o].strip().upper()] = (cat, pr)
    return out


# ───────────────────────── Metabase: остатки ─────────────────────────
def _nk(x) -> str:
    """Ключ для сравнения названий колонок: без регистра, пробелов и знаков."""
    return re.sub(r"[^0-9a-zа-я]", "", str(x).lower())


def _mb_cfg():
    """url (без /api...), заголовок с ключом и секция [metabase] из Secrets."""
    try:
        cfg = st.secrets["metabase"]
        url = re.sub(r"/api(/.*)?$", "", str(cfg["url"]).strip().rstrip("/"))
        hdr = {"x-api-key": str(cfg["api_key"]).strip()}
    except Exception:
        raise ValueError("В Secrets нет блока [metabase] с url и api_key.")
    return cfg, url, hdr


@st.cache_data(ttl=3600, show_spinner=False)
def mb_meta(tid: int):
    """db_id и поля таблицы: {нормализованное имя: (id, тип)}."""
    _, url, hdr = _mb_cfg()
    r = requests.get(f"{url}/api/table/{tid}/query_metadata", headers=hdr, timeout=60)
    r.raise_for_status()
    m = r.json()
    flds = {}
    for f in m.get("fields", []):
        flds[_nk(f.get("name"))] = (f["id"], str(f.get("base_type", "")))
        flds.setdefault(_nk(f.get("display_name")), (f["id"], str(f.get("base_type", ""))))
    return m["db_id"], flds


def mb_run(db: int, query: dict) -> pd.DataFrame:
    """Выполняет MBQL-запрос в Metabase и возвращает таблицу."""
    _, url, hdr = _mb_cfg()
    q = {"database": db, "type": "query", "query": query}
    r = requests.post(f"{url}/api/dataset/json", headers=hdr,
                      data={"query": json.dumps(q), "format_rows": "false"}, timeout=300)
    r.raise_for_status()
    data = r.json()
    if isinstance(data, dict):
        raise ValueError(str(data.get("error") or data)[:300])
    return pd.DataFrame(data)


def mb_table(key: str, default_id: int) -> pd.DataFrame:
    """Читает таблицу Metabase целиком (используется для Magnit Stock)."""
    cfg, _, _ = _mb_cfg()
    tid = int(cfg.get(key, default_id) or 0)
    if not tid:
        raise ValueError(f"Не указан ID таблицы: добавьте {key} в блок [metabase] в Secrets.")
    db, _f = mb_meta(tid)
    raw = mb_run(db, {"source-table": tid})
    if raw.empty:
        raise ValueError("Таблица вернула 0 строк.")
    return raw


def _to_date(ser: pd.Series) -> pd.Series:
    s = ser.astype(str).str.strip()
    out = pd.to_datetime(s.str[:10], format="%Y-%m-%d", errors="coerce")
    miss = out.isna()
    if miss.any():
        out[miss] = pd.to_datetime(s[miss], dayfirst=True, errors="coerce")
    return out.dt.date


def _num(ser: pd.Series) -> pd.Series:
    s = ser.astype(str).str.replace("\u00a0", "", regex=False).str.replace(" ", "", regex=False)
    return pd.to_numeric(s.str.replace(",", ".", regex=False), errors="coerce").fillna(0)


def _sku(ser: pd.Series) -> pd.Series:
    return ser.astype(str).str.strip().str.replace(r"\.0$", "", regex=True)


def _stock_tid() -> int:
    cfg, _, _ = _mb_cfg()
    return int(cfg.get("stock_table", MB_STOCK_T) or MB_STOCK_T)


def _norm_stock(raw: pd.DataFrame) -> pd.DataFrame:
    """Приводит выгрузку остатков к колонкам: date / sku / name / cell / inv / zone / loc / avail.
    Индекс строк сохраняется (по нему можно отобрать те же строки в исходной выгрузке)."""
    cols = {_nk(c): c for c in raw.columns}
    need = {"snapshotdate": "Snapshot Date", "skuwmsid": "Sku Wms ID", "name": "Name",
            "cellbarcode": "Cell Barcode", "amountinlocation": "Amount In Location"}
    miss = [t for k, t in need.items() if k not in cols]
    if miss:
        raise ValueError("В таблице нет колонок: " + ", ".join(miss) + ". Есть: " + ", ".join(map(str, raw.columns)))
    c_inv, c_zone, c_av = cols.get("originalinvoicenumber"), cols.get("zoneprefix"), cols.get("amountavailable")
    cell = raw[cols["cellbarcode"]].astype(str).str.strip()
    d = pd.DataFrame({
        "date": _to_date(raw[cols["snapshotdate"]]),
        "sku": _sku(raw[cols["skuwmsid"]]),
        "name": raw[cols["name"]].astype(str).str.strip(),
        "cell": cell,
        "inv": raw[c_inv].astype(str).str.strip() if c_inv else "",
        "zone": raw[c_zone].astype(str).str.strip() if c_zone else cell.str.split("-").str[0],
        "loc": _num(raw[cols["amountinlocation"]]),
        "avail": _num(raw[c_av]) if c_av else _num(raw[cols["amountinlocation"]]),
    })
    return d[d["date"].notna() & (d["sku"] != "")]


@st.cache_data(ttl=3600, show_spinner=False)
def load_stock_catalog() -> pd.DataFrame:
    """Каталог товаров (sku, name) за последние MB_CAT_DAYS дней — лёгкий запрос, группировка на стороне Metabase."""
    tid = _stock_tid()
    db, f = mb_meta(tid)
    since = (datetime.now(TZ).date() - timedelta(days=MB_CAT_DAYS)).isoformat()
    query = {"source-table": tid,
             "aggregation": [["count"]],
             "breakout": [["field", f["skuwmsid"][0], None], ["field", f["name"][0], None]],
             "filter": [">=", ["field", f["snapshotdate"][0], None], since]}
    raw = mb_run(db, query)
    if raw.empty:
        raise ValueError("За последние дни в таблице остатков нет данных.")
    cols = {_nk(c): c for c in raw.columns}
    d = pd.DataFrame({"sku": _sku(raw[cols["skuwmsid"]]), "name": raw[cols["name"]].astype(str).str.strip()})
    return d[d["sku"] != ""].drop_duplicates("sku").reset_index(drop=True)


@st.cache_data(ttl=1800, show_spinner=False)
def load_stock_rows(skus: tuple, d_from: date, d_to: date):
    """Снимки остатков только по выбранным товарам и периоду — фильтрует Metabase.
    Возвращает (нормализованная таблица, исходные строки Metabase со всеми колонками)."""
    tid = _stock_tid()
    db, f = mb_meta(tid)
    sku_id, sku_type = f["skuwmsid"]
    numeric = ("Integer" in sku_type or "Number" in sku_type) and all(s.isdigit() for s in skus)
    vals = [int(s) for s in skus] if numeric else list(skus)
    flt = ["and",
           [">=", ["field", f["snapshotdate"][0], None], d_from.isoformat()],
           ["<=", ["field", f["snapshotdate"][0], None], d_to.isoformat()],
           ["=", ["field", sku_id, None], *vals]]
    raw = mb_run(db, {"source-table": tid, "filter": flt})
    if raw.empty:
        return pd.DataFrame(columns=["date", "sku", "name", "cell", "inv", "zone", "loc", "avail"]), raw
    d = _norm_stock(raw)
    return d.reset_index(drop=True), raw.loc[d.index].reset_index(drop=True)


@st.cache_data(ttl=600, show_spinner=False)
def load_magnit() -> pd.DataFrame:
    return mb_table("magnit_table", MB_MAGNIT_T)


# ───────────────────────── Остаток позиции по Magnit Stock ─────────────────────────
LIST_HEADERS = {"внешний id позиции", "товар", "sku wms id", "name", "sku", "название"}


def _mpick(cols: dict, *keys):
    """Первая найденная колонка из списка вариантов названий."""
    for k in keys:
        if k in cols:
            return cols[k]
    return None


def _key(x) -> str:
    """Ключ для сравнения значений: без регистра, лишних пробелов и «.0» на конце."""
    s = str(x).strip().lower().replace("ё", "е")
    s = re.sub(r"\.0$", "", s)
    return re.sub(r"\s+", " ", s)


def zkey(z) -> str:
    """Ключ зоны для сравнения: без пробелов, в верхнем регистре («1m-buf» = «1M-BUF»)."""
    return re.sub(r"\s+", "", str(z)).upper()


RES_KEYS = {zkey(z) for z in RESERVE_ZONES}
NORES_KEYS = {zkey(z) for z in NO_RESERVE_ZONES}
ZONE_ORDER = {zkey(z): i for i, z in enumerate(RESERVE_ZONES + NO_RESERVE_ZONES)}
KIND_TEXT = {"yes": "можно зарезервировать", "no": "Заказ не зарезервируется", "?": "зона не указана в правилах"}


def magnit_norm(raw: pd.DataFrame) -> pd.DataFrame:
    """Magnit Stock -> sku / ext / wms / name / zone / cell / inv / loc / res / avail / kind (индекс как в raw)."""
    cols = {_nk(c): c for c in raw.columns}
    c_ext = _mpick(cols, "внешнийидпозиции", "внешнийид", "externalid", "externalitemid")
    c_wms = _mpick(cols, "skuwmsid", "skuwms")
    c_name = _mpick(cols, "name", "товар", "наименование")
    c_zone = _mpick(cols, "zoneprefix", "зона")
    c_loc = _mpick(cols, "amountinlocation")
    c_res = _mpick(cols, "reserved", "резерв")
    c_av = _mpick(cols, "amountavailable")
    c_cell = _mpick(cols, "cellbarcode", "ячейка")
    c_inv = _mpick(cols, "originalinvoicenumber", "накладная")
    miss = [t for t, c in (("Внешний ID позиции / Sku Wms ID", c_ext or c_wms), ("name", c_name),
                           ("zone_prefix", c_zone), ("Amount In Location", c_loc)) if c is None]
    if miss:
        raise ValueError("В Magnit Stock нет колонок: " + ", ".join(miss) + ". Есть: " + ", ".join(map(str, raw.columns)))

    def txt(c):
        if c is None:
            return pd.Series("", index=raw.index)
        return raw[c].fillna("").astype(str).str.strip().str.replace(r"\.0$", "", regex=True)

    ext, wms = txt(c_ext), txt(c_wms)
    loc = _num(raw[c_loc])
    avail = _num(raw[c_av]) if c_av else loc
    res = _num(raw[c_res]) if c_res else (loc - avail)
    d = pd.DataFrame({"sku": ext.where(ext != "", wms), "ext": ext, "wms": wms, "name": txt(c_name),
                      "zone": txt(c_zone), "cell": txt(c_cell), "inv": txt(c_inv),
                      "loc": loc, "res": res, "avail": avail}, index=raw.index)
    zk = d["zone"].map(zkey)
    d["kind"] = np.where(zk.isin(RES_KEYS), "yes", np.where(zk.isin(NORES_KEYS), "no", "?"))
    return d[d["sku"] != ""]


def magnit_find(m: pd.DataFrame, items: list):
    """Ищет позиции по Внешнему ID, Sku Wms ID или названию. Возвращает (найденные строки, не найденные значения)."""
    kext, kwms, kname = m["ext"].map(_key), m["wms"].map(_key), m["name"].map(_key)
    total = pd.Series(False, index=m.index)
    missing = []
    for it in items:
        k = _key(it)
        if not k:
            continue
        mask = (kext == k) | (kwms == k) | (kname == k)
        if not mask.any() and len(k) >= 4:                      # запасной вариант: часть названия
            mask = kname.str.contains(k, regex=False)
        if mask.any():
            total |= mask
        else:
            missing.append(it)
    return m[total], missing


def magnit_rep(sub: pd.DataFrame, raw_all: pd.DataFrame, missing: list) -> dict:
    """Собирает отчёт в том же формате, что использует stock_now_tables / stock_now_excel."""
    g = sub.copy()
    zk = g["zone"].map(zkey)
    g["comment"] = g["kind"].map(KIND_TEXT)
    g["o"] = zk.map(ZONE_ORDER).fillna(99)
    g = g.sort_values(["sku", "o", "zone", "cell"]).reset_index(drop=True)
    return dict(date=datetime.now(TZ).date(), rows=g,
                raw=raw_all.loc[sub.index].reset_index(drop=True), missing=missing)


def read_item_list(up) -> list:
    """Список позиций из колонки A загруженного файла (xlsx или csv). Строка-заголовок пропускается."""
    if up.name.lower().endswith(".csv"):
        d = pd.read_csv(up, header=None, dtype=str, sep=None, engine="python")
    else:
        d = pd.read_excel(up, header=None, dtype=str)
    out = [str(x).strip() for x in d.iloc[:, 0].dropna()]
    out = [x for x in out if x]
    if out and _key(out[0]) in LIST_HEADERS:
        out = out[1:]
    return list(dict.fromkeys(out))


def stock_now_tables(rep: dict):
    """(по зонам: товар × зона, по товарам: итоги можно / нельзя зарезервировать)."""
    g = rep["rows"]
    if g.empty:
        return g.iloc[0:0], g.iloc[0:0]
    z = (g.groupby(["sku", "name", "zone", "kind", "o"], as_index=False)
         .agg(loc=("loc", "sum"), res=("res", "sum"), avail=("avail", "sum"), cells=("cell", "nunique"))
         .sort_values(["sku", "o", "zone"]).reset_index(drop=True))
    sk = g.groupby("sku").agg(name=("name", "first"), loc=("loc", "sum"), res=("res", "sum"), avail=("avail", "sum"))
    for col, kind in (("can", "yes"), ("no", "no"), ("other", "?")):
        sk[col] = g[g["kind"] == kind].groupby("sku")["avail"].sum().reindex(sk.index).fillna(0)
    return z, sk.reset_index()


def stock_now_excel(rep: dict) -> bytes:
    """Excel: остаток позиции в формате выгрузки (зона / ID / товар / в ячейке / резерв / доступно / комментарий)
    с итогами «Можно зарезервировать» и «Не зарезервируется», сводка по зонам, итоги по товарам, все данные."""
    g = rep["rows"]
    z, sk = stock_now_tables(rep)
    cols = ["zone_prefix", "Внешний ID позиции", "name", "Ячейка", "Накладная",
            "Amount In Location", "Reserved", "Amount Available", "Комментарии"]
    data = [[r.zone, r.sku, r.name, r.cell, r.inv, r.loc, r.res, r.avail, r.comment] for r in g.itertuples(index=False)]
    if data and not sk.empty:
        data.append([None] * 9)
        for r in sk.itertuples(index=False):
            data.append([None, r.sku, r.name, None, None, None, "Можно зарезервировать", r.can, None])
            data.append([None, r.sku, r.name, None, None, None, "Не зарезервируется", r.no, None])
            if r.other:
                data.append([None, r.sku, r.name, None, None, None, "Зона не в правилах", r.other, None])
    main = pd.DataFrame(data, columns=cols)

    zone_out = z.rename(columns={"sku": "Внешний ID позиции", "name": "Товар", "zone": "Зона", "loc": "В ячейках",
                                 "res": "Резерв", "avail": "Доступно", "cells": "Ячеек"})
    if not zone_out.empty:
        zone_out["Можно резервировать"] = zone_out["kind"].map(KIND_TEXT)
        zone_out = zone_out[["Внешний ID позиции", "Товар", "Зона", "Можно резервировать", "Ячеек",
                             "В ячейках", "Резерв", "Доступно"]]
    sk_out = sk.rename(columns={"sku": "Внешний ID позиции", "name": "Товар", "loc": "В ячейках", "res": "Резерв",
                                "avail": "Доступно", "can": "Можно зарезервировать",
                                "no": "Не зарезервируется", "other": "Зона не в правилах"})
    params = pd.DataFrame([
        ["Источник", "Magnit Stock (Metabase)"],
        ["Сформировано", f"{datetime.now(TZ):%d.%m.%Y %H:%M}"],
        ["Зоны с резервом", ", ".join(RESERVE_ZONES)],
        ["Зоны без резерва", ", ".join(NO_RESERVE_ZONES)],
        ["Можно зарезервировать", "сумма «Amount Available» по зонам с резервом"],
        ["Не зарезервируется", "сумма «Amount Available» по зонам без резерва"],
        ["Reserved", "резерв из Magnit Stock (или «В ячейках» − «Доступно», если колонки нет)"],
    ], columns=["Параметр", "Значение"])
    sheets = {"Остаток позиции": main, "Сводка по зонам": zone_out, "Итого по товарам": sk_out}
    if rep.get("raw") is not None and not rep["raw"].empty:
        sheets["Все данные"] = rep["raw"]
    if rep.get("missing"):
        sheets["Не найдено"] = pd.DataFrame({"Не найдено в Magnit Stock": rep["missing"]})
    sheets["Правила"] = params
    return to_xlsx(sheets)


# ───────────────────────── Metabase: «Хаб отгрузка ББ» ─────────────────────────
def _hk(x) -> str:
    """Ключ названия колонки: без регистра, пробелов и знаков, «ё» = «е»."""
    return re.sub(r"[^0-9a-zа-я]", "", str(x).lower().replace("ё", "е"))


def _route_name(s: str) -> str:
    """«Ташкент Бродвей 1M-0026817, 1M-0026994» -> «Ташкент Бродвей» (название маршрута без номеров заказов)."""
    s = str(s).strip()
    m = re.search(r"\s+\w{1,3}-\d{4,}", s)
    name = s[:m.start()].strip() if m else s
    return name or s


@st.cache_data(ttl=300, show_spinner=False)
def load_hub() -> pd.DataFrame:
    """«Хаб отгрузка ББ» (Metabase, таблица 61): одна строка = одна коробка (биг-бэг).
    Колонки результата: route_full (Маршрут как в таблице), route (название), bar (Штрихкод ББ),
    date (Дата отправки), ok (True — принят)."""
    cfg, _, _ = _mb_cfg()
    tid = int(cfg.get("hub_table", MB_HUB_T) or MB_HUB_T)
    db, f = mb_meta(tid)

    raw = None
    key = next((k for k in f if k.startswith("датаотправ")), None)   # читаем только последние HUB_DAYS дней
    if key and "Date" in f[key][1]:
        since = (datetime.now(TZ).date() - timedelta(days=HUB_DAYS)).isoformat()
        try:
            raw = mb_run(db, {"source-table": tid,
                              "filter": [">=", ["field", f[key][0], None], since]})
        except Exception:
            raw = None
    if raw is None:
        raw = mb_run(db, {"source-table": tid})
    if raw.empty:
        raise ValueError("Таблица «Хаб отгрузка ББ» вернула 0 строк.")

    cols = {_hk(c): c for c in raw.columns}

    def find(*keys, no=()):
        for k, c in cols.items():
            if all(x in k for x in keys) and not any(x in k for x in no):
                return c
        return None

    c_route = cols.get("маршрут") or find("маршрут", no=("обновл",)) or find("route")
    c_acc = find("статус", "приемк") or find("accept")
    c_bar = find("штрихкод") or find("barcode")
    c_date = find("датаотправ") or find("senddate")
    if not all([c_route, c_acc, c_bar, c_date]):
        raise ValueError("В «Хаб отгрузка ББ» не найдены колонки: Маршрут, Статус приёмки, Штрихкод ББ, Дата отправки. "
                         "Есть: " + ", ".join(map(str, raw.columns)))

    status = raw[c_acc].fillna("").astype(str).str.strip().str.lower()
    d = pd.DataFrame({
        "route_full": raw[c_route].fillna("").astype(str).str.strip(),
        "bar": raw[c_bar].fillna("").astype(str).str.strip(),
        "date": _to_date(raw[c_date].fillna("")),
        "ok": status.str.startswith("принят"),
    })
    d["route"] = d["route_full"].map(_route_name)
    d = d[d["date"].notna() & (d["route_full"] != "")]
    has = d["bar"] != ""
    d = pd.concat([d[has].drop_duplicates(["bar", "date"], keep="last"), d[~has]], ignore_index=True)
    return d.sort_values(["date", "route", "bar"]).reset_index(drop=True)


def hub_by_date(h: pd.DataFrame) -> pd.DataFrame:
    """По датам отправки: ok — принято коробок, no — не принято, total — всего."""
    g = h.groupby("date").agg(ok=("ok", "sum"), total=("ok", "size"))
    g["no"] = g["total"] - g["ok"]
    return g.sort_index()


def hub_card_dates(g: pd.DataFrame, today: date) -> list:
    """Какие даты показать на карточке: со вчера и дальше (если таких нет — последние даты)."""
    ds = [d for d in g.index if d >= today - timedelta(days=1)]
    if not ds:
        ds = list(g.index)[-3:]
    return ds[:HUB_CARD_DATES]


def hub_card(g, today: date, err: str = "") -> str:
    """Карточка «Хаб ББ»: мелкая таблица Дата / Принято коробок / Не принято коробок."""
    if err or g is None or g.empty:
        body = "<div class='none'>Нет данных из Metabase</div>"
    else:
        rows = "".join(
            f"<div class='hr'><span>{fmt_date(d)}, {WEEKDAYS[d.weekday()]}</span>"
            f"<b class='g'>{n(g.loc[d, 'ok'])}</b><b class='r'>{n(g.loc[d, 'no'])}</b></div>"
            for d in hub_card_dates(g, today))
        body = ("<div class='hr hh'><span>Дата</span><span>Принято коробок</span><span>Не принято коробок</span></div>"
                + rows)
    return (f"<div class='kpi left nov'><div class='ico'>{ic('box', '#FFFFFF', 20)}</div>"
            f"<div class='mid'><div class='l'>Хаб ББ</div>"
            f"<div class='s'>приёмка коробок · нажмите, чтобы открыть маршруты</div></div>"
            f"<div class='hubt'>{body}</div></div>")


def _xl_safe(df: pd.DataFrame) -> pd.DataFrame:
    """Очень длинные целые (ID по 19 цифр) Excel округляет — пишем их текстом."""
    df = df.copy()
    for c in df.columns:
        s = df[c]
        if pd.api.types.is_integer_dtype(s) and len(s) and s.abs().max() >= 10 ** 15:
            df[c] = s.astype(str)
    return df


def to_xlsx(sheets: dict) -> bytes:
    from openpyxl.styles import Font, PatternFill
    out = BytesIO()

    def _cell(v):
        if isinstance(v, (list, dict, tuple, set)):
            return json.dumps(v, ensure_ascii=False, default=str)
        return v

    with pd.ExcelWriter(out, engine="openpyxl", date_format="DD.MM.YYYY",
                        datetime_format="DD.MM.YYYY HH:MM") as xw:
        for name, frame in sheets.items():
            nm = name[:31]
            frame = _xl_safe(frame).copy()
            # уникальные строковые названия колонок
            seen, cols = {}, []
            for c in frame.columns:
                c = str(c)
                seen[c] = seen.get(c, 0) + 1
                cols.append(c if seen[c] == 1 else f"{c} ({seen[c]})")
            frame.columns = cols
            for c in cols:
                if frame[c].dtype == object:
                    frame[c] = frame[c].map(_cell)
            frame.to_excel(xw, sheet_name=nm, index=False)
            ws = xw.sheets[nm]
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DCE8F2")
            for i, col in enumerate(cols, 1):
                sample = frame[col].head(300).map(lambda x: "" if pd.isna(x) else str(x)) \
                    if frame[col].dtype != object else frame[col].head(300).map(lambda x: "" if x is None else str(x))
                w = max([len(col)] + [len(x) for x in sample]) + 2
                ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(w, 10), 60)
    return out.getvalue()


def stock_report(h: pd.DataFrame, skus: list, d1: date, d2: date, raw: pd.DataFrame = None) -> dict:
    """Хронология движений, резервы и остатки по выбранным товарам между ежедневными снимками.
    Резерв = «В ячейках» (Amount In Location) − «Доступно» (Amount Available)."""
    all_dates = sorted(h["date"].unique())
    prev = [x for x in all_dates if x < d1]
    dates = prev[-1:] + [x for x in all_dates if d1 <= x <= d2]
    base = h[h["sku"].isin(skus)]
    sub = base[base["date"].isin(dates)]
    names = base.groupby("sku")["name"].first()
    keys = ["sku", "cell", "inv"]

    # остаток по дням: в ячейках / доступно / резерв
    if sub.empty:
        daily = pd.DataFrame(columns=["sku", "date", "loc", "avail"])
    else:
        daily = sub.groupby(["sku", "date"]).agg(loc=("loc", "sum"), avail=("avail", "sum"))
        idx = pd.MultiIndex.from_product([sorted(sub["sku"].unique()), dates], names=["sku", "date"])
        daily = daily.reindex(idx, fill_value=0).reset_index()
    daily["res"] = daily["loc"] - daily["avail"]

    # движения между соседними снимками (по ячейкам и накладным)
    parts = []
    if len(dates) >= 2 and not sub.empty:
        g = sub.groupby(keys + ["date"], as_index=False).agg(loc=("loc", "sum"), avail=("avail", "sum"),
                                                              zone=("zone", "first"))
        pq = (g.pivot_table(index=keys, columns="date", values="loc", aggfunc="sum", fill_value=0)
              .reindex(columns=dates, fill_value=0))
        pa = (g.pivot_table(index=keys, columns="date", values="avail", aggfunc="sum", fill_value=0)
              .reindex(columns=dates, fill_value=0).reindex(pq.index, fill_value=0))
        zone_of = g.groupby(keys)["zone"].first()
        for a, b in zip(dates[:-1], dates[1:]):
            ch = pd.DataFrame({"was": pq[a], "now": pq[b], "aw": pa[a], "an": pa[b]})
            ch["delta"] = ch["now"] - ch["was"]
            ch["rw"] = ch["was"] - ch["aw"]
            ch["rn"] = ch["now"] - ch["an"]
            ch = ch[(ch["delta"] != 0) | (ch["rw"] != ch["rn"])].reset_index()
            if ch.empty:
                continue
            ch["date"], ch["prev"] = b, a
            ch["zone"] = [zone_of.get((r.sku, r.cell, r.inv), "") for r in ch.itertuples()]
            parts.append(ch)
    cols_mv = ["Дата снимка", "Предыдущий снимок", "Sku Wms ID", "Товар", "Зона", "Ячейка", "Накладная",
               "Было в ячейке", "Стало в ячейке", "Изменение", "Доступно было", "Доступно стало",
               "Резерв было", "Резерв стало", "Тип движения"]
    if parts:
        mv = pd.concat(parts, ignore_index=True)
        mv["type"] = np.select([mv["delta"] == 0, mv["was"] == 0, mv["now"] == 0, mv["delta"] > 0],
                               ["Изменение резерва", "Поступление в ячейку", "Ячейка опустела", "Увеличение"],
                               "Уменьшение")
        mv["name"] = mv["sku"].map(names)
        mv = mv.sort_values(["sku", "date", "cell"]).reset_index(drop=True)
        mv_out = mv.rename(columns={"date": "Дата снимка", "prev": "Предыдущий снимок", "sku": "Sku Wms ID",
                                    "name": "Товар", "zone": "Зона", "cell": "Ячейка", "inv": "Накладная",
                                    "was": "Было в ячейке", "now": "Стало в ячейке", "delta": "Изменение",
                                    "aw": "Доступно было", "an": "Доступно стало",
                                    "rw": "Резерв было", "rn": "Резерв стало",
                                    "type": "Тип движения"})[cols_mv]
    else:
        mv = pd.DataFrame(columns=["sku", "delta"])
        mv_out = pd.DataFrame(columns=cols_mv)

    # итоги по товарам
    first, last = (dates[0], dates[-1]) if dates else (None, None)
    d_loc = daily.set_index(["sku", "date"])["loc"] if not daily.empty else pd.Series(dtype=float)
    d_av = daily.set_index(["sku", "date"])["avail"] if not daily.empty else pd.Series(dtype=float)
    rows = []
    for sku in skus:
        m = mv[(mv["sku"] == sku) & (mv["delta"] != 0)] if not mv.empty else mv
        inc = float(m.loc[m["delta"] > 0, "delta"].sum()) if len(m) else 0.0
        out_ = float(-m.loc[m["delta"] < 0, "delta"].sum()) if len(m) else 0.0
        start = float(d_loc.get((sku, first), 0)) if first else 0.0
        end = float(d_loc.get((sku, last), 0)) if last else 0.0
        av_end = float(d_av.get((sku, last), 0)) if last else 0.0
        rows.append({"Sku Wms ID": sku, "Товар": names.get(sku, ""), "Остаток на начало": start,
                     "Остаток на конец": end, "Доступно на конец": av_end, "Резерв на конец": end - av_end,
                     "Поступило": inc, "Убыло": out_, "Чистое изменение": end - start, "Движений": int(len(m))})
    summ = pd.DataFrame(rows)

    cols_daily = ["Дата снимка", "Sku Wms ID", "Товар", "В ячейках", "Доступно", "Резерв"]
    daily_out = daily.copy()
    if not daily_out.empty:
        daily_out["Товар"] = daily_out["sku"].map(names)
        daily_out = daily_out.rename(columns={"date": "Дата снимка", "sku": "Sku Wms ID", "loc": "В ячейках",
                                              "avail": "Доступно", "res": "Резерв"})[cols_daily]
    else:
        daily_out = pd.DataFrame(columns=cols_daily)

    # последний остаток по ячейкам (на последнюю дату снимка)
    cols_last = ["Дата снимка", "Sku Wms ID", "Товар", "Зона", "Ячейка", "Накладная", "В ячейках", "Доступно", "Резерв"]
    lst = sub[sub["date"] == last] if last else sub.iloc[0:0]
    if not lst.empty:
        lst = (lst.groupby(keys, as_index=False).agg(zone=("zone", "first"), loc=("loc", "sum"), avail=("avail", "sum")))
        lst["res"] = lst["loc"] - lst["avail"]
        lst["date"] = last
        lst["name"] = lst["sku"].map(names)
        last_out = (lst.sort_values(["sku", "cell"])
                    .rename(columns={"date": "Дата снимка", "sku": "Sku Wms ID", "name": "Товар", "zone": "Зона",
                                     "cell": "Ячейка", "inv": "Накладная", "loc": "В ячейках",
                                     "avail": "Доступно", "res": "Резерв"})[cols_last])
    else:
        last_out = pd.DataFrame(columns=cols_last)

    # все данные по товару (все колонки Metabase) + резерв
    all_out = pd.DataFrame()
    if raw is not None and not raw.empty and len(raw) == len(h):
        order = h.sort_values(["sku", "date", "cell"]).index
        all_out = raw.loc[order].copy()
        all_out["Резерв (В ячейке − Доступно)"] = (h["loc"] - h["avail"]).loc[order].values
        all_out = all_out.reset_index(drop=True)

    params = pd.DataFrame([
        ["Период", f"{d1:%d.%m.%Y} — {d2:%d.%m.%Y}"],
        ["Начальный снимок (база)", f"{first:%d.%m.%Y}" if first else "—"],
        ["Конечный снимок", f"{last:%d.%m.%Y}" if last else "—"],
        ["Снимков в расчёте", ", ".join(f"{x:%d.%m}" for x in dates) or "—"],
        ["Товаров", len(skus)],
        ["Сформировано", f"{datetime.now(TZ):%d.%m.%Y %H:%M}"],
        ["Резерв", "Резерв = «В ячейках» (Amount In Location) − «Доступно» (Amount Available). "
                   "Например, в ячейке 10, доступно 5 — резерв 5."],
        ["Примечание", "Движения — разница остатка «В ячейках» между соседними снимками, а не точное время операций. "
                       "Перемещение между ячейками видно как «убыло» в одной и «поступило» в другой."],
    ], columns=["Параметр", "Значение"])
    return dict(dates=dates, movements=mv_out, summary=summ, daily=daily_out, last=last_out,
                all=all_out, params=params)


def fmt_date(d) -> str:
    return d.strftime("%d.%m") if pd.notna(d) and d is not None else "—"


def n(x) -> str:
    return f"{int(x):,}".replace(",", " ")


def esc(x) -> str:
    return html.escape(str(x))


def day_summary(df: pd.DataFrame) -> pd.DataFrame:
    g = df.groupby("Дата").agg(Заказов=("Заказ", "count"), Собрано=("Собран", "sum"), Штук=("Штук", "sum"))
    g["Штук собрано"] = df[df["Собран"]].groupby("Дата")["Штук"].sum().reindex(g.index).fillna(0).astype(int)
    g["Осталось"] = g["Заказов"] - g["Собрано"]
    g["Процент"] = (g["Собрано"] / g["Заказов"] * 100).round().astype(int)
    return g.sort_index()


# ── нормализация названий магазинов (регистр, пробелы, дефисы, латиница-двойники) ──
_LOOKALIKE = str.maketrans("eaopcxykmt", "еаорсхукмт")


def norm_store(x: str) -> str:
    x = str(x).lower().replace("ё", "е").translate(_LOOKALIKE)
    return re.sub(r"[\s\-_.]+", "", x)


EXCLUDED_KEYS = {norm_store(x) for x in EXCLUDE_STORES}


@st.cache_data(ttl=CACHE_SECONDS, show_spinner=False)
def load_coords() -> dict:
    """Вкладка «Кординаты»: name / latitude / longitude -> {нормализованное имя: (lat, lon)}.
    Явно неверные координаты (вне Узбекистана, например долгота = широта) пропускаются."""
    sheet = _client().open_by_key(ARCHIVE_SHEET_ID)
    sheets = sheet.worksheets()
    ws = (next((w for w in sheets if str(w.id) == COORDS_GID), None)
          or next((w for w in sheets if "ордина" in w.title.lower()), None))
    if ws is None:
        raise ValueError("Не найдена вкладка с координатами («Кординаты»).")
    rows = ws.get_all_values()
    head = [" ".join(h.split()).lower() for h in rows[0]]

    def col(keys, default):
        for i, h in enumerate(head):
            if any(k in h for k in keys):
                return i
        return default

    i_name = col(["name", "назван"], 0)
    i_lat = col(["latitude", "широт", "lat"], 3)
    i_lon = col(["longitude", "долгот", "lon"], 4)

    def num(s):
        try:
            return float(str(s).replace(",", ".").replace(" ", ""))
        except ValueError:
            return None

    out = {}
    for r in rows[1:]:
        if len(r) <= max(i_name, i_lat, i_lon):
            continue
        name, lat, lon = r[i_name].strip(), num(r[i_lat]), num(r[i_lon])
        if not name or lat is None or lon is None:
            continue
        if not (35 <= lat <= 46 and 55 <= lon <= 75):
            continue
        out[norm_store(name)] = (lat, lon)
    return out


def split_routing(frame: pd.DataFrame):
    """Возвращает (сводная по городу/магазину/адресу, исключённые заказы)."""
    mask = frame["Магазин"].map(lambda s: norm_store(s) in EXCLUDED_KEYS)
    keep, dropped = frame[~mask], frame[mask]
    g = (keep.groupby(["Город", "Магазин", "Адрес"], as_index=False)["Штук"].sum()
         .sort_values(["Город", "Магазин", "Адрес"]).reset_index(drop=True))
    return g, dropped


def archive_stats(raw: pd.DataFrame) -> dict:
    """Средние по архиву: сколько штук / точек / во сколько уезжала машина №1, №2, … и к какой машине
    обычно относится магазин. Машина = рейс внутри дня (по времени фактической отгрузки)."""
    raw = raw.copy()
    raw.columns = [" ".join(str(c).split()) for c in raw.columns]

    def find(*keys):
        return next((c for c in raw.columns if all(k in c.lower() for k in keys)), None)

    c_shop, c_qty, c_trip = find("название"), find("кол-во"), find("рейс")
    c_stat, c_fact = find("статус"), find("факт")
    if not all([c_shop, c_qty, c_trip, c_fact]):
        raise ValueError("В архиве не найдены колонки: название магазина, кол-во штук, рейс, дата отгрузки факт.")

    d = pd.DataFrame({
        "shop": raw[c_shop].str.strip(),
        "qty": pd.to_numeric(raw[c_qty].str.replace(r"\s", "", regex=True), errors="coerce").fillna(0),
        "trip": raw[c_trip].str.strip(),
        "fact": pd.to_datetime(raw[c_fact].str.strip(), format="%d.%m.%Y %H:%M", errors="coerce"),
    })
    if c_stat:
        d["status"] = raw[c_stat].str.strip().str.upper()
    d = d[d["fact"].notna() & (d["trip"] != "")]
    d = d[d["fact"].dt.date >= ARCHIVE_FROM]
    if c_stat:
        shipped = d[d["status"] == "ОТГРУЖЕН"]
        d = shipped if not shipped.empty else d
    if d.empty:
        raise ValueError(f"В архиве нет отгрузок с {ARCHIVE_FROM:%d.%m.%Y}.")

    d["day"] = d["fact"].dt.date
    trips = (d.groupby(["day", "trip"]).agg(qty=("qty", "sum"), stores=("shop", "nunique"), t=("fact", "min"))
             .reset_index())
    trips["tn"] = pd.to_numeric(trips["trip"], errors="coerce").fillna(999)
    trips = trips.sort_values(["day", "t", "tn"])
    trips["k"] = trips.groupby("day").cumcount() + 1          # какая по счёту машина в этот день
    trips["min"] = trips["t"].dt.hour * 60 + trips["t"].dt.minute
    by_k = trips.groupby("k").agg(qty=("qty", "mean"), stores=("stores", "mean"), time=("min", "mean"))

    d = d.merge(trips[["day", "trip", "k"]], on=["day", "trip"])
    d["key"] = d["shop"].map(norm_store)
    return dict(
        qty={int(k): float(v) for k, v in by_k["qty"].items()},
        stores={int(k): float(v) for k, v in by_k["stores"].items()},
        time={int(k): float(v) for k, v in by_k["time"].items()},
        rank={k: float(v) for k, v in d.groupby("key")["k"].mean().items()},
        avg_all=float(trips["qty"].mean()),
        n_days=int(trips["day"].nunique()),
        date_from=trips["day"].min(), date_to=trips["day"].max(),
    )


def get_archive_stats() -> dict:
    """Средние пересчитываются при каждом чтении архива (кэш = CACHE_SECONDS, кнопка «Обновить» сбрасывает)."""
    return archive_stats(load_ship_tabs()["archive"])


_SHIP_COLS = ["Заказ", "Магазин", "Штук", "Дата", "Рейс", "Отгружен", "Факт"]


def parse_ship(raw) -> pd.DataFrame:
    """Строки вкладки отгрузок -> Заказ, Магазин, Штук, Дата (план), Рейс, Отгружен, Факт."""
    if raw is None or raw.empty:
        return pd.DataFrame(columns=_SHIP_COLS)
    raw = raw.copy()
    raw.columns = [" ".join(str(c).split()) for c in raw.columns]

    def find(*keys):
        return next((c for c in raw.columns if all(k in c.lower() for k in keys)), None)

    c_order, c_shop, c_qty = find("№"), find("название"), find("кол-во")
    c_plan, c_trip, c_stat, c_fact = find("дата", "план"), find("рейс"), find("статус"), find("факт")
    if not all([c_order, c_shop, c_qty, c_plan]):
        return pd.DataFrame(columns=_SHIP_COLS)
    d = pd.DataFrame({
        "Заказ": raw[c_order].str.strip(),
        "Магазин": raw[c_shop].str.strip(),
        "Штук": pd.to_numeric(raw[c_qty].str.replace(r"\s", "", regex=True), errors="coerce").fillna(0).astype(int),
        "Дата": pd.to_datetime(raw[c_plan], dayfirst=True, errors="coerce").dt.date,
        "Рейс": raw[c_trip].str.strip() if c_trip else "",
        "Отгружен": (raw[c_stat].str.strip().str.upper() == "ОТГРУЖЕН") if c_stat else False,
        "Факт": (pd.to_datetime(raw[c_fact].str.strip(), format="%d.%m.%Y %H:%M", errors="coerce")
                 if c_fact else pd.NaT),
    })
    return d[d["Заказ"].str.match(r"^\s*\w{1,3}-\d{4,}") & d["Дата"].notna()].reset_index(drop=True)


def build_ship(tabs: dict) -> pd.DataFrame:
    """Все заказы из «Маршруты New» + «Архив» без дублей (приоритет у «Маршруты New»), без исключённых магазинов."""
    d = pd.concat([parse_ship(tabs.get("new")), parse_ship(tabs.get("archive"))], ignore_index=True)
    d = d.drop_duplicates("Заказ", keep="first")
    return d[~d["Магазин"].map(lambda x: norm_store(x) in EXCLUDED_KEYS)].reset_index(drop=True)


def pick_ship_date(ship: pd.DataFrame, today: date):
    ds = sorted(ship["Дата"].unique())
    if not ds:
        return None
    if today in ds:
        return today
    prev = [x for x in ds if x <= today]
    return prev[-1] if prev else ds[0]


def ship_summary(ship: pd.DataFrame, d: date) -> dict:
    s_ = ship[ship["Дата"] == d]
    tr = (s_[s_["Рейс"] != ""].groupby("Рейс")
          .agg(qty=("Штук", "sum"), done=("Отгружен", "all"), t=("Факт", "min")).reset_index())
    tr["tn"] = pd.to_numeric(tr["Рейс"], errors="coerce").fillna(999)
    tr = tr.sort_values(["tn", "Рейс"])
    return dict(total=len(s_), done=int(s_["Отгружен"].sum()),
                pcs=int(s_["Штук"].sum()), pcs_done=int(s_.loc[s_["Отгружен"], "Штук"].sum()),
                trips=tr.to_dict("records"))


def sync_icon(ship: pd.DataFrame, d: date) -> str:
    """Значок синхронизации: ✓ — все отгруженные за дату заказы уже есть во вкладке «Статус по отгрузкам»,
    вращающийся — часть ещё не попала. Если вкладка недоступна или отгруженных нет — значка нет."""
    try:
        have = load_status_orders()
    except Exception:
        return ""
    shipped = ship[(ship["Дата"] == d) & ship["Отгружен"]]["Заказ"].str.strip().str.upper()
    if shipped.empty:
        return ""
    missing = int((~shipped.isin(have)).sum())
    if missing == 0:
        return "<span class='sync ok' title='Все данные попали в таблицу «Статус по отгрузкам»'>✓</span>"
    return (f"<span class='sync spin' title='В таблицу «Статус по отгрузкам» не попало заказов: {missing} "
            f"из {len(shipped)}. Данные обновляются'></span>")


def hhmm(minutes: float) -> str:
    m = int(round(minutes))
    return f"{m // 60:02d}:{m % 60:02d}"


def deadline_minutes(addr: str):
    m = re.search(r"\bдо\s*(\d{1,2})[:.](\d{2})", str(addr).lower())
    return int(m.group(1)) * 60 + int(m.group(2)) if m else None


def ordinal(k: int) -> str:
    return f"{ORDINALS[k - 1]} машина" if k <= len(ORDINALS) else f"{k}-я машина"


def plan_trips(pivot: pd.DataFrame, stats: dict, limit: int = 0) -> list:
    """Пример раскладки точек по машинам: магазины идут в том порядке, как их обычно развозили
    (по архиву), и набираются, пока машина не достигнет среднего объёма для её номера.
    limit > 0 — жёсткий потолок штук на одну машину (магазин, который сам больше лимита, едет отдельно и помечается)."""
    def target(k):
        t = stats["qty"].get(k) or stats["avg_all"]
        return min(t, limit) if limit else t

    rows = []
    for r in pivot.itertuples(index=False):
        dl = deadline_minutes(r.Адрес)
        hot = dl is not None and dl < 12 * 60                     # «ДО 09:45» и т.п. — срочные
        rk = stats["rank"].get(norm_store(r.Магазин), 98.0)       # новых магазинов в архиве нет — в хвост
        rows.append(dict(city=r.Город, shop=r.Магазин, addr=r.Адрес, qty=int(r.Штук),
                         rk=min(rk, 1.0) if hot else rk, dl=dl, hot=hot))

    cities = sorted({x["city"] for x in rows}, key=lambda c: (norm_store(c) != norm_store("Ташкент"), c))
    machines = []
    for city in cities:                                           # другой город — всегда отдельная машина
        items = sorted((x for x in rows if x["city"] == city),
                       key=lambda x: (x["rk"], x["dl"] if x["dl"] is not None else 9999, x["shop"]))
        local, cur, cur_sum = [], [], 0
        for it in items:
            t = target(len(machines) + len(local) + 1)
            over_limit = bool(limit) and cur_sum + it["qty"] > limit
            if cur and (cur_sum >= t or over_limit or
                        (cur_sum + it["qty"] > t * 1.15 and cur_sum >= t * 0.6)):
                local.append(cur)
                cur, cur_sum = [], 0
            cur.append(it)
            cur_sum += it["qty"]
        if cur:
            local.append(cur)
        if len(local) >= 2:                                       # маленький «хвост» присоединяем к предыдущей
            last, prev = sum(x["qty"] for x in local[-1]), sum(x["qty"] for x in local[-2])
            k_last = len(machines) + len(local)
            if (last < 0.35 * target(k_last) and prev + last <= 1.35 * target(k_last - 1)
                    and (not limit or prev + last <= limit)):
                local[-2] = local[-2] + local.pop()
        machines.extend(local)

    out = []
    for i, st_ in enumerate(machines, 1):
        q = sum(x["qty"] for x in st_)
        out.append(dict(n=i, stores=st_, qty=q, target=target(i), time=stats["time"].get(i),
                        over=bool(limit) and q > limit))
    return out


def machine_meta(m: dict) -> str:
    return f"{n(m['qty'])} шт · точек: {len(m['stores'])}" + (" · ⚠ выше лимита" if m.get("over") else "")


# ── маршруты на картах: точки в том порядке, как стоят в машине ──
def _chunks(pts: list) -> list:
    """Делит точки на куски до 10 (лимит ссылок карт). Возвращает [(origin|None, [точки])]."""
    res = [(None, pts[:10])]
    s = 9
    while s < len(pts) - 1:
        res.append((pts[s], pts[s + 1:s + 10]))
        s += 9
    return res


def gmaps_url(origin, pts: list) -> str:
    url = f"https://www.google.com/maps/dir/?api=1&travelmode=driving&destination={pts[-1]}"
    if origin:
        url += f"&origin={origin}"
    if len(pts) > 1:
        url += "&waypoints=" + "%7C".join(pts[:-1])
    return url


def ymaps_url(origin, pts: list) -> str:
    # пустая первая точка = «от моего местоположения»
    return "https://yandex.ru/maps/?rtt=auto&rtext=" + "~".join([origin or ""] + pts)


def map_buttons(m: dict, coords: dict) -> str:
    pts, missing = [], []
    for s_ in m["stores"]:
        c = coords.get(norm_store(s_["shop"]))
        if c:
            pts.append(f"{c[0]:.6f},{c[1]:.6f}")
        else:
            missing.append(s_["shop"])
    out = ""
    if pts:
        parts = _chunks(pts)
        many = len(parts) > 1
        links = []
        for name, fn in (("Google Maps", gmaps_url), ("Яндекс Карты", ymaps_url)):
            for i, (origin, chunk) in enumerate(parts, 1):
                label = f"{name} {i}" if many else name
                links.append(f"<a class='mapbtn' href='{fn(origin, chunk)}' target='_blank' "
                             f"rel='noopener noreferrer'>{label}</a>")
        out += f"<span class='maps'>{''.join(links)}</span>"
    if missing:
        out += f"<span class='chip wait'>нет координат: {esc(', '.join(missing))}</span>"
    return out


def machine_chips(m: dict, coords: dict = None) -> str:
    chips = (f"<span class='chip'><b>{n(m['qty'])}</b> шт</span>"
             f"<span class='chip'>Точек <b>{len(m['stores'])}</b></span>")
    if m.get("over"):
        chips += "<span class='chip wait'>⚠ выше лимита</span>"
    bar = ""
    if m.get("target"):                                           # без лимита полоску загрузки не рисуем
        pct = round(m["qty"] / m["target"] * 100)
        hi = " hi" if pct > 115 else ""
        bar = (f"<span class='load{hi}'><i style='width:{min(pct, 100)}%'></i></span>"
               f"<span class='loadpct'>{pct}%</span>")
    extra = map_buttons(m, coords) if coords else ""
    return f"<div class='chips'>{chips}{bar}{extra}</div>"


def trips_table(machines: list, coords: dict = None) -> str:
    parts = []
    for m in machines:
        parts.append(f"<tr class='car'><td class='carname'>{ordinal(m['n'])}</td>"
                     f"<td class='carmeta' colspan='2'>{machine_chips(m, coords)}</td></tr>"
                     "<tr class='band thin'><td colspan='3'></td></tr>")
        for s_ in m["stores"]:
            parts.append(f"<tr class='{'hot' if s_['hot'] else ''}'><td class='shop'>{esc(s_['shop'])}</td>"
                         f"<td>{esc(s_['addr'])}</td><td class='num'>{n(s_['qty'])}</td></tr>")
        parts.append("<tr class='band'><td colspan='3'></td></tr>")
    return f"<div class='cars-wrap'><table class='cars'>{''.join(parts)}</table></div>"


def trips_excel(machines: list, day_label: str, limit: int = 0, example: bool = True) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Рейсы (пример)" if example else "Рейсы"
    if example:
        note = f"ПРИМЕР маршрутизации на {day_label}. Это расчёт по средним значениям архива, а не факт."
    else:
        note = f"Рейсы на {day_label} (собраны вручную)."
    if limit:
        note += f" Лимит на машину: {limit} шт."
    ws.append([note])
    ws.merge_cells("A1:C1")
    ws["A1"].font = Font(bold=True, color="C00000" if example else "14232E")
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    yellow, blue = PatternFill("solid", fgColor="FFEB3B"), PatternFill("solid", fgColor="1565C0")
    hot_fill = PatternFill("solid", fgColor="FBEFD2")

    def band(h):
        ws.append([None, None, None])
        for c in ws[ws.max_row]:
            c.fill = blue
        ws.row_dimensions[ws.max_row].height = h

    for m in machines:
        ws.append([ordinal(m["n"]), machine_meta(m), None])
        ws.cell(row=ws.max_row, column=1).fill = yellow
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        band(6)
        for s_ in m["stores"]:
            ws.append([s_["shop"], s_["addr"], s_["qty"]])
            row = ws.max_row
            ws.cell(row=row, column=1).font = Font(bold=True)
            ws.cell(row=row, column=3).number_format = "#,##0"
            if s_["hot"]:
                for c in ws[row]:
                    c.fill = hot_fill
                    c.font = Font(bold=True)
        band(10)
    ws.append(["Итого", None, sum(m["qty"] for m in machines)])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=3).number_format = "#,##0"
    for col, w in zip("ABC", (24, 80, 12)):
        ws.column_dimensions[col].width = w
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# ───────────────────────── Excel ─────────────────────────
def backlog_excel(frame: pd.DataFrame, extra: pd.DataFrame = None) -> bytes:
    """Файл из двух вкладок.
    «Бэклог» — все несобранные заказы по датам: ДАТА и КОЛ-ВО (шт); новые (не маршрутизированные) заказы добавляются.
    «Сводная» — все даты по неделям (пн–вс): заказов, всего шт, собрано шт, осталось шт, % собрано + итоги по неделям."""
    from openpyxl.styles import Font, PatternFill, Alignment

    # ── вкладка «Бэклог» ──
    parts = [frame.loc[~frame["Собран"], ["Дата", "Штук"]]]
    if extra is not None and not extra.empty:
        parts.append(extra.loc[extra["Дата"].notna(), ["Дата", "Штук"]])
    b = (pd.concat(parts, ignore_index=True).groupby("Дата")["Штук"].sum().reset_index()
         .rename(columns={"Дата": "ДАТА", "Штук": "КОЛ-ВО"}).sort_values("ДАТА"))

    # ── данные для «Сводной»: все заказы + новые (они не собраны) ──
    src = frame[["Дата", "Штук", "Собран"]].copy()
    if extra is not None and not extra.empty:
        ex = extra.loc[extra["Дата"].notna(), ["Дата", "Штук"]].assign(Собран=False)
        src = pd.concat([src, ex], ignore_index=True)
    src = src[src["Дата"].notna()].copy()
    src["Заказов"] = 1
    src["Собрано_шт"] = src["Штук"].where(src["Собран"].astype(bool), 0)
    g = (src.groupby("Дата").agg(Заказов=("Заказов", "sum"), Штук=("Штук", "sum"), Собрано=("Собрано_шт", "sum"))
         .reset_index().sort_values("Дата"))
    g["Нед"] = g["Дата"].map(lambda d: d - timedelta(days=d.weekday()))

    out = BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        b.to_excel(xw, sheet_name="Бэклог", index=False)
        ws = xw.sheets["Бэклог"]
        last = len(b) + 1
        ws.cell(row=last + 1, column=1, value="ИТОГО")
        ws.cell(row=last + 1, column=2, value=f"=SUM(B2:B{last})")
        for c in ws[1] + ws[last + 1]:
            c.font = Font(bold=True)
        for r in range(2, last + 1):
            ws.cell(row=r, column=1).number_format = "DD.MM.YYYY"
        for r in range(2, last + 2):
            ws.cell(row=r, column=2).number_format = "#,##0"
        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 14

        # ── вкладка «Сводная» ──
        w2 = xw.book.create_sheet("Сводная")
        w2.append(["Неделя", "Дата", "День", "Заказов", "Всего, шт", "Собрано, шт", "Осталось, шт", "% собрано"])
        hfill = PatternFill("solid", fgColor="DCE8F2")
        wfill = PatternFill("solid", fgColor="EBDDF5")
        for c in w2[1]:
            c.font, c.fill = Font(bold=True), hfill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        w2.row_dimensions[1].height = 30

        def put_formulas(i):
            w2.cell(row=i, column=7, value=f"=E{i}-F{i}")
            w2.cell(row=i, column=8, value=f"=IF(E{i}=0,0,F{i}/E{i})")

        sub_rows = []
        for wk, part in g.groupby("Нед"):
            wlabel = f"{wk:%d.%m}–{wk + timedelta(days=6):%d.%m}"
            first = w2.max_row + 1
            for r in part.itertuples(index=False):
                w2.append([wlabel, r.Дата, WEEKDAYS[r.Дата.weekday()], int(r.Заказов), int(r.Штук), int(r.Собрано)])
                i = w2.max_row
                put_formulas(i)
                w2.cell(row=i, column=2).number_format = "DD.MM.YYYY"
            lastr = w2.max_row
            w2.append([f"Итого {wlabel}", None, None,
                       f"=SUM(D{first}:D{lastr})", f"=SUM(E{first}:E{lastr})", f"=SUM(F{first}:F{lastr})"])
            i = w2.max_row
            put_formulas(i)
            sub_rows.append(i)
            for c in w2[i]:
                c.font, c.fill = Font(bold=True), wfill

        if sub_rows:
            w2.append(["ОБЩИЙ ИТОГ", None, None,
                       "=" + "+".join(f"D{x}" for x in sub_rows),
                       "=" + "+".join(f"E{x}" for x in sub_rows),
                       "=" + "+".join(f"F{x}" for x in sub_rows)])
            i = w2.max_row
            put_formulas(i)
            for c in w2[i]:
                c.font, c.fill = Font(bold=True), hfill

        for r in range(2, w2.max_row + 1):
            for col in (4, 5, 6, 7):
                w2.cell(row=r, column=col).number_format = "#,##0"
            w2.cell(row=r, column=8).number_format = "0%"
        for col, w in zip("ABCDEFGH", (22, 14, 8, 11, 13, 14, 14, 12)):
            w2.column_dimensions[col].width = w
        w2.freeze_panes = "A2"
    return out.getvalue()


def routing_excel(g: pd.DataFrame) -> bytes:
    """Сводная в том же виде, что и в Excel: Город / Название магазина / Адрес / Сумма штук / Общий итог."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Маршрутизация"
    head = ["Город", "Название магазина", "Адрес магазина", "Сумма по полю кол-во штук в заказе"]
    ws.append(head)
    fill = PatternFill("solid", fgColor="DCE8F2")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    prev_city = None
    for r in g.itertuples(index=False):
        ws.append([r.Город if r.Город != prev_city else None, r.Магазин, r.Адрес, int(r.Штук)])
        if r.Город != prev_city:
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
        prev_city = r.Город
    last = ws.max_row
    ws.append(["Общий итог", None, None, f"=SUM(D2:D{last})"])
    for c in ws[ws.max_row]:
        c.font, c.fill = Font(bold=True), fill
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=4).number_format = "#,##0"
    for col, w in zip("ABCD", (22, 26, 70, 22)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _st_ver() -> tuple:
    try:
        return tuple(int(x) for x in re.findall(r"\d+", st.__version__)[:2])
    except Exception:
        return (0, 0)


def dl_button(label: str, data: bytes, file_name: str, key: str):
    """Компактная кнопка скачивания уже готового файла (без перезапуска страницы при нажатии)."""
    kw = dict(label=label, data=data, file_name=file_name, mime=XLSX_MIME, key=key)
    if _st_ver() >= (1, 43):
        kw["on_click"] = "ignore"
    try:
        st.download_button(**kw)
    except TypeError:
        kw.pop("on_click", None)
        st.download_button(**kw)


def xl_button(label: str, build, file_name: str, key: str):
    """Компактная кнопка скачивания Excel. Файл собирается безопасно (ошибка не роняет страницу),
    имя файла только латиницей, без перезапуска страницы при нажатии."""
    try:
        data = build()
    except Exception as e:
        st.error(f"Не удалось собрать Excel-файл: {e}. Проверьте, что в requirements.txt есть openpyxl.")
        return
    dl_button(label, data, file_name, key)


# ───────────────────────── Логотипы ─────────────────────────
def _trim_logo(raw: bytes) -> bytes:
    """Обрезает пустые поля вокруг логотипа (прозрачные или белые), чтобы знак занимал всю высоту."""
    try:
        from PIL import Image, ImageChops
        im = Image.open(BytesIO(raw)).convert("RGBA")
        alpha = im.split()[3]
        bbox = alpha.getbbox() if alpha.getextrema()[0] < 255 else None   # есть прозрачность
        if bbox is None:                                                  # сплошной фон — ищем отличие от белого
            bg = Image.new("RGBA", im.size, (255, 255, 255, 255))
            diff = ImageChops.difference(im, bg).convert("L").point(lambda v: 255 if v > 12 else 0)
            bbox = diff.getbbox()
        if not bbox:
            return raw
        pad = 2
        box = (max(bbox[0] - pad, 0), max(bbox[1] - pad, 0),
               min(bbox[2] + pad, im.width), min(bbox[3] + pad, im.height))
        out = BytesIO()
        im.crop(box).save(out, "PNG")
        return out.getvalue()
    except Exception:
        return raw


@st.cache_data(show_spinner=False)
def _logo_data_uri(names: tuple = tuple(LOGO_FILES)) -> str:
    """Логотип из файла рядом с app.py -> data-URI. Нет файла — пустая строка."""
    bases = [Path(__file__).resolve().parent, Path.cwd()]
    for base in bases:
        for name in names:
            p = base / name
            if p.exists():
                ext = p.suffix.lstrip(".").lower()
                data = p.read_bytes()
                if ext == "svg":
                    return "data:image/svg+xml;base64," + base64.b64encode(data).decode()
                return "data:image/png;base64," + base64.b64encode(_trim_logo(data)).decode()
    return ""


def logo_html(h: int = None) -> str:
    """Логотипы M Cosmetic и партнёра по очереди плавно двигаются вверх по вертикальной ленте.
    Один логотип — мягко «парит» с неоновой подсветкой. h — высота окна, px.
    По клику на логотип (скрипт внизу страницы) открывается «Сводка по дням»."""
    h = int(h or LOGO_HEIGHT)
    uris = [u for u in (_logo_data_uri(tuple(LOGO_FILES)), _logo_data_uri(tuple(LOGO2_FILES))) if u]

    if not uris:    # файлов нет — нарисованная копия
        sz = int(h * 0.8)
        icon = (f"<svg width='{sz}' height='{sz}' viewBox='0 0 30 30' fill='none' xmlns='http://www.w3.org/2000/svg'>"
                "<rect x='2' y='2' width='26' height='26' rx='6' stroke='currentColor' stroke-width='3'/>"
                "<path d='M9 12h9l-3-3M21 18h-9l3 3' stroke='currentColor' stroke-width='3' "
                "stroke-linecap='round' stroke-linejoin='round'/></svg>")
        return (f"<div class='brand' title='На главную'><span class='bi'>{icon}</span>"
                f"<span class='bn' style='font-size:{h / 38:.2f}rem'>M COSMETIC</span></div>")

    def item(u):
        return (f"<div class='lt-item' style='height:{h}px'>"
                f"<img src='{u}' alt='logo' style='height:{h}px;width:auto;max-width:100%'></div>")

    if len(uris) == 1:
        return (f"<div class='brand' title='На главную'>"
                f"<div class='lt-win lt-one' style='height:{h}px'>{item(uris[0])}</div></div>")

    # несколько логотипов: лента из N штук + копия первого в конце, чтобы цикл был бесшовным
    k = len(uris)
    step = LOGO_HOLD + LOGO_SLIDE
    total = k * step
    frames = []
    for i in range(k):
        t0 = i * step / total * 100
        t1 = (i * step + LOGO_HOLD) / total * 100
        frames.append(f"{t0:.3f}%,{t1:.3f}%{{transform:translateY(-{i * h}px)}}")
    frames.append(f"100%{{transform:translateY(-{k * h}px)}}")
    css = ("<style>"
           f"@keyframes ltslide{h}{{" + "".join(frames) + "}"
           ".lt-win{overflow:hidden}"
           f".lt-track{h}{{animation:ltslide{h} {total:.2f}s cubic-bezier(.65,0,.35,1) infinite}}"
           ".lt-item{display:flex;align-items:center}"
           "</style>")
    items = "".join(item(u) for u in uris + [uris[0]])
    return (f"{css}<div class='brand' title='На главную'><div class='lt-win' style='height:{h}px'>"
            f"<div class='lt-track{h}'>{items}</div></div></div>")


# ───────────────────────── Мелкие виджеты ─────────────────────────
def tap(key: str, card_html: str, on_click, args=(), iframe=None):
    """Карточка, на которую можно нажать (поверх неё невидимая кнопка).
    iframe=(документ, высота) — карточка рисуется в отдельном блоке, где работают скрипты (таймер)."""
    with st.container(key=f"tap_{key}"):
        if iframe:
            components.html(iframe[0], height=iframe[1])
        else:
            st.markdown(card_html, unsafe_allow_html=True)
        st.button("открыть", key=f"btn_{key}", on_click=on_click, args=args)


def go_orders(mode: str, d):
    st.session_state.page = PAGES[1]
    st.session_state.f_mode = mode
    st.session_state.f_date = d


def set_oth(name):
    """«Прочие заказы»: открыть магазин (name) или вернуться к общему списку (None)."""
    st.session_state.oth_store = name


def set_page(p):
    st.session_state.page = p
    st.session_state.hub_open = False


def open_hub():
    """Карточка «Хаб ББ»: открыть подробности по маршрутам."""
    st.session_state.hub_open = True


def close_hub():
    st.session_state.hub_open = False


def toggle_theme():
    st.session_state.dark = not st.session_state.dark


def badge(collected: bool, new: bool = False) -> str:
    if new:
        return '<span class="badge no">не маршрутизирован</span>'
    return '<span class="badge ok">собран</span>' if collected else '<span class="badge no">не собран</span>'


def orders_table(frame: pd.DataFrame, limit: int = 1500, cp: bool = False) -> str:
    head = ("<tr><th>Дата</th><th>Заказ</th><th>Магазин</th><th>Город</th>"
            + ("<th>Категория</th><th>Приоритет</th>" if cp else "")
            + "<th class='num'>Штук</th><th>Статус</th><th>WMS</th></tr>")
    body = []
    for r in frame.head(limit).to_dict("records"):
        extra = ""
        if cp:
            pr = r.get("Приоритет") or ""
            pr_html = f"<span class='badge no'>{esc(pr)}</span>" if pr else "—"
            extra = f"<td>{esc(r.get('Категория') or '—')}</td><td>{pr_html}</td>"
        body.append(
            f"<tr><td>{fmt_date(r['Дата'])}</td><td><b>{esc(r['Заказ'])}</b></td><td>{esc(r['Магазин'])}</td>"
            f"<td>{esc(r['Город'] or '—')}</td>{extra}<td class='num'>{n(r['Штук'])}</td>"
            f"<td>{badge(r['Собран'], bool(r.get('Новый', False)))}</td>"
            f"<td>{esc(r['Статус WMS'] or '—')}</td></tr>")
    return f"<div class='tbl-wrap'><table class='tbl'>{head}{''.join(body)}</table></div>"


def routing_table(g: pd.DataFrame) -> str:
    head = ("<tr><th>Город</th><th>Название магазина</th><th>Адрес магазина</th>"
            "<th class='num'>Сумма по полю кол-во штук в заказе</th></tr>")
    body, prev = [], None
    for r in g.itertuples(index=False):
        city = (esc(r.Город) or "—") if r.Город != prev else ""
        body.append(f"<tr><td class='city'>{city}</td><td class='shop'>{esc(r.Магазин)}</td>"
                    f"<td>{esc(r.Адрес)}</td><td class='num'>{n(r.Штук)}</td></tr>")
        prev = r.Город
    body.append(f"<tr class='total'><td colspan='3'>Общий итог</td><td class='num'>{n(g['Штук'].sum())}</td></tr>")
    return f"<div class='tbl-wrap'><table class='tbl'>{head}{''.join(body)}</table></div>"


# ── карточка «Отгрузка сегодня» с живым таймером и едущей «Газелью» (отдельный документ со своими стилями и скриптом) ──
SHIP_DOC = Template("""<!doctype html><html><head><meta charset="utf-8">
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
html, body { margin:0; padding:0; background:transparent; font-family:'Inter', sans-serif; color:$ink; }
.ship { background:linear-gradient(135deg, $head, $card); border:1px solid $line; border-radius:16px;
        box-shadow:0 0 0 1px rgba(139,61,255,.15), 0 10px 30px $glow;
        padding:10px 14px; margin:0 0 8px; box-sizing:border-box; }
.sh-top { display:flex; justify-content:space-between; align-items:baseline; gap:10px; }
.sh-t { font-weight:800; font-size:.98rem; color:$ink; display:flex; align-items:center; gap:8px; }
.pdot { width:8px; height:8px; border-radius:50%; background:$g1; box-shadow:0 0 8px $g1; }
.sh-t small { font-weight:500; color:$mute; font-size:.76rem; margin-left:6px; }
.sh-p { font-weight:800; font-size:1.4rem; letter-spacing:-0.02em; color:$ink; }
.sync { display:inline-flex; width:18px; height:18px; border-radius:50%; align-items:center; justify-content:center;
        margin-left:8px; vertical-align:middle; font-size:.7rem; font-weight:800; box-sizing:border-box; }
.sync.ok { background:#1E9E5C; color:#fff; }
.sync.spin { border:3px solid $track; border-top-color:$left; animation:syncspin .9s linear infinite; }
@keyframes syncspin { to { transform:rotate(360deg); } }
.bar { height:26px; border-radius:99px; background:$track; overflow:hidden; position:relative; margin:0; }
.bar > i { display:block; height:100%; border-radius:99px; background-color:#1E9E5C;
  background-image:linear-gradient(135deg, rgba(255,255,255,.22) 25%, transparent 25%, transparent 50%,
                                   rgba(255,255,255,.22) 50%, rgba(255,255,255,.22) 75%, transparent 75%),
                   linear-gradient(90deg, #1E9E5C, #1E9E5C);
  background-size:32px 32px, 100% 100%; animation:move 1s linear infinite; transition:background .6s; }
.bar.g > i { background-image:linear-gradient(135deg, rgba(255,255,255,.22) 25%, transparent 25%, transparent 50%,
                                   rgba(255,255,255,.22) 50%, rgba(255,255,255,.22) 75%, transparent 75%),
                   linear-gradient(90deg, $g1, $g2, #3B82F6); box-shadow:0 0 18px $glow; }
.bar.y > i { background-image:linear-gradient(135deg, rgba(255,255,255,.22) 25%, transparent 25%, transparent 50%,
                                   rgba(255,255,255,.22) 50%, rgba(255,255,255,.22) 75%, transparent 75%),
                   linear-gradient(90deg, #F2B705, #F2B705); }
.bar.r > i { background-image:linear-gradient(135deg, rgba(255,255,255,.22) 25%, transparent 25%, transparent 50%,
                                   rgba(255,255,255,.22) 50%, rgba(255,255,255,.22) 75%, transparent 75%),
                   linear-gradient(90deg, #D32F2F, #D32F2F); }
@keyframes move { from { background-position:0 0, 0 0; } to { background-position:32px 0, 0 0; } }
.bar-t { position:absolute; left:0; right:0; top:0; bottom:0; display:flex; align-items:center; justify-content:center;
  gap:10px; color:#fff; font-weight:800; font-size:.92rem; font-variant-numeric:tabular-nums; letter-spacing:.02em;
  text-shadow:0 1px 3px rgba(0,0,0,.6); }
.bar-t small { font-size:.74rem; font-weight:700; }
.sh-m { font-size:.78rem; color:$mute; margin-top:7px; }
.sh-m b { color:$ink; }
$truck_css
</style></head><body>
<div class="ship">
  <div class="sh-top"><div class="sh-t">Отгрузка сегодня <span class="pdot"></span><small>$date</small>$clock$sync</div><div class="sh-p">$pct%</div></div>
  <div class="lane">$truck<div class="bar g" id="dl-bar"><i style="width:$pct%"></i><span class="bar-t" id="dl-time">--:--:--</span></div></div>
  <div class="sh-m">Отгружено <b>$done</b> из $total заказов · <b>$pcs_done</b> из $pcs шт · осталось <b>$left</b> зак. ($pcs_left шт)</div>
</div>
<script>
(function () {
  var DONE = $all_done, TARGET = 17 * 3600;
  var fmt = new Intl.DateTimeFormat('en-GB', {timeZone: 'Asia/Tashkent', hour: '2-digit', minute: '2-digit',
                                               second: '2-digit', hour12: false});
  function pad(x) { return (x < 10 ? '0' : '') + x; }
  function nowSec() {
    var v = {}; fmt.formatToParts(new Date()).forEach(function (x) { v[x.type] = +x.value; });
    return (v.hour % 24) * 3600 + v.minute * 60 + v.second;
  }
  function tick() {
    var bar = document.getElementById('dl-bar'), box = document.getElementById('dl-time');
    var s = nowSec(), left = TARGET - s, cls, h;
    if (DONE) { cls = 'g'; h = '✓ Все заказы отгружены'; }
    else {
      cls = s < 12 * 3600 ? 'g' : (s < 16 * 3600 ? 'y' : 'r');
      var a = Math.abs(left);
      var t = (left < 0 ? '+' : '') + pad(Math.floor(a / 3600)) + ':' + pad(Math.floor(a % 3600 / 60)) + ':' + pad(a % 60);
      h = '<small>' + (left >= 0 ? 'До 17:00 осталось' : 'Срок 17:00 прошёл') + '</small>' + t;
    }
    bar.className = 'bar ' + cls;
    box.innerHTML = h;
  }
  tick();
  setInterval(tick, 1000);
})();
</script></body></html>""")

SHIP_IFRAME_H = 140   # высота карточки с таймером и машинкой, px


def ship_doc(d: date, pct: int, sm: dict, sync: str = "") -> str:
    done_all = sm["total"] > 0 and sm["done"] >= sm["total"]
    vals = dict(P)
    vals.update(date=f"{fmt_date(d)}, {WEEKDAYS[d.weekday()]}", pct=pct, sync=sync, clock=ic("clock", P["mute"], 15),
                done=n(sm["done"]), total=n(sm["total"]), pcs_done=n(sm["pcs_done"]), pcs=n(sm["pcs"]),
                left=n(sm["total"] - sm["done"]), pcs_left=n(sm["pcs"] - sm["pcs_done"]),
                all_done="true" if sm["done"] >= sm["total"] else "false",
                truck=truck_html(pct, done_all), truck_css=TRUCK_CSS)
    return SHIP_DOC.safe_substitute(vals)


# ── пустые короба ──
def month_label(key: str) -> str:
    y, m = key.split("-")
    return f"{MONTHS[int(m) - 1].capitalize()} {y}"


def simple_table(head: list, rows: list, num_cols=()) -> str:
    th = "".join(f"<th class='{'num' if i in num_cols else ''}'>{esc(h)}</th>" for i, h in enumerate(head))
    body = []
    for r in rows:
        tds = "".join(f"<td class='{'num' if i in num_cols else ''}'>{c}</td>" for i, c in enumerate(r))
        body.append(f"<tr>{tds}</tr>")
    return f"<div class='tbl-wrap'><table class='tbl'><tr>{th}</tr>{''.join(body)}</table></div>"


def kpi_box(value, label, sub="", cls="ok", spark=None, delta=None, icon=None, status=None) -> str:
    """Карточка: иконка-плитка, число, подпись; справа — процент и мини-график (если переданы);
    status — [(название, штук, цвет)] мелким шрифтом внизу карточки."""
    icon = icon or KPI_ICON.get(cls, "box")
    sub_html = f"<div class='s'>{sub}</div>" if sub else ""
    trend = ""
    if spark or delta is not None:
        dh = ""
        if delta is not None:
            dh = f"<span class='delta {'up' if delta >= 0 else 'dn'}'>{'↑' if delta >= 0 else '↓'} {delta:+.0f}%</span>"
        sp = spark_svg(spark, SPARK_COL.get(cls, "#E0338F"), zlib.crc32(str(label).encode()) % 10 ** 6) if spark else ""
        trend = f"<div class='trend'>{dh}{sp}</div>"
    sts = ""
    if status:
        if value is None:
            items = "".join(f"<span class='stp'><span class='nm'><i style='background:{c}'></i>{esc(nm)}</span>"
                            f"<b>{n(v)}</b></span>" for nm, v, c in status)
            sts = f"<div class='sts'>{items}</div>"
        else:
            items = "".join(f"<span class='stp'><i style='background:{c}'></i>{esc(nm)} <b>{n(v)}</b></span>"
                            for nm, v, c in status)
            sts = f"<div class='sts'><span class='stt'>шт по статусам</span>{items}</div>"
    v_html = "" if value is None else f"<div class='v'>{value}</div>"
    return (f"<div class='kpi {cls}{' nov' if value is None else ''}'><div class='ico'>{ic(icon, '#FFFFFF', 20)}</div>"
            f"<div class='mid'>{v_html}<div class='l'>{esc(label)}</div>{sub_html}</div>{trend}{sts}</div>")


# ───────────────────────── Ручные рейсы: хранение ─────────────────────────
# plan = {"2026-10-15": {"1": ["Магазин А", "Магазин Б"], "2": [...]}, ...}
def _plan_rows(plan: dict) -> list:
    rows = [["Дата", "Машина", "Магазин"]]
    for dk in sorted(plan):
        for k in sorted(plan[dk], key=int):
            shops = plan[dk][k]
            if shops:
                rows += [[dk, str(k), s] for s in shops]
            else:
                rows.append([dk, str(k), ""])
    return rows


def _plan_from_rows(rows: list) -> dict:
    plan = {}
    for r in rows[1:]:
        if len(r) < 2 or not r[0].strip() or not r[1].strip().isdigit():
            continue
        dk, k = r[0].strip(), str(int(r[1].strip()))
        s = r[2].strip() if len(r) > 2 else ""
        lst = plan.setdefault(dk, {}).setdefault(k, [])
        if s and s not in lst:
            lst.append(s)
    return plan


def plan_load():
    """Возвращает (plan, где_хранится): 'table' — вкладка в таблице, 'file' — файл на сервере, '' — пока нигде."""
    try:
        sh = _client_rw().open_by_key(SHEET_ID)
        ws = next((w for w in sh.worksheets() if w.title == PLAN_TAB), None)
        if ws is not None:
            return _plan_from_rows(ws.get_all_values()), "table"
    except Exception:
        pass
    try:
        with open(PLAN_FILE, encoding="utf-8") as f:
            return json.load(f), "file"
    except Exception:
        return {}, ""


def plan_save(dk: str):
    """Сохраняет рейсы выбранной даты (остальные даты подтягиваются из хранилища, чтобы не затереть чужие правки)."""
    mine = st.session_state.plan.get(dk, {})
    remote, _ = plan_load()
    remote[dk] = mine
    st.session_state.plan = remote
    try:
        sh = _client_rw().open_by_key(SHEET_ID)
        ws = next((w for w in sh.worksheets() if w.title == PLAN_TAB), None)
        if ws is None:
            ws = sh.add_worksheet(PLAN_TAB, rows=1000, cols=3)
        rows = _plan_rows(remote)
        ws.clear()
        ws.update(values=rows, range_name="A1", value_input_option="RAW")
        st.session_state.plan_backend = "table"
        return
    except Exception:
        pass
    try:
        with open(PLAN_FILE, "w", encoding="utf-8") as f:
            json.dump(remote, f, ensure_ascii=False)
        st.session_state.plan_backend = "file"
    except Exception:
        st.session_state.plan_backend = "none"


def ensure_plan():
    if "plan" not in st.session_state:
        st.session_state.plan, st.session_state.plan_backend = plan_load()


def stores_of(pivot: pd.DataFrame) -> dict:
    """Магазины выбранной даты: {название: {city, addr, qty}} (одинаковые названия суммируются)."""
    out = {}
    for r in pivot.itertuples(index=False):
        if r.Магазин in out:
            out[r.Магазин]["qty"] += int(r.Штук)
        else:
            out[r.Магазин] = dict(city=r.Город, addr=r.Адрес, qty=int(r.Штук))
    return out


def _on_stores(dk, k, key):
    st.session_state.plan.setdefault(dk, {})[k] = list(st.session_state[key])
    plan_save(dk)


def _on_add_machine(dk):
    p = st.session_state.plan.setdefault(dk, {"1": []})
    nk = str(max(int(x) for x in p) + 1)
    p[nk] = []
    st.session_state[f"sel_{dk}"] = nk
    plan_save(dk)


def _on_del_machine(dk):
    p = st.session_state.plan.setdefault(dk, {"1": []})
    if len(p) > 1:
        p.pop(str(max(int(x) for x in p)), None)
        st.session_state[f"sel_{dk}"] = str(max(int(x) for x in p))
        plan_save(dk)


def _on_limit():
    st.session_state.lim_val = int(st.session_state.get("trip_limit_w") or 0)


def plan_machines(dk: str, stores: dict, limit: int) -> list:
    """Готовые рейсы в том же формате, что использует таблица/Excel/кнопки карт."""
    out = []
    mach = st.session_state.plan.get(dk, {})
    for k in sorted(mach, key=int):
        lst = [s for s in mach[k] if s in stores]
        if not lst:
            continue
        items = []
        for s in lst:
            dl = deadline_minutes(stores[s]["addr"])
            items.append(dict(shop=s, addr=stores[s]["addr"], qty=stores[s]["qty"],
                              hot=dl is not None and dl < 12 * 60))
        q = sum(x["qty"] for x in items)
        out.append(dict(n=int(k), stores=items, qty=q, target=limit, over=bool(limit) and q > limit))
    return out


def plan_ui(dk: str, stores: dict):
    """Ручная раскладка: выбираем машину, добавляем в неё магазины, смотрим счётчик."""
    ensure_plan()
    plan = st.session_state.plan
    mach = plan.setdefault(dk, {"1": []})
    if not mach:
        mach["1"] = []
    ks = sorted(mach, key=int)

    st.markdown("<div class='sec'>Рейсы: собрать вручную</div>", unsafe_allow_html=True)
    sel_key = f"sel_{dk}"
    if st.session_state.get(sel_key) not in ks:
        st.session_state[sel_key] = ks[0]

    c1, c2, c3 = st.columns([3, 1.3, 1.6], vertical_alignment="bottom")
    with c1:
        k = st.selectbox("Машина", ks, key=sel_key, format_func=lambda x: ordinal(int(x)))
    with c2:
        st.button("➕ Машина", on_click=_on_add_machine, args=(dk,), width="stretch", key=f"add_{dk}")
    with c3:
        st.button("🗑 Удалить последнюю", on_click=_on_del_machine, args=(dk,), width="stretch",
                  key=f"del_{dk}", disabled=len(ks) < 2)

    st.number_input("Лимит штук на одну машину (0 — без лимита)", min_value=0, step=100,
                    value=int(st.session_state.get("lim_val", 0)), key="trip_limit_w", on_change=_on_limit)
    limit = int(st.session_state.get("lim_val", 0))

    other = {s for kk, v in mach.items() if kk != k for s in v}
    options = [s for s in stores if s not in other]
    ms_key = f"ms_{dk}_{k}"
    st.session_state[ms_key] = [s for s in mach[k] if s in options]
    st.multiselect(f"Магазины в машине: {ordinal(int(k)).lower()} (начните вводить название)", options,
                   key=ms_key, on_change=_on_stores, args=(dk, k, ms_key),
                   format_func=lambda s: f"{s} · {n(stores[s]['qty'])} шт")

    # ── счётчик ──
    cur = [s for s in mach[k] if s in stores]
    cur_q = sum(stores[s]["qty"] for s in cur)
    assigned = {s for v in mach.values() for s in v if s in stores}
    total_q = sum(v["qty"] for v in stores.values())
    assigned_q = sum(stores[s]["qty"] for s in assigned)
    left = [s for s in stores if s not in assigned]
    left_q = total_q - assigned_q

    sub = f"точек: {len(cur)}"
    if limit:
        rem = limit - cur_q
        sub += f" · до лимита {n(rem)} шт" if rem >= 0 else f" · выше лимита на {n(-rem)} шт"
    k1, k2, k3 = st.columns(3)
    with k1:
        st.markdown(kpi_box(n(cur_q), f"в машине: {ordinal(int(k)).lower()}", sub,
                            "left" if limit and cur_q > limit else "ok"), unsafe_allow_html=True)
    with k2:
        st.markdown(kpi_box(n(assigned_q), "распределено всего", f"из {n(total_q)} шт · точек: {len(assigned)}", "all"),
                    unsafe_allow_html=True)
    with k3:
        st.markdown(kpi_box(n(left_q), "осталось распределить", f"точек: {len(left)}", "left" if left else "ok"),
                    unsafe_allow_html=True)

    if left:
        st.caption("Не распределено: " + ", ".join(f"{s} ({n(stores[s]['qty'])})" for s in left))
    else:
        st.success("Все точки распределены по машинам. Готовый результат смотрите в «Рейсы».")

    be = st.session_state.get("plan_backend", "")
    if be == "table":
        st.caption(f"Изменения сохраняются автоматически (вкладка «{PLAN_TAB}» в таблице). Готовый файл — в разделе «Рейсы».")
    elif be == "file":
        st.caption("Изменения сохраняются в файл на сервере: он может пропасть при перезапуске приложения. "
                   "Дайте сервисному аккаунту доступ «Редактор» к таблице, чтобы рейсы хранились в ней.")
    elif be == "none":
        st.warning("Не удалось сохранить рейсы: дайте сервисному аккаунту доступ «Редактор» к таблице.")
    else:
        st.caption("Изменения сохраняются автоматически. Готовый файл — в разделе «Рейсы».")


# ───────────────────────── Загрузка и шапка ─────────────────────────
try:
    df = load()
except Exception as e:
    st.error("Не удалось прочитать таблицу. Проверьте, что таблица открыта для почты "
             "сервисного аккаунта (роль «Читатель») и ключ добавлен в Secrets.")
    st.exception(e)
    st.stop()

now = datetime.now(TZ)

# верхняя панель: сколько человек онлайн + дата
ONLINE_WINDOW = 40   # сколько секунд после последнего сигнала человек считается «онлайн»
BEAT_EVERY = 15      # как часто вкладка подаёт сигнал, сек


@st.cache_resource
def _presence2() -> dict:
    return {}        # общий для всех сессий: {id сессии: {"first": когда зашёл, "last": последний сигнал, "vid": посетитель}}


# ── посетители: кто и когда заходил (для окна по клику на «онлайн») ──
def parse_ua(ua: str) -> dict:
    """Браузер, ОС и тип устройства по User-Agent."""
    low = (ua or "").lower()
    if "telegram" in low:
        br = "Telegram"
    elif "edg/" in low or "edga/" in low or "edgios/" in low:
        br = "Edge"
    elif "opr/" in low or "opera" in low:
        br = "Opera"
    elif "yabrowser" in low:
        br = "Яндекс"
    elif "samsungbrowser" in low:
        br = "Samsung"
    elif "firefox" in low or "fxios" in low:
        br = "Firefox"
    elif "chrome" in low or "crios" in low:
        br = "Chrome"
    elif "safari" in low:
        br = "Safari"
    else:
        br = "Браузер"
    if "android" in low:
        os_ = "Android"
    elif "iphone" in low or "ipad" in low or "ipod" in low:
        os_ = "iOS"
    elif "windows" in low:
        os_ = "Windows"
    elif "mac os x" in low or "macintosh" in low:
        os_ = "macOS"
    elif "cros" in low:
        os_ = "ChromeOS"
    elif "linux" in low:
        os_ = "Linux"
    else:
        os_ = "ОС"
    if "ipad" in low or "tablet" in low or ("android" in low and "mobile" not in low):
        dev = "tablet"
    elif "mobi" in low or "iphone" in low:
        dev = "mobile"
    else:
        dev = "desktop"
    return {"browser": br, "os": os_, "device": dev}


DEVICE_RU = {"mobile": "Телефон", "tablet": "Планшет", "desktop": "Компьютер"}


def ensure_client():
    """Один раз за сессию: id посетителя (cookie браузера), имя из Telegram (cookie), браузер/ОС из заголовков."""
    if "cinfo" in st.session_state:
        return
    ua, tz, lang, ck = "", "", "", {}
    try:
        hdr = st.context.headers
        ua = hdr.get("User-Agent", "") or ""
        lang = (hdr.get("Accept-Language", "") or "").split(",")[0][:8]
    except Exception:
        pass
    try:
        ck = st.context.cookies
    except Exception:
        ck = {}
    try:
        tz = str(getattr(st.context, "timezone", "") or "")
    except Exception:
        tz = ""
    vid = str(ck.get("mc_vid", "") or "").strip()
    if not re.fullmatch(r"[0-9a-f]{16}", vid):
        vid = uuid.uuid4().hex[:16]
    name = unquote(str(ck.get("mc_tg", "") or ""))[:60]
    st.session_state.vid = vid
    st.session_state.cinfo = dict(ua=parse_ua(ua), tz=tz, lang=lang, name=name)


@st.cache_resource
def _vis() -> dict:
    """Общая для всех сессий история посетителей: {id посетителя: имя, последний визит, сессии [начало, конец, id сессии]}."""
    h = {"data": {}, "lock": threading.Lock(), "saved": 0.0, "dirty": False}
    try:
        with open(VISITORS_FILE, encoding="utf-8") as f:
            h["data"] = json.load(f).get("data", {})
    except Exception:
        pass
    return h


def _vis_save(h: dict, t: float):
    cut = t - VIS_KEEP_DAYS * 86400
    for k in [k for k, v in h["data"].items() if v.get("last", 0) < cut]:
        h["data"].pop(k, None)
    for v in h["data"].values():
        v["sessions"] = [s for s in v.get("sessions", []) if s[1] >= cut]
    try:
        with open(VISITORS_FILE, "w", encoding="utf-8") as f:
            json.dump({"data": h["data"]}, f, ensure_ascii=False)
    except Exception:
        pass
    h["saved"], h["dirty"] = t, False


def _vis_touch(sid: str, vid: str, t: float):
    """Отмечает сигнал от открытой вкладки: продлевает текущую сессию посетителя или начинает новую."""
    ci = st.session_state.get("cinfo") or {}
    h, s8 = _vis(), sid[:8]
    with h["lock"]:
        v = h["data"].get(vid)
        if v is None:
            v = h["data"][vid] = {"name": "", "first": t, "last": t, "sessions": []}
        v["last"] = t
        if ci.get("name"):
            v["name"] = ci["name"]
        for k in ("ua", "tz", "lang"):
            if ci.get(k):
                v[k] = ci[k]
        ss = v["sessions"]
        for s in reversed(ss):
            if s[2] == s8:
                s[1] = t
                break
        else:
            ss.append([t, t, s8])
            del ss[:-200]
            h["dirty"] = True
        if h["dirty"] or t - h["saved"] > 120:
            _vis_save(h, t)


def online_now() -> list:
    """Отмечает текущего человека и возвращает список [(время входа, id сессии)] по порядку входа."""
    if "sid" not in st.session_state:
        st.session_state.sid = uuid.uuid4().hex
    reg, t = _presence2(), time.time()
    sid = st.session_state.sid
    vid = st.session_state.get("vid") or sid[:16]
    rec = reg.get(sid)
    if rec is None:
        reg[sid] = {"first": t, "last": t, "vid": vid}
    else:
        rec["last"] = t
    _vis_touch(sid, vid, t)
    for k in [k for k, v in reg.items() if t - v["last"] > ONLINE_WINDOW]:
        reg.pop(k, None)
    return sorted((v["first"], k) for k, v in reg.items())


def dur_txt(sec: float) -> str:
    s = max(int(sec), 0)
    if s < 60:
        return f"{s} сек"
    m = s // 60
    if m < 60:
        return f"{m} мин"
    return f"{m // 60} ч {m % 60} мин"


def _plural(k, forms) -> str:
    k = abs(int(k))
    if 11 <= k % 100 <= 14:
        return forms[2]
    r = k % 10
    return forms[0] if r == 1 else (forms[1] if 2 <= r <= 4 else forms[2])


def ago_txt(sec: float) -> str:
    s = max(int(sec), 0)
    if s < 60:
        return f"{s} {_plural(s, ('секунду', 'секунды', 'секунд'))} назад"
    m = s // 60
    if m < 60:
        return f"{m} {_plural(m, ('минуту', 'минуты', 'минут'))} назад"
    hh = m // 60
    if hh < 48:
        return f"{hh} {_plural(hh, ('час', 'часа', 'часов'))} назад"
    dd = hh // 24
    return f"{dd} {_plural(dd, ('день', 'дня', 'дней'))} назад"


def clock_txt(sec: float) -> str:
    s = max(int(sec), 0)
    if s >= 3600:
        return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}"
    return f"{s // 60}:{s % 60:02d}"


def _avatar(vid: str, name: str) -> str:
    """Аватар посетителя: с именем — цветной кружок с инициалами, гость — «узор» из точек (по id посетителя)."""
    dg = hashlib.md5(vid.encode()).digest()
    hue = dg[0] * 360 // 256
    if name:
        ini = "".join(w[0] for w in name.replace("@", "").split()[:2]).upper() or "?"
        return (f"<div class='vav' style='background:linear-gradient(135deg,hsl({hue},70%,52%),"
                f"hsl({(hue + 50) % 360},70%,45%))'>{esc(ini)}</div>")
    col = f"hsl({hue},62%,58%)"
    dots = "".join(f"<circle cx='{8 + c * 5}' cy='{8 + r * 5}' r='2' fill='{col}'/>"
                   for r in range(5) for c in range(5) if dg[1 + r * 3 + min(c, 4 - c)] & 1)
    return f"<div class='vav' style='background:#fff'><svg width='34' height='34' viewBox='0 0 36 36'>{dots}</svg></div>"


VIS_PERIODS = ["Сейчас", "Сегодня", "7 дней", "30 дней"]


def _online_body():
    """Содержимое окна «Посетители»: карточки (посетители, среднее время сессии, сессий на посетителя) и таблица."""
    st.session_state.setdefault("vis_per", "7 дней")
    st.segmented_control("Период", VIS_PERIODS, key="vis_per", label_visibility="collapsed")
    per = st.session_state.get("vis_per") or "7 дней"

    t_now = time.time()
    hv, reg = _vis(), _presence2()
    live = {k: r for k, r in list(reg.items()) if t_now - r["last"] <= ONLINE_WINDOW}
    live_vids = {(r.get("vid") or k) for k, r in live.items()}
    live_s8 = {k[:8] for k in live}
    midnight = datetime.now(TZ).replace(hour=0, minute=0, second=0, microsecond=0).timestamp()
    since = {"Сегодня": midnight, "7 дней": t_now - 7 * 86400, "30 дней": t_now - 30 * 86400}.get(per, 0)
    with hv["lock"]:
        data = {k: dict(v, sessions=[list(s) for s in v.get("sessions", [])]) for k, v in hv["data"].items()}

    me = st.session_state.get("vid")
    rows, durs, total_s = [], [], 0
    for vid, v in data.items():
        if per == "Сейчас":
            ss = [s for s in v["sessions"] if s[2] in live_s8]
        else:
            ss = [s for s in v["sessions"] if s[1] >= since]
        if not ss:
            continue
        total_s += len(ss)
        durs += [max(s[1] - s[0], 0) for s in ss]
        rows.append(dict(vid=vid, name=v.get("name", ""), sess=len(ss), last=v.get("last", ss[-1][1]),
                         first=v.get("first", ss[0][0]), dur=max(ss[-1][1] - ss[-1][0], 0),
                         online=vid in live_vids, ua=v.get("ua") or {}, tz=v.get("tz", ""), lang=v.get("lang", "")))
    rows.sort(key=lambda r: -r["last"])

    if not rows:
        st.info("За выбранный период посетителей нет.")
    else:
        avg_t = sum(durs) / len(durs) if durs else 0
        avg_s = total_s / len(rows)
        avg_s_txt = f"{avg_s:.0f}" if abs(avg_s - round(avg_s)) < 0.05 else f"{avg_s:.1f}".replace(".", ",")

        def card(label, value, icon):
            return (f"<div class='vk'><div><div class='vkl'>{label}</div><div class='vkv'>{value}</div></div>"
                    f"<span class='vki'>{ic(icon, 'currentColor', 17)}</span></div>")

        out = ["<div class='vk-wrap'>",
               card("Посетители", n(len(rows)), "users"),
               card("Ср. время сессии", f"{clock_txt(avg_t)} мин", "timer"),
               card("Ср. сессий на посетителя", avg_s_txt, "hourglass"),
               "</div><div class='vtbl'>",
               "<div class='vrow vh'><div>Посетитель</div><div></div><div>Последняя активность</div><div>Детали</div></div>"]
        for r in rows[:100]:
            name = r["name"] or "Гость"
            you = "<span class='vme'>(вы)</span>" if r["vid"] == me else ""
            sess = (f"<span class='vpill s'>{ic('hourglass', 'currentColor', 12)} {r['sess']} "
                    f"{_plural(r['sess'], ('сессия', 'сессии', 'сессий'))}</span>")
            if r["online"]:
                last = "<span class='vpill on'>● онлайн</span>"
            else:
                last = f"<span class='vpill t'>{ic('history', 'currentColor', 12)} {ago_txt(t_now - r['last'])}</span>"
            ua = r["ua"]
            dev = ua.get("device", "desktop")
            tz_txt = " · ".join(x for x in (r["tz"], r["lang"]) if x) or "нет данных"
            det = (f"<span class='vd' title='Первый визит: {datetime.fromtimestamp(r['first'], TZ):%d.%m.%Y %H:%M}'>"
                   f"{ic('calendar', 'currentColor', 15)}</span>"
                   f"<span class='vd' title='Длительность последней сессии: {clock_txt(r['dur'])}'>"
                   f"{ic('history', 'currentColor', 15)}</span>"
                   f"<span class='vd' title='Часовой пояс и язык: {esc(tz_txt)}'>{ic('globe', 'currentColor', 15)}</span>"
                   f"<span class='vd' title='{DEVICE_RU.get(dev, '')}'>"
                   f"{ic('monitor' if dev == 'desktop' else 'smartphone', 'currentColor', 15)}</span>"
                   f"<span class='vtx'>{esc(ua.get('os', ''))} · {esc(ua.get('browser', ''))}</span>")
            out.append(f"<div class='vrow'><div class='vuser'>{_avatar(r['vid'], r['name'])}"
                       f"<span class='vname'>{esc(name)}{you}</span></div><div>{sess}</div><div>{last}</div>"
                       f"<div class='vdet'>{det}</div></div>")
        out.append("</div>")
        st.markdown("".join(out), unsafe_allow_html=True)

    st.caption("Посетитель определяется по cookie браузера, сессия — одно открытие дашборда. Имя берётся из Telegram "
               "(если дашборд открыт из бота) и появляется со следующего открытия. Время сессии считается по сигналам "
               f"открытой вкладки (раз в {BEAT_EVERY} сек). История хранится {VIS_KEEP_DAYS} дней в файле на сервере.")
    st.button("Обновить", key="vis_refresh")


_dlg = getattr(st, "dialog", None) or getattr(st, "experimental_dialog", None)


def online_dialog():
    """Открывает окно «Посетители» (вызывается после клика на «N онлайн»)."""
    if _dlg is None:
        st.info("Окно посетителей работает в Streamlit 1.37 и новее.")
        return
    _dlg("Посетители", width="large")(_online_body)()


# ── история изменений (для колокольчика): что добавилось, собрано, отгружено ──
EV_META = {"added": ("добавлено заказов", "#E0338F"), "routed": ("маршрутизировано", "#8B3DFF"),
           "collected": ("собрано", "#34C27A"), "shipped": ("отгружено", "#3B9CFF"),
           "removed": ("убрано из таблицы", "#8C93B5")}


@st.cache_resource
def _hist() -> dict:
    """Общая для всех сессий память: последний снимок заказов и список событий."""
    h = {"snap": None, "events": [], "since": None, "cur": {}, "lock": threading.Lock()}
    try:
        with open(HISTORY_FILE, encoding="utf-8") as f:
            d = json.load(f)
        h["events"], h["since"] = d.get("events", []), d.get("since")
        s = d.get("snap")
        if s:
            h["snap"] = dict(pcs={k: int(v) for k, v in s["pcs"].items()}, orders=set(s["orders"]),
                             new=set(s["new"]), collected=set(s["collected"]), shipped=set(s["shipped"]))
    except Exception:
        pass
    return h


def _hist_save(h: dict):
    try:
        s = h["snap"]
        with open(HISTORY_FILE, "w", encoding="utf-8") as f:
            json.dump({"since": h["since"], "events": h["events"][-HISTORY_KEEP:],
                       "snap": {"pcs": s["pcs"], "orders": sorted(s["orders"]), "new": sorted(s["new"]),
                                "collected": sorted(s["collected"]), "shipped": sorted(s["shipped"])}},
                      f, ensure_ascii=False)
    except Exception:
        pass


def record_history(df: pd.DataFrame, ship, new_orders):
    """Сравнивает текущие данные с прошлым снимком и дописывает события.
    ship / new_orders = None — источник сейчас недоступен: берём их из прошлого снимка, чтобы не плодить ложные события."""
    h, now_ = _hist(), datetime.now(TZ)
    pcs = {}
    if ship is not None and not ship.empty:
        pcs.update({k: int(v) for k, v in zip(ship["Заказ"], ship["Штук"])})
    if new_orders is not None:
        pcs.update({k: int(v) for k, v in zip(new_orders["Заказ"], new_orders["Штук"])})
    pcs.update({k: int(v) for k, v in zip(df["Заказ"], df["Штук"])})
    shipped_today, shipped_today_pcs = 0, 0
    if ship is not None and not ship.empty:
        try:
            m = ship["Отгружен"] & ship["Факт"].notna() & (ship["Факт"].dt.date == now_.date())
            shipped_today, shipped_today_pcs = int(m.sum()), int(ship.loc[m, "Штук"].sum())
        except Exception:
            pass
    with h["lock"]:
        prev = h["snap"]
        cur = dict(pcs=pcs, orders=set(df["Заказ"]),
                   new=set(new_orders["Заказ"]) if new_orders is not None else (prev["new"] if prev else set()),
                   collected=set(df.loc[df["Собран"], "Заказ"]),
                   shipped=(set(ship.loc[ship["Отгружен"], "Заказ"]) if ship is not None and not ship.empty
                            else (prev["shipped"] if prev else set())))
        new_cnt = len(new_orders) if new_orders is not None else len(cur["new"])
        h["cur"] = dict(total=len(df), left=int((~df["Собран"]).sum()) + new_cnt,
                        pcs_left=int(df.loc[~df["Собран"], "Штук"].sum()) + (
                            int(new_orders["Штук"].sum()) if new_orders is not None else 0),
                        new_cnt=new_cnt, shipped_today=shipped_today, shipped_today_pcs=shipped_today_pcs)
        if prev is None:                                    # первый запуск: просто запоминаем исходное состояние
            h["since"], h["snap"] = now_.isoformat(), cur
            _hist_save(h)
            return
        known_prev, known_cur = prev["orders"] | prev["new"], cur["orders"] | cur["new"]
        added = known_cur - known_prev
        removed = known_prev - known_cur
        routed = (cur["orders"] & prev["new"]) - prev["orders"]
        collected = (cur["collected"] - prev["collected"]) - added
        shipped = (cur["shipped"] - prev["shipped"]) - added
        found = False
        for kind, ids in (("added", added), ("routed", routed), ("collected", collected),
                          ("shipped", shipped), ("removed", removed)):
            if ids:
                found = True
                h["events"].append(dict(t=now_.isoformat(), k=kind, n=len(ids),
                                        pcs=sum(cur["pcs"].get(i, prev["pcs"].get(i, 0)) for i in ids)))
        h["events"] = h["events"][-HISTORY_KEEP:]
        changed = found or cur["orders"] != prev["orders"] or cur["new"] != prev["new"] \
            or cur["collected"] != prev["collected"] or cur["shipped"] != prev["shipped"]
        h["snap"] = cur
        if changed:
            _hist_save(h)


def _auto_track():
    """Подтягивает данные (из кэша) и фиксирует изменения — работает, пока дашборд открыт хотя бы у кого-то."""
    try:
        d = load()
        try:
            sh = build_ship(load_ship_tabs())
        except Exception:
            sh = None
        try:
            no = load_new_orders()
            no = no[~no["Заказ"].isin(d["Заказ"])]
        except Exception:
            no = None
        record_history(d, sh, no)
    except Exception:
        pass


def bell_html() -> str:
    """Колокольчик: при наведении — история изменений и сводка за сегодня."""
    h = _hist()
    with h["lock"]:
        events, cur, since = list(h["events"]), dict(h["cur"]), h["since"]
    nw = datetime.now(TZ)

    def tt(e):
        return datetime.fromisoformat(e["t"]).astimezone(TZ)

    today = [e for e in events if tt(e).date() == nw.date()]
    sums = {k: [0, 0] for k in EV_META}
    for e in today:
        sums[e["k"]][0] += e["n"]
        sums[e["k"]][1] += e["pcs"]
    dot = "<i class='bdot'></i>" if any((nw - tt(e)).total_seconds() < 86400 for e in events) else ""

    now_box = ""
    if cur:
        now_box = (f"Сейчас: не собрано <b>{n(cur['left'])}</b> зак. · бэклог <b>{n(cur['pcs_left'])}</b> шт"
                   f" · не маршрутизировано <b>{n(cur['new_cnt'])}</b>")
    td = (f"За сегодня: добавлено <b>{n(sums['added'][0])}</b> · собрано <b>{n(sums['collected'][0])}</b> "
          f"({n(sums['collected'][1])} шт) · маршрутизировано <b>{n(sums['routed'][0])}</b>")
    if cur:
        td += f"<br>Отгружено сегодня: <b>{n(cur['shipped_today'])}</b> зак. · <b>{n(cur['shipped_today_pcs'])}</b> шт"
    rows = []
    for e in reversed(events[-40:]):
        t = tt(e)
        label, col = EV_META.get(e["k"], (e["k"], "#8C93B5"))
        stamp = f"{t:%H:%M}" if t.date() == nw.date() else f"{t:%d.%m %H:%M}"
        rows.append(f"<div class='ev'><span class='t'>{stamp}</span><i style='background:{col}'></i>"
                    f"<span><b>{n(e['n'])}</b> {label} <small>· {n(e['pcs'])} шт</small></span></div>")
    if not rows:
        rows.append("<div class='empty'>Изменений пока не было. Они появятся здесь, как только в таблице "
                    "что-то добавится, соберётся или отгрузится.</div>")
    since_txt = f"{datetime.fromisoformat(since).astimezone(TZ):%d.%m %H:%M}" if since else "—"
    pop = (f"<div class='bell-pop'><div class='bp-h'>История изменений</div>"
           f"<div class='bp-box'>{now_box}<br>{td}</div>"
           f"<div class='bp-list'>{''.join(rows)}</div>"
           f"<div class='bp-f'>Изменения фиксируются при обновлении данных (раз в {CACHE_SECONDS // 60} мин, пока "
           f"дашборд открыт). Отслеживание с {since_txt}.</div></div>")
    return f"<div class='bell'>{ic('bell', P['mute'], 18)}{dot}{pop}</div>"


@st.fragment(run_every=BEAT_EVERY)
def topright():
    _auto_track()
    online_now()
    nw = datetime.now(TZ)
    # считаем людей (а не вкладки): два окна одного человека — один онлайн
    n_on = len({(r.get("vid") or k) for k, r in list(_presence2().items())}) or 1
    st.markdown(
        f"<div class='tb-right'><span class='live onl' title='Нажмите, чтобы увидеть, кто сейчас в дашборде'>"
        f"<i class='dot'></i><span>{n_on} онлайн</span></span><span class='sep'></span>"
        f"<span class='tb-date'>{nw:%d.%m.%Y}{ic('calendar', P['mute'], 15)}</span><span class='sep'></span>"
        f"{bell_html()}</div>",
        unsafe_allow_html=True)


def on_gsearch():
    """Поиск сверху: номер заказа -> «Поиск заказа», иначе товар/SKU -> «Движения по остаткам»."""
    q_ = (st.session_state.get("gsearch") or "").strip()
    if not q_:
        return
    st.session_state.gsearch = ""
    st.session_state.hub_open = False
    if re.match(r"^\s*(\w{1,3}-\d+|\d{3,9})\s*$", q_):
        st.session_state.page = PAGES[2]
        st.session_state.q_order = q_
    else:
        st.session_state.page = PAGES[4]
        st.session_state.mv_q = q_
        st.session_state.mg_q = q_


ensure_client()

# ── боковое меню: логотип, разделы с иконками, статус, профиль ──
with st.sidebar:
    st.markdown(logo_html(), unsafe_allow_html=True)
    st.markdown("<div style='height:18px'></div>", unsafe_allow_html=True)
    for i_, p_ in enumerate(PAGES):
        st.button(p_, key=f"nav_{i_}", width="stretch",
                  type="primary" if st.session_state.get("page", PAGES[0]) == p_ else "secondary",
                  on_click=set_page, args=(p_,))
    st.markdown("<div class='sidecard'><b><i></i>Данные актуальные</b></div>"
                "<div class='contact'><div class='av'>КО</div><div class='ct'>"
                f"<b>{esc(CONTACT_NAME)}</b>"
                f"<a href='tel:{esc(CONTACT_PHONE)}'>{esc(CONTACT_PHONE)}</a>"
                f"<a class='tg' href='{esc(CONTACT_TG_URL)}' target='_blank' rel='noopener noreferrer'>"
                f"{ic('send', 'currentColor', 13)} {esc(CONTACT_TG)}</a></div></div>"
                f"<a class='addr' href='{esc(ADDRESS_URL)}' target='_blank' rel='noopener noreferrer'>"
                f"{ic('pin', 'currentColor', 15)}<span>{esc(ADDRESS_TEXT)}</span></a>",
                unsafe_allow_html=True)

# ── верхняя панель ──
with st.container(key="topbar"):
    tb0, tb1, tb2, tb3 = st.columns([1.5, 4.5, 3.1, .55], vertical_alignment="center")
    with tb0:
        st.markdown(logo_html(LOGO_TOP_HEIGHT), unsafe_allow_html=True)
    with tb1:
        st.text_input("Поиск", key="gsearch", placeholder="Поиск по заказам, товарам, SKU…",
                      label_visibility="collapsed", on_change=on_gsearch)
    with tb2:
        topright()
    with tb3:
        st.button("Тема", key="theme_btn", on_click=toggle_theme, help="Переключить день / ночь")

# невидимые кнопки: скрипт ниже «нажимает» их, когда кликают на «N онлайн» или на логотип
if st.button("Онлайн", key="onl_open"):
    online_dialog()
st.button("На главную", key="go_home", on_click=set_page, args=(PAGES[0],))

# скрипт в родительской странице: Ctrl+K, клик на «онлайн»/логотип, cookie посетителя и имя из Telegram
components.html("""<script>try{var w=window.parent,d=w.document;
if(!d.__ck){d.__ck=1;
d.addEventListener('keydown',function(e){if((e.ctrlKey||e.metaKey)&&e.key.toLowerCase()==='k'){
e.preventDefault();var i=d.querySelector('.st-key-gsearch input');if(i){i.focus();i.select();}}});
d.addEventListener('click',function(e){var t=e.target;if(!t||!t.closest)return;
var q=t.closest('.onl')?'.st-key-onl_open button':(t.closest('.brand')?'.st-key-go_home button':'');
if(q){var b=d.querySelector(q);if(b){b.click();}}});}
var vid="__VID__",m=d.cookie.match(/(?:^|; )mc_vid=([^;]*)/);
if(!m||m[1]!==vid){d.cookie='mc_vid='+vid+'; max-age=31536000; path=/; SameSite=Lax';}
var nm='';try{var p=new URLSearchParams(w.location.hash.replace(/^#/,'')),dt=p.get('tgWebAppData');
if(dt){var u=JSON.parse(new URLSearchParams(dt).get('user')||'{}');
nm=((u.first_name||'')+' '+(u.last_name||'')).trim()||(u.username?'@'+u.username:'');}}catch(e){}
if(nm){var en=encodeURIComponent(nm),c2=d.cookie.match(/(?:^|; )mc_tg=([^;]*)/);
if(!c2||c2[1]!==en){d.cookie='mc_tg='+en+'; max-age=31536000; path=/; SameSite=Lax';}}
}catch(e){}</script>""".replace("__VID__", str(st.session_state.get("vid", ""))), height=0)

h1, h2, h3 = st.columns([6, 1.4, 1.5], vertical_alignment="bottom")
with h1:
    st.markdown(f"""<div class="top"><span class="crumb">{ic('home', P['mute'], 14)} Склад</span><h1>Сборка заказов</h1>
    <span class="upd">Обновлено {now:%H:%M} · Данные кэшируются на {CACHE_SECONDS // 60} мин</span></div>""",
                unsafe_allow_html=True)
with h2:
    st.button("День" if st.session_state.dark else "Ночь", key="btn_theme2", on_click=toggle_theme, width="stretch")
with h3:
    if st.button("Обновить", key="btn_refresh", width="stretch", type="primary"):
        st.cache_data.clear()
        st.session_state.pop("plan", None)      # рейсы тоже перечитаем из хранилища
        st.rerun()

if df.empty:
    st.info("В таблице нет заказов.")
    st.stop()

df["Новый"] = False
summary = day_summary(df)
dates = list(summary.index)

ship, ship_err = None, ""
try:
    ship = build_ship(load_ship_tabs())
except Exception as e:                      # дашборд работает и без таблицы отгрузок
    ship_err = str(e)

# новые (не маршрутизированные) заказы; те, что уже есть в основной таблице, не считаем повторно
new_orders, new_err = pd.DataFrame(columns=["Заказ", "Статус", "Дата", "Штук", "Магазин"]), ""
try:
    new_orders = load_new_orders()
    new_orders = new_orders[~new_orders["Заказ"].isin(df["Заказ"])].reset_index(drop=True)
except Exception as e:                      # дашборд работает и без вкладки «Новые заказы»
    new_err = str(e)
new_cnt = len(new_orders)
new_pcs = int(new_orders["Штук"].sum()) if new_cnt else 0

# общий список для страницы «Заказы»: обычные + новые (не собраны, статус «не маршрутизирован»)
new_rows = (new_orders.rename(columns={"Статус": "Статус WMS"})
            .assign(Адрес="", Город="", Сборка="", Собран=False, Новый=True))
df_all = pd.concat([df, new_rows], ignore_index=True)
all_dates = sorted(set(dates) | {x for x in new_orders["Дата"] if pd.notna(x)})

# состояние фильтров (до создания виджетов)
if st.session_state.get("f_date") not in [ALL] + all_dates:
    st.session_state.pop("f_date", None)
if "f_date" not in st.session_state:
    st.session_state.f_date = next((x for x in all_dates if x >= now.date()), all_dates[-1])
st.session_state.setdefault("f_mode", "Все")
st.session_state.setdefault("page", PAGES[0])
st.session_state.setdefault("trip_limit", 0)
st.session_state.setdefault("hub_open", False)

page = st.segmented_control("Раздел", PAGES, key="page", label_visibility="collapsed") or PAGES[0]
date_label = lambda x: "Все даты" if x == ALL else f"{fmt_date(x)}, {WEEKDAYS[x.weekday()]}"

if page != PAGES[0]:                        # ушли со «Сводки» — подробности хаба закрываем
    st.session_state.hub_open = False
hub_open = bool(st.session_state.hub_open) and page == PAGES[0]


# ───────────────────────── Хаб ББ: подробности по маршрутам ─────────────────────────
def hub_view():
    st.button("← Назад к сводке", key="hub_back", on_click=close_hub)
    st.markdown("<div class='sec'>Хаб отгрузка ББ: маршруты</div>", unsafe_allow_html=True)
    try:
        hub = load_hub()
    except Exception as e:
        st.warning("Не удалось прочитать «Хаб отгрузка ББ» из Metabase. Проверьте блок [metabase] в Secrets "
                   "(url, api_key) и доступ к таблице 61.")
        st.caption(str(e))
        return
    if hub.empty:
        st.info("В «Хаб отгрузка ББ» нет данных.")
        return

    ds = sorted(hub["date"].unique())
    today = now.date()
    future = [d_ for d_ in ds if d_ >= today]
    default = today if today in ds else (future[0] if future else ds[-1])
    options = [ALL] + ds
    if st.session_state.get("hub_date") not in options:
        st.session_state.hub_date = default
    st.selectbox("Дата отправки", options, key="hub_date", format_func=date_label)
    sel = st.session_state.hub_date

    sub = hub if sel == ALL else hub[hub["date"] == sel]
    ok_n = int(sub["ok"].sum())
    tot_n = len(sub)
    no_n = tot_n - ok_n
    pct = round(ok_n / tot_n * 100) if tot_n else 0

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        st.markdown(kpi_box(n(ok_n), "принято коробок", "", "ok"), unsafe_allow_html=True)
    with c2:
        st.markdown(kpi_box(n(no_n), "не принято коробок", "", "left" if no_n else "ok"), unsafe_allow_html=True)
    with c3:
        st.markdown(kpi_box(n(tot_n), "тотал, коробок", "", "all"), unsafe_allow_html=True)
    with c4:
        st.markdown(kpi_box(f"{pct}%", "тотал в процентах", "принято от общего числа", "blue"),
                    unsafe_allow_html=True)

    st.markdown("<div class='sec'>По маршрутам</div>", unsafe_allow_html=True)
    keys = ["route"] if sel != ALL else ["date", "route"]
    g = sub.groupby(keys).agg(ok=("ok", "sum"), tot=("ok", "size")).reset_index()
    g["no"] = g["tot"] - g["ok"]
    g["pct"] = (g["ok"] / g["tot"] * 100).round().astype(int)
    g = g.sort_values(["no", "route"], ascending=[False, True])

    head = (["Дата"] if sel == ALL else []) + ["Маршрут", "Принято коробок", "Не принято коробок", "Тотал", "", "% принято"]
    base_num = 2 if sel == ALL else 1
    rows = []
    for r in g.itertuples(index=False):
        bar = f"<div class='bxbar'><i style='width:{r.pct}%'></i></div>"
        cells = ([f"{fmt_date(r.date)}, {WEEKDAYS[r.date.weekday()]}"] if sel == ALL else []) + [
            f"<b>{esc(r.route)}</b>", n(r.ok), f"<b>{n(r.no)}</b>" if r.no else n(r.no), n(r.tot), bar,
            f"<b>{r.pct}%</b>"]
        rows.append(cells)
    rows.append(([""] if sel == ALL else []) + ["<b>Итого</b>", f"<b>{n(ok_n)}</b>", f"<b>{n(no_n)}</b>",
                                                f"<b>{n(tot_n)}</b>", "", f"<b>{pct}%</b>"])
    st.markdown(simple_table(head, rows, num_cols={base_num, base_num + 1, base_num + 2, base_num + 4}),
                unsafe_allow_html=True)

    if no_n == 0:
        st.success("Все коробки приняты.")
    label = "all" if sel == ALL else f"{sel:%Y-%m-%d}"

    def _build():
        bad = sub[~sub["ok"]].sort_values(["date", "route_full", "bar"])
        x = pd.DataFrame({"Маршрут": bad["route_full"].values,
                          "Не принятые штрих коды ББ": bad["bar"].values,
                          "Дата отправки": bad["date"].values})
        return to_xlsx({"Не принятые ББ": x})

    xl_button("Выгрузить не принятые ББ в Excel", _build, f"hub_ne_prinyato_{label}.xlsx", "dl_hub")
    st.caption("Принято — статус приёмки «Принят», остальные коробки считаются не принятыми. "
               "Excel: Маршрут, Не принятые штрих коды ББ, Дата отправки (одна строка = один штрихкод).")


if hub_open:
    hub_view()

# ───────────────────────── 1. Сводка по дням ─────────────────────────
if page == PAGES[0] and not hub_open:
    total = len(df)
    pcs_total = int(df["Штук"].sum())
    pcs_left = int(df.loc[~df["Собран"], "Штук"].sum()) + new_pcs      # бэклог = несобранное + новые не маршрутизированные

    if ship is not None and not ship.empty:
        sd = pick_ship_date(ship, now.date())
        sm = ship_summary(ship, sd)
        sync_html = sync_icon(ship, sd)
        pct_ship = round(sm["pcs_done"] / sm["pcs"] * 100) if sm["pcs"] else 0
        is_today = (sd == now.date()) and sm["total"] > 0
        all_done = sm["total"] > 0 and sm["done"] >= sm["total"]
        truck_stop = (not is_today) or all_done          # прошлый день или всё отгружено — машина стоит
        ship_html = f"""<div class="ship"><div class="sh-top">
          <div class="sh-t">{'Отгрузка сегодня' if sd == now.date() else 'Последняя отгрузка'}
          <span class="pdot"></span><small>{fmt_date(sd)}, {WEEKDAYS[sd.weekday()]}</small>{ic('clock', P['mute'], 15)}{sync_html}</div><div class="sh-p">{pct_ship}%</div></div>
          <div class="lane">{truck_html(pct_ship, truck_stop)}<div class="bar"><i style="width:{pct_ship}%"></i></div></div>
          <div class="sh-m">Отгружено <b>{n(sm['done'])}</b> из {n(sm['total'])} заказов ·
          <b>{n(sm['pcs_done'])}</b> из {n(sm['pcs'])} шт · осталось <b>{n(sm['total'] - sm['done'])}</b> зак.
          ({n(sm['pcs'] - sm['pcs_done'])} шт)</div></div>"""
        frame = (ship_doc(sd, pct_ship, sm, sync_html), SHIP_IFRAME_H) if is_today else None   # сегодня: линия + таймер + «Газель»
        if sd in dates:
            tap("ship", ship_html, go_orders, ("Все", sd), iframe=frame)
        elif frame:
            components.html(frame[0], height=frame[1])
        else:
            st.markdown(ship_html, unsafe_allow_html=True)
    elif ship_err:
        st.caption("Статус отгрузки недоступен: проверьте доступ к таблице отгрузок.")

    # «Хаб ББ» из Metabase (кэш 5 минут)
    hub_g, hub_err = None, ""
    try:
        hub_g = hub_by_date(load_hub())
    except Exception as e:                  # дашборд работает и без Metabase
        hub_err = str(e)

    ser = summary.tail(14)
    s_pcs = (ser["Штук"] - ser["Штук собрано"]).tolist()
    k1, k2, k3 = st.columns(3)
    with k1:
        tap("all", kpi_box(None, "WMS Сборка шт", "", "all", status=status_breakdown(df_all)),
            go_orders, ("Все", ALL))
    with k2:
        tap("hub", hub_card(hub_g, now.date(), hub_err), open_hub)
    with k3:
        tap("pcs", kpi_box(n(pcs_left), "бэклог, шт", "", "blue", spark=s_pcs),
            go_orders, ("Не собрано", ALL))
    st.caption("Нажмите на карточку, чтобы открыть подробности.")
    if hub_err:
        st.caption(f"Хаб ББ недоступен: {hub_err}")

    # ── новые заказы по датам доставки: нажатие открывает список «Не собрано» на эту дату ──
    if new_err:
        st.caption(f"Новые заказы недоступны: {new_err}")
    elif new_cnt:
        has_date = new_orders["Дата"].map(pd.notna)
        groups = [(d_, g_) for d_, g_ in new_orders[has_date].groupby("Дата")]
        groups.sort(key=lambda x: x[0])
        if (~has_date).any():
            groups.append((None, new_orders[~has_date]))
        for d_, g_ in groups:
            if d_ is None:
                d_html, key, target = "—<small>без даты</small>", "new_none", ALL
            else:
                d_html = f"{fmt_date(d_)}<small>{WEEKDAYS[d_.weekday()]}, {d_:%Y}</small>"
                key, target = f"new_{d_:%Y%m%d}", d_
            tap(key, f"""<div class="day {'urg' if d_ is not None and (d_ - now.date()).days <= 3 else 'amb'}">
              <div class="d">{d_html}</div>
              <div class="mid"><span class="badge no">не маршрутизировано</span>
              <div class="meta">заказов <b>{n(len(g_))}</b> · <b>{n(g_['Штук'].sum())}</b> шт</div></div>
              <div class="pct">{n(g_['Штук'].sum())}<small>шт</small></div><div class="chev">{ic('chev', P['mute'], 16)}</div></div>""",
                go_orders, ("Не собрано", target))

    xl_button("Выгрузить бэклог в Excel",
              lambda: backlog_excel(df, new_orders), f"backlog_{now:%Y-%m-%d}.xlsx", "dl_backlog")

    show_all = st.toggle(f"Показать все дни (по умолчанию только дни с готовностью ниже {HIDE_FROM_DAYS_PCT}%)",
                         value=False)
    shown = summary if show_all else summary[summary["Процент"] < HIDE_FROM_DAYS_PCT]
    hidden = len(summary) - len(shown)
    if hidden:
        st.caption(f"Скрыто дней с готовностью {HIDE_FROM_DAYS_PCT}% и выше: {hidden}")
    if shown.empty:
        st.success(f"Все дни собраны на {HIDE_FROM_DAYS_PCT}% и выше.")

    for d, r in shown.iterrows():
        full = "full" if r["Осталось"] == 0 else ""
        left_pcs = r["Штук"] - r["Штук собрано"]
        tap(f"day_{d:%Y%m%d}", f"""<div class="day {full}">
          <div class="d">{fmt_date(d)}<small>{WEEKDAYS[d.weekday()]}, {d:%Y}</small></div>
          <div class="mid"><div class="bar"><i style="width:{r['Процент']}%"></i></div>
          <div class="meta">собрано <b>{r['Собрано']}</b> зак. ({n(r['Штук собрано'])} шт) · осталось <b>{r['Осталось']}</b> зак.
          ({n(left_pcs)} шт) · всего {r['Заказов']} зак., {n(r['Штук'])} шт</div></div>
          <div class="pct">{r['Процент']}%<small>собрано</small></div><div class="chev">{ic('chev', P['mute'], 16)}</div></div>""",
            go_orders, ("Собрано" if r["Осталось"] == 0 else "Не собрано", d))

# ───────────────────────── 2. Заказы ─────────────────────────
if page == PAGES[1]:
    st.selectbox("Дата отгрузки", [ALL] + all_dates, key="f_date", format_func=date_label)
    st.segmented_control("Показать", MODES, key="f_mode")
    mode = st.session_state.f_mode or "Все"

    sub = df_all if st.session_state.f_date == ALL else df_all[df_all["Дата"] == st.session_state.f_date]
    sub_main = sub[~sub["Новый"]]           # для маршрутизации и рейсов — только маршрутизированные заказы
    done_pcs = int(sub.loc[sub["Собран"], "Штук"].sum())
    left_pcs = int(sub.loc[~sub["Собран"], "Штук"].sum())
    st.caption(f"За выбранную дату: собрано {n(done_pcs)} шт · осталось {n(left_pcs)} шт · "
               f"заказов {len(sub)} · всего {n(sub['Штук'].sum())} шт")
    if mode == "Прочие заказы":
        # статистика по магазинам из EXCLUDE_STORES: все даты сразу, без разбивки по дням; собранные и не собранные
        oth = df_all.copy()
        oth["_k"] = oth["Магазин"].map(norm_store)
        oth = oth[oth["_k"].isin(EXCLUDED_KEYS)]

        def _sorted(f):
            return (f.assign(_d=f["Дата"].map(lambda x: x if pd.notna(x) else date.max))
                    .sort_values(["_d", "Заказ"]).drop(columns="_d"))

        sel = st.session_state.get("oth_store")
        if sel not in EXCLUDE_STORES:
            sel = None

        if sel is None:
            st.markdown("<div class='sec'>Прочие заказы по магазинам</div>", unsafe_allow_html=True)
            cols = st.columns(3)
            for i, name in enumerate(EXCLUDE_STORES):
                part = oth[oth["_k"] == norm_store(name)]
                nc, dn = part[~part["Собран"]], part[part["Собран"]]
                with cols[i % 3]:
                    tap(f"oth_{i}", kpi_box(n(len(nc)), name,
                        f"не собрано заказов · {n(nc['Штук'].sum())} шт<br>"
                        f"собрано: {n(len(dn))} зак. · {n(dn['Штук'].sum())} шт",
                        "left" if len(nc) else "ok", icon="box"),
                        set_oth, (name,))
            st.caption("Нажмите на карточку, чтобы открыть список заказов магазина.")
        else:
            part = oth[oth["_k"] == norm_store(sel)]
            is_uz = norm_store(sel) == norm_store("Узуммаркет")
            if is_uz:
                try:
                    cpmap = load_cat_prio()
                except Exception as e:
                    cpmap = {}
                    st.warning("Не удалось прочитать категории и приоритеты. Откройте таблицу для почты сервисного "
                               "аккаунта (роль «Читатель»).")
                    st.caption(str(e))
                key_ = part["Заказ"].str.strip().str.upper()
                part = part.assign(Категория=key_.map(lambda k: cpmap.get(k, ("", ""))[0]),
                                   Приоритет=key_.map(lambda k: cpmap.get(k, ("", ""))[1]))
            nc, dn = part[~part["Собран"]], part[part["Собран"]]
            st.button("← Все прочие заказы", on_click=set_oth, args=(None,), key="oth_back")
            st.markdown(f"<div class='sec'>{esc(sel)}</div>", unsafe_allow_html=True)
            o1, o2, o3, o4 = st.columns(4)
            with o1:
                st.markdown(kpi_box(n(len(nc)), "не собрано заказов", "", "left" if len(nc) else "ok"),
                            unsafe_allow_html=True)
            with o2:
                st.markdown(kpi_box(n(nc["Штук"].sum()), "не собрано, шт", "", "left" if len(nc) else "ok"),
                            unsafe_allow_html=True)
            with o3:
                st.markdown(kpi_box(n(len(dn)), "собрано заказов", "", "ok"), unsafe_allow_html=True)
            with o4:
                st.markdown(kpi_box(n(dn["Штук"].sum()), "собрано, шт", "", "ok"), unsafe_allow_html=True)

            if part.empty:
                st.info("Заказов этого магазина нет.")
            else:
                st.markdown("<div class='sec'>Не собранные заказы</div>", unsafe_allow_html=True)
                if nc.empty:
                    st.success("Не собранных заказов нет.")
                else:
                    st.markdown(orders_table(_sorted(nc), cp=is_uz), unsafe_allow_html=True)
                st.markdown("<div class='sec'>Собранные заказы</div>", unsafe_allow_html=True)
                if dn.empty:
                    st.info("Собранных заказов нет.")
                else:
                    st.markdown(orders_table(_sorted(dn), cp=is_uz), unsafe_allow_html=True)
    elif mode == "Маршрутизация":
        pivot, dropped = split_routing(sub_main)
        if pivot.empty:
            st.info("Нет заказов для сводной.")
        else:
            # ручная раскладка по машинам (по выбранной дате)
            if st.session_state.f_date == ALL:
                st.info("Чтобы собрать рейсы, выберите конкретную дату отгрузки.")
            else:
                plan_ui(f"{st.session_state.f_date:%Y-%m-%d}", stores_of(pivot))

            # сводная по выбранной дате: все заказы даты, без исключённых магазинов
            st.markdown("<div class='sec'>Сводная по точкам</div>", unsafe_allow_html=True)
            st.caption(f"В сводной: {n(len(sub_main) - len(dropped))} заказов · {n(pivot['Штук'].sum())} шт · "
                       f"точек: {len(pivot)}. Не включены магазины: {', '.join(EXCLUDE_STORES)}"
                       + (f" (исключено {n(len(dropped))} заказов, {n(dropped['Штук'].sum())} шт)." if len(dropped) else "."))
            label = "all" if st.session_state.f_date == ALL else f"{st.session_state.f_date:%Y-%m-%d}"
            xl_button("Скачать сводную в Excel", lambda: routing_excel(pivot),
                      f"marshrutizatsiya_{label}.xlsx", "dl_routing")
            st.markdown(routing_table(pivot), unsafe_allow_html=True)
    elif mode == "Рейсы":
        if st.session_state.f_date == ALL:
            st.info("Выберите конкретную дату отгрузки: рейсы собираются на один день.")
        else:
            pivot, dropped = split_routing(sub_main)
            if pivot.empty:
                st.info("Нет заказов для рейсов.")
            else:
                ensure_plan()
                dk = f"{st.session_state.f_date:%Y-%m-%d}"
                stores = stores_of(pivot)
                limit = int(st.session_state.get("lim_val", 0))
                machines = plan_machines(dk, stores, limit)
                assigned = {s_["shop"] for m in machines for s_ in m["stores"]}
                left = [s for s in stores if s not in assigned]

                if not machines:
                    st.info("Рейсы на эту дату ещё не собраны. Откройте «Маршрутизация», выберите машину "
                            "и добавьте в неё магазины.")
                else:
                    coords = {}
                    try:
                        coords = load_coords()
                    except Exception as e:
                        st.caption(f"Кнопки карт недоступны: не удалось прочитать координаты ({e}).")
                    st.markdown(f"<div class='chips big'><span class='chip'>Машин <b>{len(machines)}</b></span>"
                                f"<span class='chip'>Точек <b>{len(assigned)}</b></span>"
                                f"<span class='chip'>Штук <b>{n(sum(m['qty'] for m in machines))}</b></span>"
                                + (f"<span class='chip'>Лимит <b>{n(limit)}</b></span>" if limit else "")
                                + "</div>", unsafe_allow_html=True)
                    if left:
                        st.warning("Не распределено по машинам: " + ", ".join(
                            f"{s} ({n(stores[s]['qty'])} шт)" for s in left))
                    xl_button("Скачать рейсы в Excel",
                              lambda: trips_excel(machines, f"{st.session_state.f_date:%d.%m.%Y}", limit, example=False),
                              f"reisy_{st.session_state.f_date:%Y-%m-%d}.xlsx", "dl_trips")
                    st.markdown(trips_table(machines, coords), unsafe_allow_html=True)
    else:
        if mode == "Не собрано":
            sub = sub[~sub["Собран"]]
        elif mode == "Собрано":
            sub = sub[sub["Собран"]]

        sub = (sub.assign(_d=sub["Дата"].map(lambda x: x if pd.notna(x) else date.max))
               .sort_values(["_d", "Заказ"]).drop(columns="_d"))
        st.markdown(f"**{mode}: {len(sub)} заказов · {n(sub['Штук'].sum())} шт**")
        if sub.empty:
            st.info("Заказов нет.")
        else:
            st.markdown(orders_table(sub), unsafe_allow_html=True)
            if len(sub) > 1500:
                st.caption("Показаны первые 1500 заказов, уточните дату.")

# ───────────────────────── 3. Поиск ─────────────────────────
if page == PAGES[2]:
    q = st.text_input("Номер заказа", key="q_order", placeholder="например, 1M-0024629 или 0024629").strip()
    if q:
        hits = df[df["Заказ"].str.contains(q, case=False, regex=False)]
        if hits.empty:
            st.warning(f"Заказ «{q}» не найден. Можно вводить только последние цифры номера.")
        else:
            st.caption(f"Найдено: {len(hits)}")
            for r in hits.head(30).to_dict("records"):
                st.markdown(f"""<div class="hit"><div class="n">{esc(r['Заказ'])} &nbsp;{badge(r['Собран'])}</div>
                <div class="s">{esc(r['Магазин'])} · {esc(r['Город'])}</div>
                Отгрузка: <b>{fmt_date(r['Дата'])}</b> · {n(r['Штук'])} шт · WMS: <b>{esc(r['Статус WMS'] or '—')}</b></div>""",
                            unsafe_allow_html=True)
    else:
        st.info("Введите номер заказа или его часть.")

# ───────────────────────── 4. Пустые короба ─────────────────────────
if page == PAGES[3]:
    bx = None
    try:
        bx = load_boxes()
    except Exception as e:
        mail = ""
        try:
            mail = st.secrets["gcp_service_account"].get("client_email", "")
        except Exception:
            pass
        st.warning("Не удалось прочитать таблицу пустых коробов. Откройте её для почты сервисного "
                   f"аккаунта {mail} (роль «Читатель»).")
        st.caption(str(e))

    if bx is not None:
        if bx.empty:
            st.info("В таблице пока нет записей.")
        else:
            months = sorted(bx["Месяц"].unique(), reverse=True)
            cur_key = f"{now:%Y-%m}"
            if st.session_state.get("b_month") not in months:
                st.session_state.b_month = cur_key if cur_key in months else months[0]
            shops = sorted(bx["Магазин"].unique(), key=lambda s: s.lower())
            if st.session_state.get("b_shop") not in [ALL] + shops:
                st.session_state.b_shop = ALL

            f1, f2 = st.columns(2)
            with f1:
                st.selectbox("Месяц", months, key="b_month", format_func=month_label)
            with f2:
                st.selectbox("Магазин", [ALL] + shops, key="b_shop",
                             format_func=lambda s: "Все магазины" if s == ALL else s)
            mk, shop = st.session_state.b_month, st.session_state.b_shop

            base = bx if shop == ALL else bx[bx["Магазин"] == shop]
            m = base[base["Месяц"] == mk]
            title = month_label(mk) + ("" if shop == ALL else f" · {shop}")
            st.markdown(f"<div class='sec'>Пустые короба: {esc(title)}</div>", unsafe_allow_html=True)

            if m.empty:
                st.info("За выбранный месяц записей нет.")
            else:
                last = m.iloc[-1]
                total_b, recs, n_shops = int(m["Коробок"].sum()), len(m), m["Магазин"].nunique()
                c1, c2, c3, c4 = st.columns(4)
                with c1:
                    st.markdown(kpi_box(n(total_b), "коробок за месяц", f"{recs} записей"), unsafe_allow_html=True)
                with c2:
                    if shop == ALL:
                        st.markdown(kpi_box(n(n_shops), "магазинов", "прислали данные", "all"), unsafe_allow_html=True)
                    else:
                        st.markdown(kpi_box(n(round(total_b / recs)), "в среднем за запись", "коробок", "all"),
                                    unsafe_allow_html=True)
                with c3:
                    top = m.groupby("Магазин")["Коробок"].sum().sort_values(ascending=False)
                    if shop == ALL:
                        st.markdown(kpi_box(esc(top.index[0]), "больше всего коробок",
                                            f"{n(top.iloc[0])} шт", "left").replace("class='v'", "class='v' style='font-size:1.2rem'"),
                                    unsafe_allow_html=True)
                    else:
                        mx = m.loc[m["Коробок"].idxmax()]
                        st.markdown(kpi_box(n(mx["Коробок"]), "максимум за запись",
                                            f"{mx['Время']:%d.%m.%Y}", "left"), unsafe_allow_html=True)
                with c4:
                    st.markdown(kpi_box(n(last["Коробок"]), "последняя запись",
                                        f"{esc(last['Магазин'])} · {last['Время']:%d.%m %H:%M}", "ok"),
                                unsafe_allow_html=True)

                if shop == ALL:
                    st.markdown("<div class='sec'>Магазины за месяц</div>", unsafe_allow_html=True)
                    g = (m.groupby("Магазин").agg(Записей=("Коробок", "count"), Коробок=("Коробок", "sum"),
                                                  Последняя=("Время", "max"))
                         .sort_values("Коробок", ascending=False).reset_index())
                    mx = max(int(g["Коробок"].max()), 1)
                    rows = [[f"<b>{esc(r.Магазин)}</b>", n(r.Записей),
                             f"<div class='bxbar'><i style='width:{r.Коробок / mx * 100:.0f}%'></i></div>",
                             f"<b>{n(r.Коробок)}</b>", f"{r.Последняя:%d.%m %H:%M}"] for r in g.itertuples(index=False)]
                    st.markdown(simple_table(["Магазин", "Записей", "", "Коробок", "Последняя запись"], rows,
                                             num_cols={1, 3}), unsafe_allow_html=True)

                    st.markdown("<div class='sec'>По дням</div>", unsafe_allow_html=True)
                    dd = (m.assign(День=m["Время"].dt.date).groupby("День")
                          .agg(Записей=("Коробок", "count"), Коробок=("Коробок", "sum"))
                          .sort_index(ascending=False).reset_index())
                    mxd = max(int(dd["Коробок"].max()), 1)
                    rows = [[f"{r.День:%d.%m}, {WEEKDAYS[r.День.weekday()]}", n(r.Записей),
                             f"<div class='bxbar'><i style='width:{r.Коробок / mxd * 100:.0f}%'></i></div>",
                             f"<b>{n(r.Коробок)}</b>"] for r in dd.itertuples(index=False)]
                    st.markdown(simple_table(["День", "Записей", "", "Коробок"], rows, num_cols={1, 3}),
                                unsafe_allow_html=True)
                else:
                    all_months = st.toggle("История за все месяцы", value=False, key="b_all")
                    hist = base if all_months else m
                    st.markdown(f"<div class='sec'>История: {esc(shop)}"
                                f"{'' if all_months else ' · ' + esc(month_label(mk))}</div>", unsafe_allow_html=True)
                    h = hist.sort_values("Время", ascending=False)
                    rows = [[f"{r.Время:%d.%m.%Y %H:%M}", f"<b>{esc(r.Магазин)}</b>", f"<b>{n(r.Коробок)}</b>"]
                            for r in h.head(1500).itertuples(index=False)]
                    st.markdown(simple_table(["Дата и время", "Магазин", "Коробок"], rows, num_cols={2}),
                                unsafe_allow_html=True)

            # сводная по месяцам (с учётом выбранного магазина)
            st.markdown("<div class='sec'>Сводка по месяцам</div>", unsafe_allow_html=True)
            ms = (base.groupby("Месяц").agg(Записей=("Коробок", "count"), Коробок=("Коробок", "sum"),
                                            Магазинов=("Магазин", "nunique"))
                  .sort_index(ascending=False).reset_index())
            mxm = max(int(ms["Коробок"].max()), 1)
            rows = [[f"<b>{esc(month_label(r.Месяц))}</b>", n(r.Записей), n(r.Магазинов),
                     f"<div class='bxbar'><i style='width:{r.Коробок / mxm * 100:.0f}%'></i></div>",
                     f"<b>{n(r.Коробок)}</b>"] for r in ms.itertuples(index=False)]
            st.markdown(simple_table(["Месяц", "Записей", "Магазинов", "", "Коробок"], rows, num_cols={1, 2, 4}),
                        unsafe_allow_html=True)

# ───────────────────────── 5. Движения по остаткам ─────────────────────────
if page == PAGES[4]:
    st.markdown("<div class='sec'>Движения по остаткам</div>", unsafe_allow_html=True)

    # ── Остаток позиции по Magnit Stock ──
    st.markdown("<div class='sec'>Остаток позиции (Magnit Stock)</div>", unsafe_allow_html=True)
    st.caption("Актуальные данные из таблицы Magnit Stock. Найдите одну позицию или загрузите список: "
               "в колонке A файла могут быть Внешний ID позиции, Товар или Sku Wms ID (можно вперемешку).")
    mi1, mi2 = st.columns(2)
    with mi1:
        mg_q = st.text_input("Одна позиция: название, Внешний ID или Sku Wms ID", key="mg_q").strip()
    with mi2:
        mg_up = st.file_uploader("Список позиций: Excel (.xlsx) или CSV, значения в колонке A",
                                 type=["xlsx", "csv"], key="mg_up")
    mg_paste = st.text_area("Или вставьте список, по одной позиции в строке", key="mg_paste", height=90)

    mg_items = [mg_q] if mg_q else []
    if mg_up is not None:
        try:
            mg_items += read_item_list(mg_up)
        except Exception as e:
            st.error(f"Не удалось прочитать файл: {e}")
    mg_items += [x.strip() for x in mg_paste.splitlines() if x.strip()]
    mg_items = list(dict.fromkeys(mg_items))

    if mg_items:
        mg_raw, mg_df = None, None
        try:
            mg_raw = load_magnit()
            mg_df = magnit_norm(mg_raw)
        except Exception as e:
            st.error(f"Не удалось прочитать Magnit Stock: {e}")
        if mg_df is not None:
            mg_sub, mg_missing = magnit_find(mg_df, mg_items)
            if mg_missing:
                st.warning(f"Не найдено позиций: {len(mg_missing)} из {len(mg_items)}. "
                           + ", ".join(esc(x) for x in mg_missing[:15]) + (" …" if len(mg_missing) > 15 else ""))
            if mg_sub.empty:
                st.info("По этим значениям в Magnit Stock ничего нет.")
            else:
                rep_m = magnit_rep(mg_sub, mg_raw, mg_missing)
                zt, sk = stock_now_tables(rep_m)

                tot_can, tot_no = int(sk["can"].sum()), int(sk["no"].sum())
                k1, k2, k3, k4 = st.columns(4)
                with k1:
                    st.markdown(kpi_box(n(tot_can), "можно зарезервировать, шт", "доступно в зонах с резервом", "ok"),
                                unsafe_allow_html=True)
                with k2:
                    st.markdown(kpi_box(n(tot_no), "не зарезервируется, шт", "доступно в зонах без резерва",
                                        "left" if tot_no else "ok"), unsafe_allow_html=True)
                with k3:
                    st.markdown(kpi_box(n(sk["loc"].sum()), "в ячейках всего, шт", f"позиций: {len(sk)}", "all"),
                                unsafe_allow_html=True)
                with k4:
                    st.markdown(kpi_box(n(sk["res"].sum()), "уже в резерве, шт", "Reserved", "blue"),
                                unsafe_allow_html=True)

                rows = [[esc(r.sku), esc(r.name), n(r.loc), n(r.res), f"<b>{n(r.can)}</b>", f"<b>{n(r.no)}</b>"]
                        for r in sk.head(500).itertuples(index=False)]
                st.markdown(simple_table(["Внешний ID позиции", "Товар", "В ячейках", "Резерв",
                                          "Можно зарезервировать", "Не зарезервируется"], rows,
                                         num_cols={2, 3, 4, 5}), unsafe_allow_html=True)

                st.markdown("<div class='sec'>По зонам</div>", unsafe_allow_html=True)
                badge_of = {"yes": "<span class='badge ok'>можно зарезервировать</span>",
                            "no": "<span class='badge no'>Заказ не зарезервируется</span>",
                            "?": "<span class='chip wait'>зона не указана в правилах</span>"}
                multi = zt["sku"].nunique() > 1
                head = (["Товар"] if multi else []) + ["Зона", "В ячейках", "Резерв", "Доступно", "Комментарий"]
                bi = 1 if multi else 0
                rows = [(([f"{esc(r.name)} · {esc(r.sku)}"] if multi else []) +
                         [f"<b>{esc(r.zone)}</b>", n(r.loc), n(r.res), f"<b>{n(r.avail)}</b>", badge_of.get(r.kind, "")])
                        for r in zt.head(500).itertuples(index=False)]
                st.markdown(simple_table(head, rows, num_cols={bi + 1, bi + 2, bi + 3}), unsafe_allow_html=True)

                try:
                    dl_button("Скачать Excel: остаток позиций (Magnit Stock)", stock_now_excel(rep_m),
                              f"ostatok_magnit_{now:%Y%m%d_%H%M}.xlsx", "dl_mg_now")
                except Exception as e:
                    st.error(f"Не удалось собрать Excel: {e}")
                st.caption("Можно зарезервировать = «Доступно» в зонах с резервом, не зарезервируется = «Доступно» в зонах "
                           "без резерва. В Excel: детали по строкам с итогами, сводка по зонам, итоги по товарам, "
                           "все данные Magnit Stock и лист «Не найдено».")

    # ── Движения по снимкам остатков ──
    st.markdown("<div class='sec'>Движения за период</div>", unsafe_allow_html=True)
    cat = None
    try:
        cat = load_stock_catalog()
    except Exception as e:
        st.warning("Не удалось прочитать остатки из Metabase. Проверьте блок [metabase] в Secrets "
                   "(url, api_key, stock_table) и доступность Metabase из интернета.")
        st.caption(str(e))

    if cat is not None:
        today = now.date()
        dmin, dmax = today - timedelta(days=MB_STOCK_DAYS), today
        st.caption(f"Товаров в каталоге: {n(len(cat))} · выбрать период можно за последние {MB_STOCK_DAYS} дней")

        q = st.text_input("Название товара или SKU WMS ID", key="mv_q",
                          placeholder="например, MIXIT или 100000000000011427").strip()
        if q:
            found = cat[cat["sku"].str.contains(q, case=False, regex=False)
                        | cat["name"].str.contains(q, case=False, regex=False)]
        else:
            found = cat.iloc[0:0]
        opts = found["sku"].tolist()[:200]
        nm = dict(zip(cat["sku"], cat["name"]))
        if q and found.empty:
            st.warning("Ничего не найдено.")
        elif q:
            st.caption(f"Найдено товаров: {len(found)}" + (" (показаны первые 200, уточните запрос)" if len(found) > 200 else ""))
        chosen = st.multiselect("Товары", opts, default=opts if len(opts) <= 10 else [],
                                format_func=lambda s: f"{nm.get(s, '')} · {s}", key=f"mv_sel_{q}")

        period = st.date_input("Период (для движений)", value=(max(dmin, dmax - timedelta(days=13)), dmax),
                               min_value=dmin, max_value=dmax, key="mv_period", format="DD.MM.YYYY")
        d1, d2 = (period[0], period[1]) if isinstance(period, (tuple, list)) and len(period) == 2 else (None, None)

        if st.button("Сформировать Excel", key="mv_go", disabled=not chosen or d1 is None):
            try:
                hist, raw_rows = load_stock_rows(tuple(sorted(chosen)), d1 - timedelta(days=7), d2)
                rep = stock_report(hist, chosen, d1, d2, raw_rows)
                if len(rep["dates"]) < 2:
                    st.session_state.mv_res = None
                    st.warning("В выбранном периоде нет снимков для сравнения. Расширьте период.")
                else:
                    sheets = {"Хронология движений": rep["movements"], "Итого по товарам": rep["summary"],
                              "Остаток по дням": rep["daily"], "Последний остаток": rep["last"]}
                    if not rep["all"].empty:
                        sheets["Все данные по товару"] = rep["all"]
                    sheets["Параметры"] = rep["params"]
                    data = to_xlsx(sheets)
                    st.session_state.mv_res = dict(sig=(tuple(chosen), d1, d2), data=data, rep=rep,
                                                   name=f"dvizheniya_{d1:%Y%m%d}_{d2:%Y%m%d}.xlsx")
            except Exception as e:
                st.session_state.mv_res = None
                st.error(f"Не удалось сформировать файл: {e}")

        res = st.session_state.get("mv_res")
        if res and res["sig"] == (tuple(chosen), d1, d2):
            rep = res["rep"]
            dl_button("Скачать Excel: движения по остаткам", res["data"], res["name"], "dl_mv")
            sm_ = rep["summary"]
            rows = [[esc(r["Sku Wms ID"]), esc(r["Товар"]), n(r["Остаток на начало"]), n(r["Остаток на конец"]),
                     n(r["Доступно на конец"]), n(r["Резерв на конец"]),
                     n(r["Поступило"]), n(r["Убыло"]), n(r["Движений"])] for _, r in sm_.iterrows()]
            st.markdown(simple_table(["Sku Wms ID", "Товар", "На начало", "На конец", "Доступно", "Резерв",
                                      "Поступило", "Убыло", "Движений"],
                                     rows, num_cols={2, 3, 4, 5, 6, 7, 8}), unsafe_allow_html=True)
            st.caption("Резерв = «В ячейках» − «Доступно». Движения считаются между ежедневными снимками. "
                       "Полная хронология, остаток по дням, последний остаток по ячейкам и все данные по товару — в файле.")

    # ── выгрузка Magnit Stock ──
    st.markdown("<div class='sec'>Magnit Stock</div>", unsafe_allow_html=True)
    st.caption("Загружает таблицу Magnit Stock из Metabase целиком и готовит файл Excel.")
    if st.button("Подготовить файл Magnit Stock", key="mg_go"):
        try:
            mg = load_magnit()
            st.session_state.mg_res = dict(
                data=to_xlsx({"Magnit Stock": mg}), rows=len(mg), name=f"magnit_stock_{now:%Y-%m-%d_%H%M}.xlsx")
        except Exception as e:
            st.session_state.mg_res = None
            st.error(f"Не удалось выгрузить Magnit Stock: {e}")
    mg_res = st.session_state.get("mg_res")
    if mg_res:
        st.caption(f"Готово: {n(mg_res['rows'])} строк.")
        dl_button("Скачать Magnit Stock (Excel)", mg_res["data"], mg_res["name"], "dl_magnit")

# ───────────────────────── 6. Аналитика ─────────────────────────
if page == PAGES[5]:
    import altair as alt
    st.markdown("<div class='sec'>Аналитика: количество штук</div>", unsafe_allow_html=True)

    src = df_all.dropna(subset=["Дата"])
    if src.empty:
        st.info("Нет данных для графика.")
    else:
        by = st.segmented_control("Группировка", ["По дням", "По неделям"], key="an_by") or "По дням"
        weekly = by == "По неделям"
        dmin_a, dmax_a = min(src["Дата"]), max(src["Дата"])
        start = max(dmin_a, dmax_a - timedelta(days=55 if weekly else 13))
        per = st.date_input("Период", value=(start, dmax_a), min_value=dmin_a, max_value=dmax_a,
                            key=f"an_period_{'w' if weekly else 'd'}", format="DD.MM.YYYY")
        if isinstance(per, (tuple, list)) and len(per) == 2:
            a1, a2 = per
            part = src[(src["Дата"] >= a1) & (src["Дата"] <= a2)].copy()
            if part.empty:
                st.info("За выбранный период заказов нет.")
            else:
                part["Статус"] = np.where(part["Собран"], "Собрано", "Не собрано")
                tot = int(part["Штук"].sum())
                done_ = int(part.loc[part["Собран"], "Штук"].sum())
                colors = alt.Scale(domain=["Собрано", "Не собрано"], range=["#1E9E5C", "#E8A33D"])

                if not weekly:
                    g = part.groupby(["Дата", "Статус"], as_index=False)["Штук"].sum()
                    g["Дата"] = pd.to_datetime(g["Дата"])
                    c1, c2, c3 = st.columns(3)
                    with c1:
                        st.markdown(kpi_box(n(tot), "всего штук за период", f"{a1:%d.%m} — {a2:%d.%m}", "all"),
                                    unsafe_allow_html=True)
                    with c2:
                        st.markdown(kpi_box(n(done_), "собрано, шт", "", "ok"), unsafe_allow_html=True)
                    with c3:
                        st.markdown(kpi_box(n(tot - done_), "не собрано, шт", "", "left"), unsafe_allow_html=True)
                    chart = (alt.Chart(g).mark_bar(size=22 if len(g) < 40 else 10)
                             .encode(x=alt.X("yearmonthdate(Дата):T", title=None,
                                             axis=alt.Axis(format="%d.%m", labelAngle=-45)),
                                     y=alt.Y("Штук:Q", title="Штук", axis=alt.Axis(format=",.0f")),
                                     color=alt.Color("Статус:N", title=None, scale=colors,
                                                     legend=alt.Legend(orient="top")),
                                     tooltip=[alt.Tooltip("yearmonthdate(Дата):T", title="Дата", format="%d.%m.%Y"),
                                              alt.Tooltip("Статус:N"),
                                              alt.Tooltip("Штук:Q", title="Штук", format=",.0f")])
                             .properties(height=400))
                else:
                    dd = pd.to_datetime(part["Дата"])
                    part["W"] = (dd - pd.to_timedelta(dd.dt.weekday, unit="D")).dt.normalize()
                    wk = (part.groupby("W")
                          .agg(Штук=("Штук", "sum"),
                               Собрано=("Штук", lambda s: int(s[part.loc[s.index, "Собран"]].sum())))
                          .reset_index().sort_values("W"))
                    wk["Не собрано"] = wk["Штук"] - wk["Собрано"]
                    wk["Неделя"] = (wk["W"].dt.strftime("%d.%m") + "–" +
                                    (wk["W"] + pd.Timedelta(days=6)).dt.strftime("%d.%m"))
                    wk["Δ"] = wk["Штук"].diff()
                    wk["Δ%"] = wk["Штук"].pct_change() * 100
                    order = wk["Неделя"].tolist()

                    last_w = wk.iloc[-1]
                    if len(wk) > 1 and pd.notna(last_w["Δ%"]):
                        arrow = "▲" if last_w["Δ"] > 0 else ("▼" if last_w["Δ"] < 0 else "■")
                        sub3 = f"{arrow} {last_w['Δ']:+,.0f} шт ({last_w['Δ%']:+.0f}%) к прошлой неделе".replace(",", " ")
                    else:
                        sub3 = "нет предыдущей недели для сравнения"
                    c1, c2, c3 = st.columns(3)
                    with c1:
                        st.markdown(kpi_box(n(tot), "всего штук за период", f"недель: {len(wk)}", "all"),
                                    unsafe_allow_html=True)
                    with c2:
                        st.markdown(kpi_box(n(wk["Штук"].mean()), "в среднем за неделю, шт", "", "ok"),
                                    unsafe_allow_html=True)
                    with c3:
                        st.markdown(kpi_box(n(last_w["Штук"]), f"последняя неделя ({last_w['Неделя']})", sub3,
                                            "left" if len(wk) > 1 and last_w["Δ"] < 0 else "ok"),
                                    unsafe_allow_html=True)

                    g = (part.groupby(["W", "Статус"], as_index=False)["Штук"].sum())
                    g["Неделя"] = (g["W"].dt.strftime("%d.%m") + "–" +
                                   (g["W"] + pd.Timedelta(days=6)).dt.strftime("%d.%m"))
                    xenc = alt.X("Неделя:N", sort=order, title=None, axis=alt.Axis(labelAngle=-30))
                    bars = (alt.Chart(g).mark_bar(size=38)
                            .encode(x=xenc,
                                    y=alt.Y("Штук:Q", title="Штук", axis=alt.Axis(format=",.0f")),
                                    color=alt.Color("Статус:N", title=None, scale=colors,
                                                    legend=alt.Legend(orient="top")),
                                    tooltip=[alt.Tooltip("Неделя:N"), alt.Tooltip("Статус:N"),
                                             alt.Tooltip("Штук:Q", title="Штук", format=",.0f")]))
                    labels = (alt.Chart(wk).mark_text(dy=-8, fontWeight="bold")
                              .encode(x=xenc, y=alt.Y("Штук:Q"), text=alt.Text("Штук:Q", format=",.0f")))
                    chart = (bars + labels).properties(height=400)

                try:
                    st.altair_chart(chart, width="stretch")
                except TypeError:
                    st.altair_chart(chart, use_container_width=True)

                if weekly:
                    def _chg(r):
                        if pd.isna(r["Δ"]):
                            return "—"
                        arrow = "▲" if r["Δ"] > 0 else ("▼" if r["Δ"] < 0 else "■")
                        pct = f" ({r['Δ%']:+.0f}%)" if pd.notna(r["Δ%"]) else ""
                        return f"{arrow} {r['Δ']:+,.0f}{pct}".replace(",", " ")
                    rows = [[f"<b>{r.Неделя}</b>", n(r.Штук), n(r.Собрано),
                             f"{(r.Собрано / r.Штук * 100 if r.Штук else 0):.0f}%", _chg(wk.loc[i])]
                            for i, r in wk.iterrows()]
                    st.markdown(simple_table(["Неделя", "Штук", "Собрано", "Доля собранного", "К прошлой неделе"],
                                             rows[::-1], num_cols={1, 2, 3}), unsafe_allow_html=True)
                    st.caption("Неделя считается с понедельника по воскресенье. Первая и последняя недели периода "
                               "могут быть неполными: для честного сравнения выбирайте период от понедельника до воскресенья.")
                else:
                    st.caption("Столбец = один день. Зелёная часть — собранные штуки, оранжевая — не собранные "
                               "(включая новые не маршрутизированные заказы). Наведите курсор на столбец, чтобы увидеть число.")
